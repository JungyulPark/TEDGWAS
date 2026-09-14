from pathlib import Path
import re,json,hashlib,sys
import numpy as np,pandas as pd
from docx import Document
O=Path(__file__).resolve().parents[1];P=O/'provenance';master=O.parents[1]/'MANUSCRIPT_TED_TRAP_v5_MASTER.md';text=master.read_text(encoding='utf-8')
mr=pd.read_csv(P/'MR_primary_canonical.csv');raw=pd.read_csv(P/'MR_all_estimators_verified.csv');co=pd.read_csv(P/'coloc_canonical_v2.csv');inst=pd.read_csv(P/'instruments_verified.csv')
allmr=pd.read_csv(O/'Supplementary_Data_2_MR.csv');allco=pd.read_csv(O/'Supplementary_Data_3_Colocalization.csv');power=pd.read_csv(O/'Supplementary_Data_4_Power.csv')
display=json.loads((P/'clinical_display_sources.json').read_text(encoding='utf-8'));tables=display['tables'];checks=[];errors=[];numeric=0
def check(ok,label):
    checks.append(label)
    if not ok:errors.append(label)
def clean(s):return re.sub(r'\s+',' ',str(s).replace('*','').replace('\u00a0',' ')).strip()
tr=str.maketrans('⁻⁰¹²³⁴⁵⁶⁷⁸⁹−','-0123456789-')
def agree(shown,value,label):
    global numeric
    numeric+=1;s=clean(shown).replace(',','').translate(tr)
    if s=='NA':check(pd.isna(value),label+' unavailable');return
    if s.startswith('<'):check(np.isfinite(value) and 0<=value<float(s[1:]),label+' below display threshold');return
    if '×10' in s:
        mant,exp=s.split('×10');actual=float(mant)*10**int(exp);unit=10**int(exp)*10**(-len(mant.split('.')[1]) if '.' in mant else 0)
    else:
        actual=float(s);unit=10**(-len(s.split('.')[1]) if '.' in s else 0)
    check(np.isfinite(value) and abs(actual-value)<=unit/2+1e-12,label)
def effect(s,r,label):
    v=re.findall(r'\d+\.\d+',s);check(len(v)==3,label+' three OR/CI values')
    for a,b in zip(v,np.exp([r.beta,r.beta-1.96*r.se,r.beta+1.96*r.se])):agree(a,b,label)
ocs={'BBJ Graves disease':'BBJ_Graves','UKB hyperthyroidism':'UKB_hyperthyroid','FinnGen Graves ophthalmopathy':'FinnGen_GO'}
def get(g,o):return mr[(mr.gene_symbol==g)&(mr.outcome==o)].iloc[0]
for name,keys in [('MANUSCRIPT_Submission',['1','2','3']),('SUPPLEMENTARY_MATERIAL',['S1','S2','S3','S4','S5'])]:
    doc=Document(O/(name+'.docx'));check(len(doc.tables)==len(keys),name+' table count')
    for tb,key in zip(doc.tables,keys):
        expected=[tables[key]['header']]+tables[key]['rows'];check(len(tb.rows)==len(expected),key+' row count')
        for row,ex in zip(tb.rows,expected):
            check(len(row.cells)==len(ex),key+' column count')
            for c,e in zip(row.cells,ex):check(clean(c.text)==clean(e),key+' DOCX cell '+clean(e))
            check(bool(row._tr.xpath('./w:trPr/w:cantSplit')),key+' row does not split')
        check(bool(tb.rows[0]._tr.xpath('./w:trPr/w:tblHeader')),key+' repeating header')
        for ex in expected:check('| '+' | '.join(map(str,ex))+' |' in text,key+' source row exists')
    alltext='\n'.join([p.text for p in doc.paragraphs]+[c.text for tb in doc.tables for row in tb.rows for c in row.cells])
    check('�' not in alltext,name+' no replacement characters')
    check(not re.search(r'\bP\s*=\s*0\.000(?:\s|[;,)]|$)',alltext),name+' no zero P value display')
for row in tables['2']['rows']:
    g=clean(row[0]);o=ocs[row[1]];r=get(g,o);label=f'Table2 {g}/{o}'
    agree(row[2],r.n_iv,label+' instruments');effect(row[3],r,label);agree(row[4],r.pvalue,label+' P')
    c=co[(co.gene==g)&(co.outcome==o)&np.isclose(co.p12,1e-5,atol=1e-12)].iloc[0]
    agree(row[5],c['PP.H2'],label+' H2');agree(row[6],c['PP.H3'],label+' H3');agree(row[7],c['PP.H4'],label+' H4')
for row in tables['3']['rows']:
    r=get(clean(row[0]),'BBJ_Graves');agree(row[1],r.n_iv,'Table3 instruments');effect(row[2],r,'Table3 effect');agree(row[3],r.pvalue,'Table3 P');check(r.pvalue<.05/2544,'Table3 discovery threshold')
check(len(tables['3']['rows'])==13,'All thirteen discovery hits retained')
for row in tables['S1']['rows']:
    g=clean(row[0]);o=ocs[row[1]];r=get(g,o);rr=raw[(raw.gene_symbol==g)&(raw.outcome==o)]
    agree(row[2],r.n_iv,'S1 instruments');agree(row[3],r.pvalue,'S1 primary P')
    for col,method in [(4,'Weighted median'),(5,'Weighted mode')]:
        v=rr[rr.method==method];agree(row[col],v.iloc[0].pvalue if len(v) else np.nan,'S1 '+method)
    agree(row[6],r.egger_intercept_p,'S1 Egger intercept');agree(row[7],r.cochran_q_p,'S1 Q')
for row in tables['S2']['rows']:
    g=clean(row[0]);o=ocs[row[1]]
    for col,p12 in [(2,1e-5),(3,5e-6),(4,1e-6)]:
        c=co[(co.gene==g)&(co.outcome==o)&np.isclose(co.p12,p12,atol=1e-12)].iloc[0];agree(row[col],c['PP.H4'],'S2 H4')
for row in tables['S3']['rows']:
    g=clean(row[0]);o=ocs[row[1]];r=get(g,o);new=allmr[(allmr.gene_symbol==g)&(allmr.outcome==o)&(allmr.scenario=='reharmonized_eqtlgen')].iloc[0]
    effect(row[2],r,'S3 reference');agree(row[3],r.pvalue,'S3 reference P');effect(row[4],new,'S3 cohort');agree(row[5],new.pvalue,'S3 cohort P')
for row in tables['S4']['rows']:
    o=ocs[row[0]];scenario='original_reference' if row[1]=='Reference' else 'reharmonized_eqtlgen';r=power[(power.outcome==o)&(power.scenario==scenario)].iloc[0]
    agree(row[2],r.n_genes,'S4 genes')
    for a,b in zip(re.findall(r'\d+\.\d+',row[3]),[r.or_median,r.or_q1,r.or_q3]):agree(a,b,'S4 detectable OR')
    agree(row[4],100*r.frac_OR1_5,'S4 OR1.5 percentage');agree(row[5],100*r.frac_OR2,'S4 OR2 percentage')
# Confirm the consolidated data retain the old primary results and full scenarios.
ref=allmr[allmr.scenario=='original_reference'].merge(mr,on=['gene_symbol','outcome'],suffixes=('_new','_old'),validate='one_to_one');check(len(ref)==7219,'All original MR estimates in consolidated file')
for col in ['beta','se','pvalue','n_iv']:check(np.allclose(ref[col+'_new'],ref[col+'_old'],atol=1e-12,rtol=1e-12),col+' primary data unchanged')
check(len(allmr)==30115,'Complete merged MR row count');check(set(allmr.scenario)=={'original_reference','paired_reference','paired_eqtlgen','reharmonized_eqtlgen'},'All MR scenarios retained')
check(len(allco)==324 and set(allco.scenario)=={'original_reference','paired_reference','paired_eqtlgen','available_eqtlgen'},'All coloc rows/scenarios retained')
check(len(power)==12,'All power scenarios retained')
check(inst.gene_symbol.nunique()==2544 and len(inst)==6135,'Instrument manifest unchanged')
refco=allco[allco.scenario=='original_reference'].merge(co,on=['gene','outcome','p12'],suffixes=('_new','_old'),validate='one_to_one');check(len(refco)==81,'All original coloc settings retained')
for k in range(5):check(np.allclose(refco[f'PP.H{k}_new'],refco[f'PP.H{k}_old'],atol=1e-12,rtol=1e-12),'Original H'+str(k)+' unchanged')
v=json.loads((P/'maf/sensitivity_verification_v2.json').read_text(encoding='utf-8'));check(v['status']=='PASS' and v['max_abs_delta_Z']<1e-8,'Completed sensitivity verification')
check(all(r['n_threshold_crossings']==0 for r in v['coloc_summary']['comparisons']),'Zero frequency threshold crossings retained')
for r in display['narrative_estimates']:
    m=get(r['gene'],r['outcome']);agree(r['P'],m.pvalue,'Narrative P');agree(r['OR'],np.exp(m.beta),'Narrative OR')
    check(r['P'] in text and r['OR'] in text and r['CI'] in text,'Narrative rendered values present')
fs=json.loads((P/'clinical_figure_sources.json').read_text(encoding='utf-8'))
for r in fs['Figure2']:
    m=get(r['gene'],r['outcome']);check(np.isclose(r['pvalue'],m.pvalue,rtol=1e-12,atol=0),'Figure2 source P');check(np.isclose(r['or'],np.exp(m.beta)),'Figure2 source OR');agree(r['displayed_p'],m.pvalue,'Figure2 display P')
check(not re.search(r'\b(?:Table|Tables) S(?:[6-9]|1[0-2])\b',text),'No old supplementary table references')
check('Figure 4.' not in text and 'Figure S3.' not in text and 'Figure S2.' in text,'No removed figure references')
check('{{' not in text and 'TODO' not in text,'No unresolved generation placeholders')
check(not any(q in text for q in ['AI-assisted','artificial intelligence','language model','Codex','OpenAI','Claude','ChatGPT']),'Remote author decision: no manuscript AI-tool declaration')
check(not any('prespecified' in line and 'outcome hierarchy' not in line for line in text.splitlines()),'No unregistered threshold described as prespecified')
check('Vattikuti S, Vattikuti S' not in text,'Duplicated reference author corrected')
check('All authors read and approved the final manuscript.' not in text,'No claim that this new revision has already been approved')
plain=text.replace('*','')
check('not a P value' in plain and 'not treatment effects' in plain,'Probability and treatment-effect interpretation retained')
check('P* = 0.0056' in text,'TNFSF14 UKB nominal result retained')
# Independently check current text rather than trusting historical metadata.
import runpy
from statistics import NormalDist
wc=runpy.run_path(str(master.parent/'scripts/26_wordcount_main_text.py'))
sections=wc['split_sections'](text)
actual_counts={'abstract_words':wc['count_words'](sections['Abstract']),
 'main_words_excluding_headings':sum(wc['count_words'](sections[s]) for s in wc['MAIN']),
 'main_words_including_headings':sum(wc['count_words'](sections[s],headings=True) for s in wc['MAIN'])}
saved_counts=json.loads((P/'word_counts.json').read_text(encoding='utf-8'))
for key,value in actual_counts.items():check(saved_counts.get(key)==value,'Current word-count record: '+key)
before,rest=text.split('## References',1);refblock,after=rest.split('## Figure Legends',1)
references={int(n):s.strip() for n,s in re.findall(r'(?m)^(\d+)\. (.+)$',refblock)}
order=[]
for match in re.finditer(r'\[([\d,–\- ]+)\]',before+'\n'+after):
    for part in match.group(1).split(','):
        ends=re.split('[–-]',part.strip())
        for n in range(int(ends[0]),int(ends[-1])+1):
            check(n in references,'Citation in reference range')
            if n not in order:order.append(n)
check(order==list(range(1,len(references)+1)),'References follow first appearance throughout manuscript')
verified=json.loads((P/'references_verified.json').read_text(encoding='utf-8'))
check(len(verified)==len(references),'Reference verification count current')
for r in verified:check(r['reference']==str(r['number'])+'. '+references.get(r['number'],''),'Reference verification text matches '+str(r['number']))
main_doc=Document(O/'MANUSCRIPT_Submission.docx')
doc_ref_text='\n'.join(p.text for p in main_doc.paragraphs)
for line in references.values():check(clean(line).replace('’', "'") in clean(doc_ref_text).replace('’', "'"),'DOCX reference matches current master')
check('does not standardize the genetic predictor itself' in text,'Standardized expression scale distinguished from genetic predictor')
check('comparing them with observed effects cannot distinguish an absent effect from limited power' in text,'No inference from observed effect versus detection limit')
for gene,outcome,shown in [('TSHR','BBJ_Graves','1.39'),('IGF1R','BBJ_Graves','0.99'),('IGF1R','UKB_hyperthyroid','0.33'),('IGF1R','FinnGen_GO','0.72')]:
    a=.05/2544 if outcome=='BBJ_Graves' else .05
    limit=(NormalDist().inv_cdf(1-a/2)+NormalDist().inv_cdf(.8))*get(gene,outcome).se
    agree(shown,limit,'Gene-specific 80% detection limit '+gene+'/'+outcome)
check('BBJ (α = 0.05/2,544)' in text and 'respectively (α = 0.05)' in text,'Outcome-specific detection thresholds disclosed')
# Expanded Results: quantitative candidate comparisons and complete screen counts.
bbj=mr[mr.outcome=='BBJ_Graves'];hits=bbj[bbj.pvalue<.05/2544]
check(sum(hits.beta<0)==7 and sum(hits.beta>0)==6,'Seven lower-odds and six higher-odds discovery genes')
for g,o,shown in [('TNFSF14','BBJ_Graves','0.994'),('IFNGR1','BBJ_Graves','0.989'),('TNFSF14','FinnGen_GO','0.017'),('IFNGR1','FinnGen_GO','0.020')]:
    r=co[(co.gene==g)&(co.outcome==o)&np.isclose(co.p12,1e-5,atol=1e-12)].iloc[0]
    agree(shown,r['PP.H4'],'Expanded candidate narrative '+g+'/'+o)
    check(shown in text,'Expanded candidate probability present')
disc=text.split('## Discussion',1)[1].split('## Declarations',1)[0]
lim=disc.split('### Limitations',1)[1].split('### Conclusions',1)[0]
check(len([p for p in lim.strip().split('\n\n') if p.strip()])==2,'Two complete limitation paragraphs')
check(len([p for p in disc.split('### Limitations')[0].strip().split('\n\n') if p.strip()])==7,'Seven interpretation paragraphs before limitations')
counts=json.loads((P/'word_counts.json').read_text(encoding='utf-8'));check(counts['abstract_words']<=250,'Abstract within working 250-word ceiling');check(counts['main_words_including_headings']<=5000,'Main within working 5000-word ceiling')
check('modest inherited contribution' not in text,'No unsupported restriction on inherited effect size')
for name,keys in [('MANUSCRIPT_Submission',['1','2','3']),('SUPPLEMENTARY_MATERIAL',['S1','S2','S3','S4','S5'])]:
    source=Document(O/(name+'.docx'))
    for index,number in enumerate(keys):
        upload=Document(O/'tables'/f'Table{number}.docx')
        check(len(upload.tables)==1,'Separate Table '+number+' exists')
        check([[c.text for c in r.cells] for r in upload.tables[0].rows]==[[c.text for c in r.cells] for r in source.tables[index].rows],'Separate Table '+number+' matches manuscript')
        for p in upload.paragraphs:
            check(not p.paragraph_format.page_break_before,'Separate Table '+number+' no leading page break')
# The complete estimator supplement retains coefficients/P values and discloses
# test-specific intervals; do not infer P values from rounded displayed CIs.
extra=pd.read_csv(O/'Supplementary_Data_5_MR_estimators.csv')
check(len(extra)==len(raw)==13039,'Complete estimator supplement row count')
joined=extra.merge(raw,on=['gene_symbol','outcome','method'],suffixes=('_new','_raw'),validate='one_to_one')
for col in ['beta','se','pvalue','n_iv']:
    check(np.allclose(joined[col+'_new'],joined[col+'_raw'],rtol=1e-12,atol=0),'All estimator '+col+' preserved')
from scipy import stats
for method, group in extra.groupby('method'):
    dfs=group.n_iv-(2 if method=='MR Egger' else 1)
    is_t=method in ['MR Egger','Weighted mode']
    prob=2*(stats.t.sf(abs(group.beta/group.se),dfs) if is_t else stats.norm.sf(abs(group.beta/group.se)))
    critical=stats.t.ppf(.975,dfs) if is_t else stats.norm.ppf(.975)
    check(np.allclose(prob,group.pvalue,rtol=1e-6,atol=1e-300),'All '+method+' P values match original test')
    check(np.allclose(group.log_or_ci_lower,group.beta-critical*group.se),'All '+method+' lower CIs use matching test')
    check(np.allclose(group.log_or_ci_upper,group.beta+critical*group.se),'All '+method+' upper CIs use matching test')
    check((group.test_distribution==('Student t' if is_t else 'normal')).all(),'All '+method+' distributions labelled')
for scenario, group in allmr.groupby('scenario'):
    valid=group.dropna(subset=['beta','se','pvalue'])
    prob=2*stats.norm.sf(abs(valid.beta/valid.se))
    check(np.allclose(prob,valid.pvalue,rtol=1e-6,atol=1e-300),'All primary/scenario P values: '+scenario)
check(np.allclose(allco[[f'PP.H{i}' for i in range(5)]].sum(axis=1),1,rtol=0,atol=1e-12),'Every posterior row sums to one')
check((allco[[f'PP.H{i}' for i in range(5)]].values>=0).all(),'All posterior probabilities non-negative')
check('top_snp field' in text and 'conditional on H4' in text,'Conditional top-SNP definition disclosed')
check('exact rule preceded inspection of results' in text,'Combined-filter timing uncertainty disclosed')
evidence=json.loads((P/'major_review_evidence.json').read_text(encoding='utf8'))
check(evidence['status']=='PASS','Local source verification passed')
for record in evidence['TSHR_scale_check']:
    for key in ['variant_log_or','variant_or','exposure_beta']:
        shown=f"{record[key]:.3f}" if key=='variant_or' else f"{record[key]:.5f}"
        agree(shown,record[key],'TSHR source scale '+key)
        check(shown.replace('-','−') in text,'TSHR source scale printed '+shown)
for record in evidence['TSHR_reference_LD']:
    if 'rs179252' in [record['SNP_A'],record['SNP_B']]:
        shown=f"{record['R2']:.6f}";agree(shown,record['R2'],'Reference LD');check(shown in text,'Reference LD printed')
check(evidence['FinnGen_H2_dominant_loci']==7,'Seven H2-dominant FinnGen loci')
check(evidence['CTLA4_FinnGen_H4']>=.8,'CTLA4 failure does not originate in FinnGen')
# Independently verified post hoc leave-one-out results: report every omission.
loo=pd.read_csv(O/'Supplementary_Data_6_Leave_one_out.csv')
rv=pd.read_csv(P/'leave_one_out_R_validation.csv')
lv=json.loads((P/'leave_one_out_verification.json').read_text(encoding='utf8'))
check(lv['status']=='PASS','LOO independent R verification passed')
for name,key in [('Supplementary_Data_6_Leave_one_out.csv','data6_sha256'),('provenance/leave_one_out_R_validation.csv','R_validation_sha256'),('provenance/posthoc_20260914/leave_one_out_estimates.csv','source_sha256')]:
    check(hashlib.sha256((O/name).read_bytes()).hexdigest()==lv[key],'LOO verified source hash '+name)
check(len(loo)==20 and len(rv)==24,'LOO complete exported and independently verified counts')
joined=loo.merge(rv,on=['gene','outcome','excluded_SNP'],validate='one_to_one',suffixes=('_py','_R'))
check(len(joined)==20,'All displayed LOO results independently matched')
for col in ['n_iv','beta','se','pvalue','OR','CI_lower','CI_upper']:
    check(np.allclose(joined[col+'_py'],joined[col+'_R'],rtol=1e-10,atol=1e-300),'LOO R/Python '+col)
check(np.allclose(2*stats.norm.sf(abs(loo.beta/loo.se)),loo.pvalue,rtol=1e-10,atol=1e-300),'LOO normal P values verified')
check(np.allclose(np.exp(loo.beta-stats.norm.ppf(.975)*loo.se),loo.CI_lower,rtol=1e-10),'LOO lower CIs verified')
check(np.allclose(np.exp(loo.beta+stats.norm.ppf(.975)*loo.se),loo.CI_upper,rtol=1e-10),'LOO upper CIs verified')
omissions=loo[loo.excluded_SNP!='None (all instruments)']
check(len(omissions)==15 and sum(omissions.gene=='IGF1R')==11 and sum(omissions.gene=='CTLA4')==4,'LOO 15 = 11 IGF1R + 4 CTLA4')
check((omissions.n_iv==omissions.original_n_iv-1).all(),'Exactly one SNP omitted per row')
check((loo.frequency_scenario=='original_reference').all(),'LOO frequency scenario disclosed')
check((omissions.loc[omissions.gene=='IGF1R','beta']>0).all(),'All eleven IGF1R omission directions retained')
brief={'BBJ':'BBJ_Graves','UKB':'UKB_hyperthyroid','FinnGen':'FinnGen_GO'}
seen=set()
for row in tables['S5']['rows']:
    key=(clean(row[0]),brief[row[1]],row[2]);seen.add(key)
    found=omissions[(omissions.gene==key[0])&(omissions.outcome==key[1])&(omissions.excluded_SNP==key[2])]
    check(len(found)==1,'S5 unique omitted SNP '+str(key));r=found.iloc[0]
    agree(row[3],r.n_iv,'S5 remaining SNPs')
    check(row[4]==('Wald ratio' if r.n_iv==1 else 'IVW'),'S5 estimator matches remaining count')
    for shown,value in zip(re.findall(r'\d+\.\d+',row[5]),[r.OR,r.CI_lower,r.CI_upper]):agree(shown,value,'S5 OR/CI '+str(key))
    agree(row[6],r.pvalue,'S5 P '+str(key))
check(seen==set(zip(omissions.gene,omissions.outcome,omissions.excluded_SNP)),'Table S5 includes every omission exactly once')
for outcome,pv,orr,lo,hi in [('BBJ_Graves','0.0751','1.49','0.96','2.30'),('UKB_hyperthyroid','0.468','1.16','0.78','1.72')]:
    r=omissions[(omissions.gene=='IGF1R')&(omissions.outcome==outcome)&(omissions.excluded_SNP=='rs2654980')].iloc[0]
    for shown,value in [(pv,r.pvalue),(orr,r.OR),(lo,r.CI_lower),(hi,r.CI_upper)]:
        agree(shown,value,'IGF1R LOO narrative');check(shown in text,'LOO narrative value present')
    check(sum(omissions[(omissions.gene=='IGF1R')&(omissions.outcome==outcome)].pvalue>=.05)==2,'Two IGF1R omissions lose nominal significance '+outcome)
for outcome,orr,lo,hi in [('UKB_hyperthyroid','1.02','0.37','2.82'),('FinnGen_GO','1.72','0.16','18.90')]:
    r=omissions[(omissions.gene=='CTLA4')&(omissions.outcome==outcome)&(omissions.excluded_SNP=='rs13030124')].iloc[0]
    for shown,value in [(orr,r.OR),(lo,r.CI_lower),(hi,r.CI_upper)]:
        agree(shown,value,'CTLA4 LOO narrative');check(shown in text,'CTLA4 narrative value present')
check('European *CTLA4* MR support was concentrated in rs13030124' in text,'Table 2 CTLA4 dependence disclosed')
check('This analysis was not repeated under substituted eQTLGen frequencies' in text,'LOO scope limitation disclosed')
check('Supplementary Data 1–6' in text.split('**Data availability.**',1)[1].split('**Author contributions.**',1)[0],'Data availability includes Data 6')
check('Multi-signal colocalization was not performed' in text,'Multi-signal analysis explicitly unperformed')
checktext=(O/'STROBE_MR_CHECKLIST.md').read_text(encoding='utf8')
check('no leave-one-out result is reported' not in checktext and 'Table S5' in checktext and 'Figure S2' in checktext,'STROBE updated to report LOO')

checklist=Document(O/'STROBE_MR_CHECKLIST.docx');check(len(checklist.tables)==1,'Separate reporting checklist exists')
result={'status':'PASS' if not errors else 'FAIL','numeric_cells_compared':numeric,'checks':len(checks),'errors':errors,'scope':'Rounded values versus original analytical outputs, complete consolidated data, DOCX table transcription and essential limitations. Visual layout, author declarations and live journal requirements require separate review.'}
(P/'integrity_audit.json').write_text(json.dumps(result,indent=2),encoding='utf-8',newline='\r\n');print(json.dumps(result,indent=2));sys.exit(bool(errors))

from pathlib import Path
import re,subprocess,tempfile,copy,shutil,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
O=Path(__file__).resolve().parents[1]
master=O.parents[1]/'MANUSCRIPT_TED_TRAP_v5_MASTER.md'
t=master.read_text(encoding='utf-8');main,sup=t.split('## Supplementary Material',1)
title=re.search(r'^# (.+)$',t,re.M).group(1)
texts={'MANUSCRIPT_Submission':main,'SUPPLEMENTARY_MATERIAL':'# Supplementary material\n\n'+title+'\n\n'+sup,'COVER_LETTER_EndocrineConnections':(O/'COVER_LETTER_EndocrineConnections.md').read_text(encoding='utf-8'),'STROBE_MR_CHECKLIST':(O/'STROBE_MR_CHECKLIST.md').read_text(encoding='utf-8')}
if sys.argv[1:]:
    assert set(sys.argv[1:]).issubset(texts), 'Unknown document name'
    texts={name:texts[name] for name in sys.argv[1:]}
pandoc=shutil.which('pandoc')
if not pandoc:
    import pypandoc
    pandoc=pypandoc.get_pandoc_path()
def setsec(sec,land=False,lines=False):
    sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
    sec.page_width=Inches(11.69 if land else 8.27);sec.page_height=Inches(8.27 if land else 11.69)
    sec.left_margin=sec.right_margin=Inches(.7 if land else .95)
    sec.top_margin=sec.bottom_margin=Inches(.65 if land else .8)
    sec.header_distance=sec.footer_distance=Inches(.3)
    for x in sec._sectPr.findall(qn('w:lnNumType')):sec._sectPr.remove(x)
    if lines:
        x=OxmlElement('w:lnNumType');x.set(qn('w:countBy'),'1');x.set(qn('w:restart'),'continuous');x.set(qn('w:distance'),'240');sec._sectPr.append(x)
def boundary(doc,p,lines):
    # The inserted section describes the portrait section ending here.
    new=OxmlElement('w:p');pr=OxmlElement('w:pPr');sp=copy.deepcopy(doc.sections[-1]._sectPr)
    for x in sp.findall(qn('w:type')):sp.remove(x)
    typ=OxmlElement('w:type');typ.set(qn('w:val'),'nextPage');sp.append(typ)
    pr.append(sp);new.append(pr);p._p.addprevious(new)
with tempfile.TemporaryDirectory(prefix='tedtrap-documents-') as scratch:
    for name,text in texts.items():
        md=Path(scratch)/(name+'.md');md.write_text(re.sub(r'(?m)^---\s*$','',text),encoding='utf-8')
        out=O/(name+'.docx');subprocess.run([pandoc,str(md),'-f','markdown+pipe_tables+tex_math_dollars','-t','docx','-o',str(out)],check=True)
        d=Document(out);ismain=name=='MANUSCRIPT_Submission';issup=name=='SUPPLEMENTARY_MATERIAL';iscover=name.startswith('COVER');ischeck=name=='STROBE_MR_CHECKLIST'
        d.core_properties.title=title if ismain else name.replace('_',' ');d.core_properties.identifier=name;d.core_properties.author='Jungyul Park; Min-Seon Kim; Kyung-Hwa Shin; Suk-Woo Yang';d.core_properties.last_modified_by=''
        for sec in d.sections:setsec(sec,ischeck,ismain)
        for nm in ['Normal','Body Text','First Paragraph','Compact','Caption','Table','List Paragraph']:
            if nm not in d.styles:continue
            st=d.styles[nm];st.font.name='Times New Roman';st.font.size=Pt(11 if iscover else 12)
            st.paragraph_format.line_spacing=1 if iscover else 2
            st.paragraph_format.space_after=Pt(5 if iscover else 0)
        for nm in ['Title','Heading 1','Heading 2','Heading 3','Heading 4']:
            if nm not in d.styles:continue
            st=d.styles[nm];st.font.name='Times New Roman';st.font.size=Pt(15 if nm=='Title' else 12);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
        intables=False
        for p in list(d.paragraphs):
            p.paragraph_format.widow_control=True
            if p.style.name=='Heading 1' and p.text in [title,'Supplementary material','Cover letter','STROBE MR reporting checklist']:p.style='Title'
            if p.text in ['Abstract','Introduction','References','Figure Legends']:p.paragraph_format.page_break_before=True
            if p.text=='Tables':p._element.getparent().remove(p._element);continue
            if re.match(r'^Table (S?[1-4])\.',p.text):
                if (ismain and p.text.startswith('Table 1.')) or (issup and p.text.startswith('Table S1.')):boundary(d,p,ismain)
                else:p.paragraph_format.page_break_before=True
                intables=True;p.paragraph_format.keep_with_next=True
            if intables or ischeck:
                p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(7)
                for r in p.runs:r.font.size=Pt(10)
            if p.style.name.startswith('List') or re.match(r'^\d+\. ',p.text):p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(6)
            if iscover and p.text.startswith(('Sincerely','Suk-Woo Yang')):p.paragraph_format.keep_with_next=True
            if iscover and p.text=='Cover letter':p.paragraph_format.space_before=Pt(0)
        if ismain or issup:setsec(d.sections[-1],True,False)
        for old in list(d.tables):
            fresh=d.add_table(rows=len(old.rows),cols=len(old.columns))
            for sr,tr in zip(old.rows,fresh.rows):
                for sc,tc in zip(sr.cells,tr.cells):
                    for el in list(tc._tc):
                        if el.tag!=qn('w:tcPr'):tc._tc.remove(el)
                    for sp in sc.paragraphs:tc._tc.append(copy.deepcopy(sp._p))
            old._tbl.addprevious(fresh._tbl);old._tbl.getparent().remove(old._tbl)
        for table in d.tables:
            heads=[c.text for c in table.rows[0].cells];cols=len(heads);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
            if cols==3:weights=[.5,1.8,7.99]
            elif cols==8 and heads[3]=='OR (95% CI)':weights=[.7,2.0,.95,1.45,1.0,1.15,1.15,1.1]
            elif cols==8:weights=[.8,2.3,.9,.85,.85,.85,1.15,1.1]
            elif cols==5 and heads[0]=='Dataset':weights=[2.7,1.8,.8,.9,3.0]
            elif cols==5 and heads[1]=='Instruments':weights=[1.0,.8,2.4,1.05,3.6]
            elif cols==5:weights=[.9,3,1.6,1.6,1.6]
            elif cols==6 and heads[0]=='Outcome':weights=[2.4,1.45,.8,2.25,1.1,1.1]
            elif cols==6 and heads[2]=='Instruments':weights=[.8,2.8,.85,2.1,1.05,1.4]
            elif cols==6:weights=[.85,2.45,2.15,1,2.15,1]
            else:weights=[1]*cols
            widths=[10.29*x/sum(weights) for x in weights]
            tw=table._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(10.29*1440)))
            for col,width in zip(table.columns,widths):col.width=Inches(width)
            for ri,row in enumerate(table.rows):
                pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if ri==0:pr.append(OxmlElement('w:tblHeader'))
                for ci,c in enumerate(row.cells):
                    c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cp=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                    for side,val in [('top','80'),('bottom','80'),('left','80'),('right','80')]:x=OxmlElement('w:'+side);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');m.append(x)
                    cp.append(m)
                    if ri==0:x=OxmlElement('w:shd');x.set(qn('w:fill'),'EEEEEE');cp.append(x)
                    for p in c.paragraphs:
                        p.style='Normal';p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=False
                        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ischeck or ci<2 or (cols==5 and ci==4) else WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:
                            r.font.name='Times New Roman';r.font.size=Pt(9.5 if cols==8 or ischeck else 10)
                            if ri==0:r.bold=True
                            elif re.search(r'\bP\b',heads[ci]) and 'PP.H' not in heads[ci]:
                                value=c.text.translate(str.maketrans('⁻⁰¹²³⁴⁵⁶⁷⁸⁹','-0123456789'))
                                try:
                                    bits=value.split('×10');pv=float(bits[0])*(10**int(bits[1]) if len(bits)==2 else 1)
                                    r.bold=pv<.05
                                except ValueError:r.bold=False
        if issup:
            for p in list(d.paragraphs):
                if p.text.startswith('Figure S1.'):
                    p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
                    np=d.add_paragraph();p._p.addnext(np._p);np.add_run().add_picture(str(O/'figures/FigureS1.png'),width=Inches(8.9));np.alignment=WD_ALIGN_PARAGRAPH.CENTER;np.paragraph_format.line_spacing=1
        seen=set()
        for sec in d.sections:
            for foot in [sec.footer,sec.first_page_footer,sec.even_page_footer]:
                if foot.part.partname in seen:continue
                seen.add(foot.part.partname)
                for x in list(foot._element):foot._element.remove(x)
                p=foot.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p._p.get_or_add_pPr().append(OxmlElement('w:suppressLineNumbers'))
                for typ in ['begin','instruction','end']:
                    rr=p.add_run()
                    if typ=='instruction':el=OxmlElement('w:instrText');el.text=' PAGE '
                    else:el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),typ)
                    rr._r.append(el)
        order=['headerReference','footerReference','footnotePr','endnotePr','type','pgSz','pgMar','paperSrc','pgBorders','lnNumType','pgNumType','cols','formProt','vAlign','noEndnote','titlePg','textDirection','bidi','rtlGutter','docGrid','printerSettings']
        for sec in d.sections:
            els=list(sec._sectPr)
            for el in els:sec._sectPr.remove(el)
            for el in sorted(els,key=lambda x:order.index(x.tag.split('}')[-1]) if x.tag.split('}')[-1] in order else 99):sec._sectPr.append(el)
        d.save(out);print(name,'tables',len(d.tables))

# Separate editable tables for the publisher's upload workflow.
(O/'tables').mkdir(exist_ok=True)
for source,keys in [('MANUSCRIPT_Submission',['1','2','3']),('SUPPLEMENTARY_MATERIAL',['S1','S2','S3','S4'])]:
    if source not in texts:continue
    for index,number in enumerate(keys):
        table_doc=Document(O/(source+'.docx'))
        target=table_doc.tables[index]._tbl
        keep=[target.getprevious(),target,target.getnext(),table_doc._element.body.sectPr]
        for el in list(table_doc._element.body):
            if el not in keep:table_doc._element.body.remove(el)
        setsec(table_doc.sections[-1],True,False)
        for p in table_doc.paragraphs:p.paragraph_format.page_break_before=False
        table_doc.core_properties.title=f'Table {number}'
        table_doc.core_properties.identifier=f'Table{number}'
        table_doc.save(O/'tables'/f'Table{number}.docx')

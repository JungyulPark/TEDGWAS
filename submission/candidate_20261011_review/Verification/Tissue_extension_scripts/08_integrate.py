from pathlib import Path
import json,hashlib,shutil,re,difflib,copy,sys
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1];A=R/'outputs/TED_Tissue_Extension_20261003';S=R/'outputs/TED_Allele_Integrated_20261003';O=R/'outputs/TED_Tissue_Integrated_20261003'
if O.exists():assert (O/'Manuscript.docx').read_bytes()==(S/'Manuscript.docx').read_bytes(),'Preserve completed review packages'
O.mkdir(exist_ok=True);V=O/'Verification';V.mkdir(exist_ok=True)
for p in S.iterdir():
 if p.name in ['Verification','PACKAGE_SHA256.json','Full_Review.pdf','READ_FIRST_KO.md','Research_planning','Research_Question_and_Contribution_KO.md']:continue
 if p.is_dir():shutil.copytree(p,O/p.name,dirs_exist_ok=True)
 else:shutil.copy2(p,O/p.name)
shutil.copy2(S/'PACKAGE_SHA256.json',V/'Previous_package_SHA256.json')
shutil.copytree(A/'Protocol',O/'Research_planning',dirs_exist_ok=True)
review=O/'Additional_analysis_figure';review.mkdir(exist_ok=True)
for p in (A/'Figures').iterdir():shutil.copy2(p,review/p.name)
def plain(s):return s.replace('*','')
def settext(p,t):
 p.clear()
 for s in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',t):
  if not s:continue
  bold=s.startswith('**') and s.endswith('**');ital=s.startswith('*') and s.endswith('*') and not bold
  r=p.add_run(s[2:-2] if bold else s[1:-1] if ital else s);r.italic=ital;r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
def after(p,t,style=None):
 e=OxmlElement('w:p');p._p.addnext(e);n=Paragraph(e,p._parent);n.style=style or p.style;settext(n,t);n.paragraph_format.widow_control=True;return n
old=(S/'Manuscript.md').read_text(encoding='utf-8');new=old;changes={}
def replace(oldp,newp):
 global new
 assert oldp in new,oldp[:70]
 new=new.replace(oldp,newp);changes[plain(oldp)]=newp
paras=old.split('\n\n')
abm=next(p for p in paras if p.startswith('**Methods:**'))
replace(abm,abm.replace('Sensitivity analyses examined instruments, allele frequencies and model assumptions.','A fixed five-tissue GTEx extension assessed both genes.'))
abr=next(p for p in paras if p.startswith('**Results:**'))
replace(abr,abr.replace('Secondary colocalization gave *TSHR* shared-variant probabilities of 0.951 in BBJ and 0.264 in FinnGen GD, with prior sensitivity; *IGF1R* never reached 0.80.','Secondary colocalization support varied across datasets. Adipose *TSHR* associations passed correction; thyroid and fibroblast *IGF1R* estimates did not reach *P* < 0.05.'))
intro=next(p for p in paras if p.startswith('To address this question'))
replace(intro,intro.replace('Our aim was','A later fixed-panel tissue extension evaluated whether these patterns extended beyond the original blood instruments. Our aim was'))
bench=next(p for p in paras if p.startswith('Frequency substitution preserved'))
replace(bench,bench.replace('(Table S3).','(Supplementary Dataset, Detection_original).'))
method='We subsequently assessed both genes in a fixed panel of five GTEx v8 tissues: thyroid, cultured fibroblasts, subcutaneous adipose, visceral adipose and whole blood [26]. The extension plan was fixed before retrieving these tissue results, after the blood analysis was known. GTEx is predominantly European ancestry, not exclusively European. Instruments required cis-eQTL *P* < 5×10⁻⁸, F > 10 and European-reference LD r² < 0.001. Ten planned gene–tissue comparisons against FinnGen GD formed the primary family (*P* < 0.005); 40 comparisons against the other four datasets formed a secondary family (*P* < 0.00125). Unestimable combinations remained in the correction denominators. Effects used the original slopes and standard errors for inverse-normal-transformed expression. Complete regional summaries were unavailable for this extension, so tissue colocalization was not performed; significant-only pairs were used solely for instrument selection. Details and all planned combinations are in Table S3 and the Supplementary Dataset.'
result='Four of ten gene–tissue combinations provided a qualifying instrument, each retaining one SNP after clumping (F = 32.1–61.1). Higher genetically proxied *TSHR* expression in subcutaneous and visceral adipose was associated with lower FinnGen GD odds (OR 0.36, 95% CI 0.31–0.43; *P* = 1.28×10⁻³⁴ and OR 0.50, 95% CI 0.43–0.58; *P* = 4.09×10⁻¹⁹, respectively), with corrected support in BBJ and UKB. Subcutaneous estimates also passed correction in both ophthalmopathy datasets; visceral estimates did not (*P* = 0.0368 and 0.186). Neither thyroid nor cultured-fibroblast *IGF1R* estimates reached *P* < 0.05 in any dataset. Their FinnGen GD estimates were OR 0.76 (95% CI 0.56–1.02; *P* = 0.0692) and OR 1.11 (95% CI 0.89–1.38; *P* = 0.359). All 20 estimates and 30 unestimable comparisons are reported (Table S3; Supplementary Dataset, Tissue_extension).'
new=new.replace('### Orbital expression and implementation','### Tissue-specific extension\n\n'+method+'\n\n### Orbital expression and implementation')
new=new.replace('### Descriptive orbital expression','### Tissue-specific extension\n\n'+result+'\n\n### Descriptive orbital expression')
pdsc=next(p for p in paras if p.startswith('The lower-odds direction for'))
replace(pdsc,pdsc+' The adipose analyses extended the lower-odds direction beyond blood, but they reused disease GWAS data and implicated the same genomic region. They do not establish an adipose-mediated mechanism or provide independent causal replications.')
igf=next(p for p in paras if p.startswith('For *IGF1R*, the defensible'))
replace(igf,igf.replace('Blood instruments may not represent regulation in the relevant orbital cell population or disease stage.','The fixed-panel extension found no nominal association for thyroid or cultured-fibroblast *IGF1R*, despite qualifying instruments. Limited support was therefore not confined to the blood analyses that could be estimated. However, the confidence intervals, single-instrument design and lack of orbital eQTLs preclude excluding a clinically relevant effect. Cultured fibroblasts are not orbital fibroblasts, and the unestimable tissue combinations are not negative findings.'))
lim=old.split('### Limitations\n\n')[1].split('\n\n')[0]
replace(lim,lim+' The GTEx extension had four estimable combinations, each with one instrument, and predominantly European rather than ancestry-homogeneous exposure data. It reused the disease datasets and lacked tissue-specific regional colocalization.')
con=next(p for p in paras if p.startswith('Genetically proxied blood *TSHR* expression showed consistent'))
replace(con,con.replace('*IGF1R* estimates indicated higher odds but provided limited statistical evidence.','The adipose extension supported the same *TSHR* direction. Blood *IGF1R* estimates indicated higher odds but provided limited statistical evidence, and thyroid/fibroblast estimates provided no nominal support.'))
# Determine first-citation order after inserting the GTEx citation, before writing.
body,refs=new.split('## References\n',1);order=[]
for group in re.findall(r'\[([\d, ]+)\]',body):
 for n in map(int,group.split(',')):
  if n not in order:order.append(n)
assert sorted(order)==list(range(1,34));mapping={n:i+1 for i,n in enumerate(order)}
def renum(t):return re.sub(r'\[([\d, ]+)\]',lambda m:'['+', '.join(str(mapping[int(x)]) for x in m[1].split(','))+']',t)
entries={int(m[1]):m[2] for m in re.finditer(r'^(\d+)\. (.+)$',refs,re.M)}
new=renum(body)+'## References\n\n'+'\n\n'.join(str(mapping[n])+'. '+entries[n] for n in order)+'\n'
doc=Document(S/'Manuscript.docx')
for p in list(doc.paragraphs):
 if p.text in changes:settext(p,changes[p.text])
 if p.text=='Orbital expression and implementation':
  before=p._p.getprevious();prev=Paragraph(before,p._parent);h=after(prev,'Tissue-specific extension','Heading 2');after(h,method,'Normal')
 if p.text=='Descriptive orbital expression':
  prev=Paragraph(p._p.getprevious(),p._parent);h=after(prev,'Tissue-specific extension','Heading 2');after(h,result,'Normal')
for p in doc.paragraphs:
 if re.search(r'\[([\d, ]+)\]',p.text):
  # Reconstruct Markdown emphasis from existing runs before changing citations.
  t=''.join('*'+r.text+'*' if r.italic and r.text else r.text for r in p.runs);settext(p,renum(t))
refs_started=False;refps=[]
for p in doc.paragraphs:
 if p.text=='References':refs_started=True;continue
 if refs_started and re.match(r'^\d+\. ',p.text):refps.append(p)
assert len(refps)==33
for p,n in zip(refps,order):settext(p,str(mapping[n])+'. '+entries[n])
doc.save(O/'Manuscript.docx');(O/'Manuscript.md').write_text(new,encoding='utf-8')
abstract=new.split('## Abstract')[1].split('**Keywords:**')[0];wc=lambda t:len(re.sub(r'\[[\d, ]+\]','',t.replace('*','')).split())
assert wc(abstract)<=250,wc(abstract)
# Supplementary methods and replacement Table S3.
sold=(S/'Supplementary_Material.md').read_text(encoding='utf-8');snew=sold.replace('Table S3 explicitly retains this original three-outcome scope;','The Detection_original worksheet retains this original three-outcome scope;')
sm1='The post hoc tissue extension retained a frozen panel of TSHR and IGF1R in five standard-release GTEx v8 tissues. QTL sample counts were 574 (thyroid), 483 (cultured fibroblasts), 581 (subcutaneous adipose), 469 (visceral adipose) and 670 (whole blood). The original v8 archive supplied significant variant–gene pairs and eGenes records for all tested genes; gene-specific release thresholds exceeded 5×10⁻⁸ for every eligible combination, permitting complete selection at the fixed threshold. Absence from the tested-gene file was distinguished from a minimum eQTL P value above the threshold. The standard release includes multiple ancestries; the overall donor cohort was 85.3% European American, not a tissue-specific percentage.'
sm2='Variants were restricted to ±1 Mb of the source GENCODE v26 TSS, P < 5×10⁻⁸, F > 10 and MAF ≥ 0.01. GTEx variant IDs identify GRCh38 REF and ALT, with slope referring to ALT. Actual effect-allele frequencies were reconstructed from each selected eGene MAF and ref_factor, not by treating MAF as ALT frequency. GTEx mappings linked GRCh38 variants to rsIDs and GRCh37 reference coordinates. PLINK clumping used 503 European reference samples, r² < 0.001 and a 10-Mb window; an independent genotype-based calculation reproduced selection. One of 20 thyroid IGF1R candidates was absent from the reference, while its strongest candidate was retained. All selected variants were non-palindromic and matched exactly in the five disease datasets; no proxies or relaxed thresholds were used.'
sm3='Each of the four estimable gene–tissue combinations retained one SNP. Wald estimates used βY/βX and SEY/|βX| with normal two-sided P values and 95% CIs; independent R calculations verified all 20 estimates. The UKB linear-to-log-odds conversion was applied once. ORs represent one unit of inverse-normal-transformed normalized expression; they are not directly comparable to an intervention or to the reconstructed eQTLGen scale. Primary adjusted P = min(1,10P) for FinnGen GD; secondary adjusted P = min(1,40P) for the other datasets. The 30 unestimable cells remained in the planned families. Single-instrument analyses cannot support heterogeneity, leave-one-out or MR-Egger diagnostics. Significant-only eQTL pairs were not used for colocalization; full regional data were not obtained from the requester-pays resource, and no tissue posterior was calculated. Tissue_extension in the single Supplementary Dataset retains all 50 cells, availability, selected and excluded candidates, harmonized alleles and source statistics.'
snew=snew.replace('### Orbital expression and reproducibility','### Fixed-panel tissue extension\n\n'+sm1+'\n\n'+sm2+'\n\n'+sm3+'\n\n### Orbital expression and reproducibility')
sd=Document(S/'Supplementary_Material.docx')
for p in list(sd.paragraphs):
 if 'Table S3 explicitly retains' in p.text:settext(p,p.text.replace('Table S3 explicitly retains','The Detection_original worksheet retains'))
 if p.text=='Orbital expression and reproducibility':
  prev=Paragraph(p._p.getprevious(),p._parent);h=after(prev,'Fixed-panel tissue extension','Heading 2');q=after(h,sm1,'Normal');q=after(q,sm2,'Normal');after(q,sm3,'Normal')
tab=sd.tables[2]
for row in list(tab._tbl)[2:]:
 if row.tag.endswith('}tr'):tab._tbl.remove(row)
# Remove all remaining rows, retaining table properties/grid.
for row in list(tab.rows):tab._tbl.remove(row._tr)
primary=pd.read_csv(A/'Tables/Tissue_MR_primary_10.csv')
labels={'Thyroid':'Thyroid','Cells_Cultured_fibroblasts':'Cultured fibroblasts','Adipose_Subcutaneous':'Subcutaneous adipose','Adipose_Visceral_Omentum':'Visceral adipose','Whole_Blood':'Whole blood'}
head=['Gene / tissue','N','IVs','OR (95% CI)','P value','Adjusted P']
widths=[3.3,.55,.5,2.0,1.35,1.4]
data=[head]
for _,r in primary.iterrows():
 yes=r.status=='estimated';status='Not tested' if r.reason=='not_tested_in_GTEx_tissue' else 'No qualifying IV'
 data.append([r.gene+' / '+labels[r.tissue],str(int(r.n_exposure)),str(int(r.n_retained_instruments)),f'{r.OR:.2f} ({r.CI_lower:.2f}–{r.CI_upper:.2f})' if yes else status,f'{r.P:.2e}' if yes and r.P<.001 else f'{r.P:.4f}' if yes else '—',f'{r.P_adjusted:.2e}' if yes and r.P_adjusted<.001 else f'{r.P_adjusted:.3f}' if yes else '—'])
for vals in data:
 cells=tab.add_row().cells
 for j,(c,t) in enumerate(zip(cells,vals)):
  c.width=Inches(widths[j]);settext(c.paragraphs[0],t);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1
  for run in p.runs:run.font.size=Pt(9.5);run.bold=vals is head
for j,col in enumerate(tab.columns):col.width=Inches(widths[j])
tab.autofit=False
for c in tab.rows[0].cells:
 pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill','EDF0F2');pr.append(shade)
cap='Table S3. Fixed-panel tissue extension with FinnGen GD as the primary disease dataset. IVs are retained instrumental variants. ORs are per unit of normalized, inverse-normal-transformed expression. Adjusted P = min(1,10P); the ten-comparison denominator includes unestimable cells. Not tested means the gene was absent from the tissue eGenes file; no qualifying IV means the minimum eQTL P did not meet 5×10⁻⁸. All four estimable combinations used one SNP (F 32.1–61.1). The 40 secondary comparisons use a separate denominator of 40 and are retained, with all source statistics, in Tissue_extension. Cultured fibroblasts and adipose tissues are not orbital tissue. Original detection benchmarks remain in Detection_original.'
for p in sd.paragraphs:
 if p.text=='Table S3. Original association detection benchmarks':settext(p,'Table S3. Tissue-specific expression and FinnGen GD')
 elif p.text.startswith('Table S3. Previously completed MR detection benchmarks'):settext(p,cap)
sd.save(O/'Supplementary_Material.docx');(O/'Supplementary_Material.md').write_text(snew,encoding='utf-8')
sd3='| '+' | '.join(head)+' |\n| '+' | '.join(['---']*len(head))+' |\n'+'\n'.join('| '+' | '.join(row)+' |' for row in data[1:])+'\n\n'+cap+'\n'
(O/'TableS3_Tissue_extension.md').write_text(sd3,encoding='utf-8')
st=Document(S/'STROBE_MR_Checklist.docx')
for row in st.tables[0].rows:
 label=row.cells[0].text;p=row.cells[1].paragraphs[0]
 if label.startswith('8 Sensitivity'):settext(p,'Frequency, windows, priors, omissions, allele audit and fixed GTEx panel. No multi-signal coloc or Steiger test.')
 if label.startswith('13 Sensitivity'):settext(p,'Tables S1–S3; Figure S2; workbook allele and tissue sheets, including all 50 planned tissue comparisons.')
 for run in p.runs:run.font.size=Pt(9.5)
st.save(O/'STROBE_MR_Checklist.docx')
for name,a,b in [('Manuscript',old,new),('Supplementary_Material',sold,snew)]:
 (V/(name+'_changes.diff')).write_text(''.join(difflib.unified_diff(a.splitlines(True),b.splitlines(True))),encoding='utf-8')
audit={'abstract_words':wc(abstract),'main_words':wc('\n'.join(x for x in new.split('## Introduction')[1].split('## Declarations')[0].splitlines() if not x.startswith('#'))),'discussion_paragraphs':len(new.split('## Discussion')[1].split('### Limitations')[0].strip().split('\n\n')),'citation_mapping':mapping,'table_S3_replaced':True,'old_detection_data_preserved_in_workbook':True,'main_tables_figures_changed':False}
(V/'integration_summary.json').write_text(json.dumps(audit,indent=2));print(json.dumps(audit,indent=2))

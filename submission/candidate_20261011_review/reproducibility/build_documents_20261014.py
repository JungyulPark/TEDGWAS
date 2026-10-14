#!/usr/bin/env python3
"""Rebuild the reviewer-facing documents of candidate_20261011_review from Markdown.

Portable replacement for the October local build: needs only pandoc, python-docx,
LibreOffice (soffice) and poppler-utils (pdfunite, pdfinfo). Sources are the repository master and the
package's Markdown files; scientific figures are reused unchanged.

    python3 submission/candidate_20261011_review/reproducibility/build_documents_20261014.py

Outputs (package root): Manuscript, Tables, Supplementary_Material and
Cover_Letter_Draft as DOCX + PDF; Tables/Table1-3.pdf; TableS3_Tissue_extension.pdf;
STROBE_MR_Checklist.pdf (from its edited DOCX); Final_Tables_Figures.pdf and
Full_Review.pdf. All_Figures.pdf and the Figures/ directory are not touched.
"""
from __future__ import annotations

import copy
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PKG = Path(__file__).resolve().parents[1]
ROOT = PKG.parents[1]
MASTER = ROOT / "MANUSCRIPT_TED_TRAP_v5_MASTER.md"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
PANDOC = shutil.which("pandoc")
FONT = "Times New Roman"


def read(p: Path) -> str:
    return p.read_text(encoding="utf-8").replace("\r\n", "\n")


def legends() -> dict[str, str]:
    out = {}
    for line in read(PKG / "Figure_Legends.md").split("\n"):
        m = re.match(r"^\*\*(Figure S?\d)\.", line)
        if m:
            out[m.group(1)] = line
    return out


# ---------------------------------------------------------------- formatting
def page(sec, landscape=False):
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (Inches(11.69), Inches(8.27)) if landscape else (Inches(8.27), Inches(11.69))
    sec.left_margin = sec.right_margin = Inches(0.7 if landscape else 1.0)
    sec.top_margin = sec.bottom_margin = Inches(0.7 if landscape else 1.0)


def line_numbers(sec):
    for x in sec._sectPr.findall(qn("w:lnNumType")):
        sec._sectPr.remove(x)
    x = OxmlElement("w:lnNumType")
    x.set(qn("w:countBy"), "1")
    x.set(qn("w:restart"), "continuous")
    x.set(qn("w:distance"), "240")
    sec._sectPr.append(x)


def page_number_footer(sec):
    p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.get_or_add_pPr().append(OxmlElement("w:suppressLineNumbers"))
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.name = FONT
    run.font.size = Pt(10)


def set_font(style):
    """Force Times New Roman, overriding theme fonts that pandoc's styles inherit."""
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)


def styles(doc, size, spacing):
    by_name = {s.name: s for s in doc.styles}
    for name in ["Normal", "Body Text", "First Paragraph", "Compact", "Table", "Block Text"]:
        if name not in by_name:
            continue
        st = by_name[name]
        set_font(st)
        st.font.size = Pt(size)
        st.paragraph_format.line_spacing = spacing
        st.paragraph_format.space_after = Pt(4 if spacing < 1.5 else 0)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        if name not in by_name:
            continue
        st = by_name[name]
        set_font(st)
        st.font.size = Pt({"Title": 15, "Heading 1": 13, "Heading 2": 12, "Heading 3": 11}[name])
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)


def border(tag, sz="8"):
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    return el


def format_tables(doc, size):
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        for x in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(x)
        b = OxmlElement("w:tblBorders")
        b.append(border("top", "12"))
        b.append(border("bottom", "12"))
        tblPr.append(b)
        for i, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    tb = OxmlElement("w:tcBorders")
                    tb.append(border("bottom", "8"))
                    tcPr.append(tb)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(size)
                        r.font.name = FONT
                        if i == 0:
                            r.font.bold = True


def small_notes(doc, size):
    """Table titles/notes outside tables: single-spaced, slightly smaller."""
    for p in doc.paragraphs:
        if p.style.name.startswith(("Heading", "Title")):
            continue
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.size = Pt(size)


# ---------------------------------------------------------------- building
def pandoc(md_text: str, out: Path, resource: Path | None = None):
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in.md"
        src.write_text(md_text, encoding="utf-8")
        cmd = [PANDOC, str(src), "-f", "markdown+pipe_tables+subscript+superscript-implicit_figures",
               "-t", "docx", "-o", str(out)]
        if resource:
            cmd += ["--resource-path", str(resource)]
        subprocess.run(cmd, check=True)


def to_pdf(docx: Path) -> Path:
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", td, str(docx)],
                       check=True, capture_output=True)
        pdf = Path(td) / (docx.stem + ".pdf")
        target = docx.with_suffix(".pdf")
        shutil.copyfile(pdf, target)
    return target


def strip_meta(doc, title):
    cp = doc.core_properties
    cp.title = title
    cp.author = "Jungyul Park; Min-Seon Kim; Kyung-Hwa Shin; Suk-Woo Yang"
    cp.last_modified_by = ""
    cp.comments = ""


def build_manuscript():
    text = read(MASTER)
    leg = legends()
    text = text.rstrip() + "\n\n## Figure legends\n\n" + "\n\n".join(leg[f"Figure {i}"] for i in (1, 2, 3)) + "\n"
    out = PKG / "Manuscript.docx"
    pandoc(text, out)
    d = Document(out)
    styles(d, 12, 2.0)
    if "Line Number" not in {s.name for s in d.styles}:
        from docx.enum.style import WD_STYLE_TYPE
        ln = d.styles.add_style("Line Number", WD_STYLE_TYPE.CHARACTER)
        ln.element.set(qn("w:styleId"), "LineNumber")
        set_font(ln)
        ln.font.size = Pt(9)
    for sec in d.sections:
        page(sec)
        line_numbers(sec)
        page_number_footer(sec)
    for p in d.paragraphs:
        if p.style.name in ("Heading 1", "Heading 2") and p.text in ("Abstract", "Introduction", "References", "Figure legends"):
            p.paragraph_format.page_break_before = True
        if p.style.name == "Heading 1" and p.text.startswith("TSHR and IGF1R"):
            p.style = d.styles["Title"]
    strip_meta(d, re.search(r"^# (.+)$", text, re.M).group(1))
    d.save(out)
    return out


def build_tables(md_text: str, out: Path, title: str):
    pandoc(md_text, out)
    d = Document(out)
    styles(d, 10, 1.1)
    for sec in d.sections:
        page(sec, landscape=True)
        page_number_footer(sec)
    first = True
    for p in d.paragraphs:
        if p.style.name in ("Heading 1", "Heading 2") and p.text.startswith("Table"):
            if not first:
                p.paragraph_format.page_break_before = True
            first = False
        if p.style.name.startswith("Heading") and p.text.startswith("C. Frequency-sensitive"):
            p.paragraph_format.page_break_before = True
    format_tables(d, 9)
    small_notes(d, 9)
    strip_meta(d, title)
    d.save(out)
    return out


def tables_sections():
    t = read(PKG / "Tables.md")
    body = t.split("\n", 1)[1]  # drop "# Main tables"
    parts = re.split(r"(?m)^(?=## Table )", body)
    return [p for p in parts if p.startswith("## Table ")]


def build_supplement():
    leg = legends()
    s3 = read(PKG / "TableS3_Tissue_extension.md").strip().split("\n")
    note = [l for l in s3 if l.startswith("Table S3.")][0].replace("Table S3. ", "", 1)
    rows = [l for l in s3 if l.startswith("|")]
    s3md = "## Table S3. Tissue-specific expression and FinnGen GD\n\n" + "\n".join(rows) + "\n\n" + note + "\n"
    text = (read(PKG / "Supplementary_Material.md").rstrip() + "\n\n"
            + read(PKG / "SupplementaryTables_S1_S2.md").strip() + "\n\n" + s3md + "\n"
            + "## Figure S1\n\n![](Figures/FigureS1.png){width=6.3in}\n\n" + leg["Figure S1"] + "\n\n"
            + "## Figure S2\n\n![](Figures/FigureS2.png){width=6.3in}\n\n" + leg["Figure S2"] + "\n")
    text = text.replace("# Supplementary material", "# Supplementary material\n\nTSHR and IGF1R in Graves disease and thyroid eye disease: genetic evidence from blood expression", 1)
    out = PKG / "Supplementary_Material.docx"
    pandoc(text, out, resource=PKG)
    d = Document(out)
    styles(d, 11, 1.5)
    for sec in d.sections:
        page(sec)
        page_number_footer(sec)
    for p in d.paragraphs:
        if p.style.name in ("Heading 1", "Heading 2") and re.match(r"^(Table|Figure) S\d", p.text):
            p.paragraph_format.page_break_before = True
    format_tables(d, 9)
    strip_meta(d, "Supplementary material")
    d.save(out)
    return out, s3md


def build_cover():
    out = PKG / "Cover_Letter_Draft.docx"
    pandoc(read(PKG / "Cover_Letter_Draft.md"), out)
    d = Document(out)
    styles(d, 10.5, 1.0)
    for sec in d.sections:
        page(sec)
        sec.left_margin = sec.right_margin = Inches(0.8)
        sec.top_margin = sec.bottom_margin = Inches(0.55)
    for p in d.paragraphs:
        p.paragraph_format.space_after = Pt(4)
        if p.style.name == "Title":
            p.paragraph_format.space_before = Pt(0)
    strip_meta(d, "Cover letter")
    d.save(out)
    return out


def merge(pdfs, out: Path):
    subprocess.run(["pdfunite", *map(str, pdfs), str(out)], check=True)


def pages(pdf: Path) -> int:
    info = subprocess.run(["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout
    return int(re.search(r"^Pages:\s+(\d+)", info, re.M).group(1))


def main() -> int:
    if not (PANDOC and SOFFICE):
        sys.exit("pandoc and LibreOffice (soffice) are required")
    ms = to_pdf(build_manuscript())
    secs = tables_sections()
    tables_docx = build_tables("\n".join(secs), PKG / "Tables.docx", "Main tables")
    tables_pdf = to_pdf(tables_docx)
    (PKG / "Tables").mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        for i, sec in enumerate(secs, 1):
            one = build_tables(sec, Path(td) / f"Table{i}.docx", f"Table {i}")
            shutil.copyfile(to_pdf(one), PKG / "Tables" / f"Table{i}.pdf")
        supp_docx, s3md = build_supplement()
        supp_pdf = to_pdf(supp_docx)
        s3 = build_tables(s3md, Path(td) / "TableS3_Tissue_extension.docx", "Table S3")
        shutil.copyfile(to_pdf(s3), PKG / "TableS3_Tissue_extension.pdf")
        s12 = build_tables(read(PKG / "SupplementaryTables_S1_S2.md"), Path(td) / "S12.docx", "Tables S1-S2")
        s12_pdf = to_pdf(s12)
        figs = [PKG / "Figures" / f"Figure{i}.pdf" for i in (1, 2, 3)]
        sfigs = [PKG / "Figures" / f"FigureS{i}.pdf" for i in (1, 2)]
        merge([tables_pdf, *figs, s12_pdf, PKG / "TableS3_Tissue_extension.pdf", *sfigs], PKG / "Final_Tables_Figures.pdf")
        merge([ms, tables_pdf, *figs, supp_pdf], PKG / "Full_Review.pdf")
    to_pdf(build_cover())
    to_pdf(PKG / "STROBE_MR_Checklist.docx")
    for name in ["Manuscript", "Tables", "Supplementary_Material", "Cover_Letter_Draft", "STROBE_MR_Checklist",
                 "Final_Tables_Figures", "Full_Review", "TableS3_Tissue_extension"]:
        print(f"{name}.pdf: {pages(PKG / (name + '.pdf'))} pages")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

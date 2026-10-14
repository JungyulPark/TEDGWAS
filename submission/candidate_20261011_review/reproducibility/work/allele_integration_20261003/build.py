from pathlib import Path
import sys,re,shutil,json,hashlib,difflib,zipfile,copy
from xml.etree import ElementTree as E
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import RGBColor,Pt
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1]
S=R/'outputs/TED_Research_Contribution_20261003'
A=R/'outputs/TED_Allele_Direction_20261003'
O=R/'outputs/TED_Allele_Integrated_20261003';V=O/'Verification'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert not O.exists(),'Preserve existing packages'
for name,digest in json.loads((S/'PACKAGE_SHA256.json').read_text(encoding='utf-8')).items():assert sha(S/name)==digest,name
assert json.loads((A/'AUDIT_RESULT.json').read_text())['all_pass']
O.mkdir();V.mkdir()
for p in S.iterdir():
    if p.name in ['Verification','PACKAGE_SHA256.json','Full_Review.pdf','READ_FIRST_KO.md','Followup_Analysis_Assessment_KO.md']:continue
    if p.is_dir():shutil.copytree(p,O/p.name)
    else:shutil.copy2(p,O/p.name)
shutil.copy2(S/'PACKAGE_SHA256.json',V/'Previous_package_SHA256.json')
shutil.copytree(A,V/'Allele_audit')
P=O/'Research_planning';P.mkdir()
for name in ['PROTOCOL_v1.md','PROTOCOL_FREEZE.json','Tissue_planned_results_50_cells.csv']:shutil.copy2(A/name,P/name)

def inline(p,t):
    for s in re.split(r'(\*[^*]+\*)',t):
        if not s:continue
        r=p.add_run(s[1:-1] if s.startswith('*') and s.endswith('*') else s)
        r.italic=s.startswith('*') and s.endswith('*');r.font.color.rgb=RGBColor(0,0,0)
def settext(p,t):
    p.clear();inline(p,t)
def after(p,t,style=None):
    e=OxmlElement('w:p');p._p.addnext(e);n=Paragraph(e,p._parent)
    n.style=style or p.style;inline(n,t);n.paragraph_format.widow_control=True
    return n
def plain(t):return t.replace('*','')

method='In a post hoc allele-direction audit, we re-extracted rs179252 and rs179247 from the original expression and disease files. Disease effects were aligned to the blood expression-increasing allele, and G-allele correlations were calculated in European and East Asian reference genotypes and independently checked with PLINK. This audit assessed allele alignment and consistency with published thymus findings, rather than thymic mediation (Supplementary Dataset, TSHR_allele_audit).'
res='The expression-increasing rs179252-G allele was associated with lower odds in all five disease datasets and was positively correlated with rs179247-G in European (r = 0.978; r² = 0.957; n = 503) and East Asian (r = 0.996; r² = 0.992; n = 504) reference samples. The latter allele also increased blood *TSHR* expression. Allele-aligned recalculation reproduced the original MR estimates (Supplementary Dataset, TSHR_allele_audit).'
disc='The lower-odds direction for *TSHR* need not contradict the pathogenic role of receptor-stimulating antibodies. Blood transcript abundance, tissue-specific expression and receptor activation are different exposures, and genetic regulatory effects can differ across tissues [26]. Our audit linked the blood expression-increasing rs179252-G allele to rs179247-G in both reference populations. The latter is the GD-protective allele associated with greater thymic transcription in published studies [27, 28] and also increased blood expression in eQTLGen. This directional consistency strengthens the plausibility of a central-tolerance explanation, but does not demonstrate that thymic expression mediates the disease association. These linked variants are not independent replications or proven causal regulators. The marked departure of the MR OR from 1 also reflects division by a small reconstructed expression coefficient. It should not be translated into the effect of experimentally increasing *TSHR* expression.'
old=(S/'Manuscript.md').read_text(encoding='utf-8');text=old
paragraphs=old.split('\n\n')
anchor_m=next(p for p in paragraphs if p.startswith('Where applicable, sensitivity assessment'))
anchor_r=next(p for p in paragraphs if p.startswith('Higher genetically proxied blood *TSHR* expression was associated with lower odds of BBJ'))
old_disc=next(p for p in paragraphs if p.startswith('The lower-odds direction for *TSHR*'))
text=text.replace(anchor_m,anchor_m+'\n\n'+method).replace(anchor_r,anchor_r+'\n\n'+res).replace(old_disc,disc)
text=text.replace('individual-level multivariable regression, genetic directionality test or multi-signal colocalization','individual-level multivariable regression, Steiger directionality test or multi-signal colocalization')
d=Document(S/'Manuscript.docx')
for p in list(d.paragraphs):
    if p.text==plain(anchor_m):after(p,method)
    elif p.text==plain(anchor_r):after(p,res)
    elif p.text==plain(old_disc):settext(p,disc)
    elif 'genetic directionality test or multi-signal colocalization' in p.text:settext(p,p.text.replace('genetic directionality test','Steiger directionality test'))
d.save(O/'Manuscript.docx');(O/'Manuscript.md').write_text(text,encoding='utf-8')
assert plain(method) in [p.text for p in d.paragraphs] and plain(res) in [p.text for p in d.paragraphs]

supp_old=(S/'Supplementary_Material.md').read_text(encoding='utf-8')
sp=next(p for p in supp_old.split('\n\n') if p.startswith('Previous TSHR clumping'))
sp_new=sp.replace('Reference r² for rs179252–rs179247 was 0.957116/0.991625 in Europeans/East Asians. These LD results do not establish conditional association or matching regulatory direction across tissues. ','').replace('The present reappraisal preserves these earlier results without rerunning them.','These earlier clumping and conditional analyses were preserved; the separate allele-direction audit below newly examined rs179252 and rs179247.')
sh='TSHR allele alignment and reference LD'
ss1='In a post hoc audit, original full eQTLGen and disease GWAS files were re-extracted for rs179252 and rs179247. Their GRCh37 positions were 14:81435985 and 14:81432546, and GRCh38 positions were 14:80969641 and 14:80966202. Both non-palindromic allele pairs matched across sources. UKB effects were reversed to the G allele after the existing scale conversion; BBJ and FinnGen required no reversal. Exposure coefficients and first-order Wald estimates were independently reconstructed. Original GWAS P values and recalculated MR P values were retained separately.'
ss2='G-allele counts at both variants were positively correlated in European (n = 503; r = 0.978322812; r² = 0.957115524) and East Asian (n = 504; r = 0.995803919; r² = 0.991625445) references, without missing genotypes. Direct BED decoding agreed with PLINK dosage output. PLINK haplotype estimates supported G–G and T–A coupling; these were inferred from unphased genotypes, not directly observed phased haplotypes. Genotype-based r² and haplotype-based estimates were kept distinct. Alignment to the G allele resolved the negative default European correlation obtained when counting T at rs179252. The TSHR_allele_audit worksheet retains both eQTL records, both variants across five disease outcomes, all four inferred haplotypes in each population, and MR reproduction. This audit does not identify a causal variant or establish mediation in thymus or orbit.'
supp=supp_old.replace(sp,sp_new+'\n\n### '+sh+'\n\n'+ss1+'\n\n'+ss2)
d=Document(S/'Supplementary_Material.docx')
p=next(p for p in d.paragraphs if p.text==sp)
settext(p,sp_new);h=after(p,sh,'Heading 2');q=after(h,ss1,'Normal');after(q,ss2,'Normal')
d.save(O/'Supplementary_Material.docx');(O/'Supplementary_Material.md').write_text(supp,encoding='utf-8')
assert len(d.tables)==3 and len(d.inline_shapes)==2

d=Document(S/'STROBE_MR_Checklist.docx')
for row in d.tables[0].rows:
    if row.cells[0].text.startswith('13 Sensitivity results'):
        settext(row.cells[1].paragraphs[0],'Tables S1–S3; Figure S2; frequency/LOO/window and TSHR allele-audit sheets. Single-IV exclusions stated.')
        for run in row.cells[1].paragraphs[0].runs:run.font.size=Pt(9.5)
d.save(O/'STROBE_MR_Checklist.docx')

for name,prev,new in [('Manuscript',old,text),('Supplementary_Material',supp_old,supp)]:
    (V/(name+'_changes.diff')).write_text(''.join(difflib.unified_diff(prev.splitlines(True),new.splitlines(True))),encoding='utf-8')

# Prepare a two-sheet editing subset so large immutable statistical worksheets
# never pass through an unrelated workbook conversion. Artifact Tool authors
# the changed readme/dictionary and the new audit sheet; OPC merge preserves
# the remaining worksheet bytes and their existing style references.
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',NS)
with zipfile.ZipFile(S/'Supplementary_Dataset.xlsx') as z:
    parts={n:z.read(n) for n in z.namelist()}
wb=E.fromstring(parts['xl/workbook.xml']);sheets=wb.find('{'+NS+'}sheets')
for elem in list(sheets)[2:]:sheets.remove(elem)
parts['xl/workbook.xml']=E.tostring(wb)
rel=E.fromstring(parts['xl/_rels/workbook.xml.rels'])
for elem in list(rel):
    if elem.attrib['Type'].endswith('/worksheet') and elem.attrib['Id'] not in ['rId1','rId2']:rel.remove(elem)
E.register_namespace('',REL)
parts['xl/_rels/workbook.xml.rels']=E.tostring(rel)
ct=E.fromstring(parts['[Content_Types].xml'])
for elem in list(ct):
    if elem.attrib.get('PartName','').startswith('/xl/worksheets/sheet') and elem.attrib['PartName'] not in ['/xl/worksheets/sheet1.xml','/xl/worksheets/sheet2.xml']:ct.remove(elem)
E.register_namespace('','http://schemas.openxmlformats.org/package/2006/content-types')
parts['[Content_Types].xml']=E.tostring(ct)
with zipfile.ZipFile(W/'edit_subset.xlsx','w',zipfile.ZIP_DEFLATED) as z:
    for n,b in parts.items():
        if n.startswith('xl/worksheets/') and n not in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:continue
        z.writestr(n,b)

blocks=[];start=5
for title,file in [('A Blood expression directions','Table_A1_eQTL_original_direction.csv'),('B Allele labelled genotype LD','Table_A2_signed_LD.csv'),('C Original and aligned disease effects','Table_A3_both_variants_five_disease_GWAS.csv'),('D Reproduced MR estimates','Table_A4_MR_reproduction.csv'),('E Inferred haplotype frequencies','Table_A2b_inferred_haplotype_frequencies.csv')]:
    df=pd.read_csv(A/'Tables'/file)
    # Raw numeric values remain numbers. These are audited analytical exports,
    # not Excel formulas replacing a statistical program.
    values=df.astype(object).where(pd.notna(df),None).values.tolist()
    blocks.append(dict(title=title,file=file,row=start,headers=list(df.columns),values=values))
    start+=len(df)+4
(W/'blocks.json').write_text(json.dumps(blocks,ensure_ascii=False,allow_nan=False),encoding='utf-8')
wc=lambda t:len(re.sub(r'\[[0-9, ]+\]','',t.replace('*','')).split())
audit={'source_package_verified':True,'abstract_words':wc(text.split('## Abstract\n')[1].split('**Keywords:**')[0]),'main_words':wc('\n'.join(l for l in text.split('## Introduction\n')[1].split('## Declarations')[0].splitlines() if not l.startswith('#'))),'discussion_paragraphs':len(text.split('## Discussion\n')[1].split('### Limitations')[0].strip().split('\n\n')),'new_primary_tables':0,'new_primary_figures':0,'new_supplementary_tables':0,'new_workbook_sheets':['TSHR_allele_audit'],'new_tissue_MR_performed':False,'references_renumbered':False}
(V/'integration_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
print(json.dumps(audit,indent=2))

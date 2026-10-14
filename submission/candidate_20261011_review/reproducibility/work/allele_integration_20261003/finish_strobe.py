from pathlib import Path
from docx import Document
from docx.shared import Pt
import os,sys,runpy
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1]
p=R/'outputs/TED_Allele_Integrated_20261003/STROBE_MR_Checklist.docx'
d=Document(p)
for row in d.tables[0].rows:
    for cell in row.cells:
        for para in cell.paragraphs:
            if para.text.startswith('Frequency, windows, priors, instrument omissions.'):
                para.clear()
                para.add_run('Frequency, windows, priors, instrument omissions and allele alignment. No multi-signal coloc or Steiger directionality test.').font.size=Pt(9.5)
            para.paragraph_format.space_before=Pt(2)
            para.paragraph_format.space_after=Pt(2)
d.save(p)
os.environ['PATH']=str(R/'work/tools/lo/LibreOffice/program')+os.pathsep+'${LOCAL_USER_HOME}/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin'+os.pathsep+os.environ['PATH']
renderer='${LOCAL_USER_HOME}/.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents/render_docx.py'
out=W/'render_v3/STROBE_MR_Checklist'
assert not list(out.glob('page-*.png'))
sys.argv=[renderer,str(p),'--output_dir',str(out),'--emit_pdf']
runpy.run_path(renderer,run_name='__main__')

from pathlib import Path
import json,re,hashlib,zipfile,shutil,csv,sys,math
from docx import Document
from pypdf import PdfReader,PdfWriter
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1]
S=R/'outputs/TED_Research_Contribution_20261003';O=R/'outputs/TED_Allele_Integrated_20261003';V=O/'Verification'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
checks=[]
def check(name,condition):
    checks.append({'check':name,'pass':bool(condition)})
    assert condition,name
old=(S/'Manuscript.md').read_text(encoding='utf-8');new=(O/'Manuscript.md').read_text(encoding='utf-8')
def between(text,a,b):return text.split(a,1)[1].split(b,1)[0]
for a,b in [('## Abstract','## Introduction'),('## Introduction','## Methods'),('### Limitations','## Declarations'),('## Declarations','## References')]:check('unchanged '+a,between(old,a,b)==between(new,a,b))
check('references identical',old.split('## References')[1]==new.split('## References')[1])
check('33 reference entries',len(re.findall(r'^\d+\. ',new.split('## References')[1],re.M))==33)
first=[]
for g in re.findall(r'\[([0-9, –-]+)\]',new.split('## References')[0]):
    for item in g.split(','):
        z=item.strip();m=re.fullmatch(r'(\d+)[–-](\d+)',z)
        values=list(range(int(m[1]),int(m[2])+1)) if m else [int(z)]
        for value in values:
            if value not in first:first.append(value)
check('citation first occurrence order 1-33',first==list(range(1,34)))
oldres=between(old,'## Results','## Discussion').strip().split('\n\n');newres=between(new,'## Results','## Discussion').strip().split('\n\n')
for i,p in enumerate(oldres):check('existing Results paragraph '+str(i)+' preserved',p in newres)
check('one new Results paragraph',len(newres)==len(oldres)+1)
old_disc=between(old,'## Discussion','### Limitations').strip().split('\n\n');new_disc=between(new,'## Discussion','### Limitations').strip().split('\n\n')
check('8 Discussion paragraphs',len(new_disc)==8)
check('only protective direction Discussion changed',sum(a!=b for a,b in zip(old_disc,new_disc))==1)
check('2 Limitations paragraphs',len(between(new,'### Limitations','### Conclusions').strip().split('\n\n'))==2)
for p in new.split('\n\n'):
    if p.startswith(('In a post hoc allele-direction audit','The expression-increasing rs179252-G','The lower-odds direction for')):
        check('DOCX/MD new paragraph '+p[:32],p.replace('*','') in [x.text for x in Document(O/'Manuscript.docx').paragraphs])
for folder in ['Tables','Figures']:
    for p in (S/folder).iterdir():
        if p.is_file():check('unchanged '+str(p.relative_to(S)),sha(p)==sha(O/p.relative_to(S)))
for file in ['Tables.docx','Tables.md','Tables.pdf','All_Figures.pdf','Figure_Legends.md','Cover_Letter_Draft.docx','Cover_Letter_Draft.md','Cover_Letter_Draft.pdf']:
    check('unchanged '+file,sha(S/file)==sha(O/file))
sd,nd=Document(S/'Supplementary_Material.docx'),Document(O/'Supplementary_Material.docx')
check('3 supplementary tables',len(nd.tables)==3)
for i,(a,b) in enumerate(zip(sd.tables,nd.tables)):check('unchanged supplementary table XML '+str(i+1),a._tbl.xml==b._tbl.xml)
check('2 supplementary figures',len(nd.inline_shapes)==2)
with zipfile.ZipFile(S/'Supplementary_Material.docx') as a,zipfile.ZipFile(O/'Supplementary_Material.docx') as b:
    for p in a.namelist():
        if p.startswith('word/media/'):check('unchanged supplementary media '+p,a.read(p)==b.read(p))
old_supp=(S/'Supplementary_Material.md').read_text(encoding='utf-8')
new_supp=(O/'Supplementary_Material.md').read_text(encoding='utf-8')
for i,paragraph in enumerate(old_supp.strip().split('\n\n')):
    if not paragraph.startswith('Previous TSHR clumping'):
        check('existing supplementary prose preserved '+str(i),paragraph in new_supp)
ld=list(csv.DictReader((V/'Allele_audit/Tables/Table_A2_signed_LD.csv').open()))
for row in ld:
    for column in ['signed_G_dosage_r','r_squared']:
        check('Results LD rounding '+row['population']+column,format(float(row[column]),'.3f') in next(p for p in newres if p.startswith('The expression-increasing')))
    check('aligned G-G positive '+row['population'],float(row['signed_G_dosage_r'])>0 and row['rs179252_counted_allele']==row['rs179247_counted_allele']=='G')
wbcheck=json.loads((V/'workbook_checks.json').read_text())
check('workbook independent checks passed',wbcheck['status']=='PASS' and wbcheck['new_numeric_cells_verified']==248)
check('workbook hash matches checked file',sha(O/'Supplementary_Dataset.xlsx')==wbcheck['sha256'])
w=load_workbook(O/'Supplementary_Dataset.xlsx',read_only=True,data_only=False)
check('workbook opens with 20 sheets',len(w.sheetnames)==20 and w.sheetnames[-1]=='TSHR_allele_audit')
check('audit G effect alleles correct',w['TSHR_allele_audit']['E7'].value=='G' and w['TSHR_allele_audit']['E8'].value=='G')
w.close()
for p in (O/'Research_planning').iterdir():check('frozen planning file unchanged '+p.name,sha(p)==sha(V/'Allele_audit'/p.name))
check('tissue MR still not executed',json.loads((V/'integration_audit.json').read_text())['new_tissue_MR_performed'] is False)
pdf_pages={}
for name,folder,count in [('Manuscript',W/'render/Manuscript',14),('Supplementary_Material',W/'render/Supplementary_Material',8),('STROBE_MR_Checklist',W/'render_v3/STROBE_MR_Checklist',1)]:
    p=folder/(name+'.pdf');pdf=PdfReader(p);pdf_pages[name]=len(pdf.pages)
    check(name+' rendered page count',len(pdf.pages)==count)
    text='\n'.join(page.extract_text() for page in pdf.pages)
    check(name+' PDF text readable',len(text)>1000 and '\ufffd' not in text)
    check(name+' updated PDF',('TSHR_allele_audit' in text) if name!='STROBE_MR_Checklist' else 'TSHR allele-audit' in text)
    shutil.copy2(p,O/p.name)
components=['Manuscript.pdf','Tables.pdf']+['Figures/Figure'+str(i)+'.pdf' for i in [1,2,3]]+['Supplementary_Material.pdf']
merged=PdfWriter();expected=[]
for component in components:
    reader=PdfReader(O/component);expected.extend(page.extract_text() for page in reader.pages)
    merged.append(reader,import_outline=False)
with (O/'Full_Review.pdf').open('wb') as f:merged.write(f)
merged.close();read=PdfReader(O/'Full_Review.pdf')
check('review PDF page count',len(read.pages)==len(expected))
for i,p in enumerate(read.pages):check('review PDF page '+str(i+1)+' preserved',p.extract_text()==expected[i])
pdf_pages['Full_Review']=len(read.pages)
result={'status':'PASS','scope':'Integration of completed allele audit; no new tissue MR; not certification of causal validity','checks_count':len(checks),'checks':checks,'pdf_page_counts':pdf_pages,'visual_review':'Every page of all three edited DOCX files inspected after rendering; changed spreadsheet sections inspected; STROBE spacing repaired to one page','workbook_validation':wbcheck}
(V/'final_checks.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
code=V/'Integration_scripts';code.mkdir(exist_ok=True)
for file in ['build.py','workbook.mjs','merge_workbook.py','render.py','finish_strobe.py','verify_and_package.py']:shutil.copy2(W/file,code/file)
manifest={str(p.relative_to(O)).replace('\\','/'):sha(p) for p in sorted(O.rglob('*')) if p.is_file() and p.name!='PACKAGE_SHA256.json'}
(O/'PACKAGE_SHA256.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8')
dest=O.with_suffix('.zip')
with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in sorted(O.rglob('*')):
        if p.is_file():z.write(p,O.name+'/'+str(p.relative_to(O)).replace('\\','/'))
with zipfile.ZipFile(dest) as z:
    assert z.testzip() is None
    for p,h in manifest.items():assert hashlib.sha256(z.read(O.name+'/'+p)).hexdigest()==h,p
print(json.dumps({'status':'PASS','checks':len(checks),'pdf_pages':pdf_pages,'package_files':len(manifest)+1,'zip_bytes':dest.stat().st_size,'zip_sha256':sha(dest)},indent=2))

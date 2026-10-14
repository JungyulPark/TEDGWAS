from pathlib import Path
import sys, re, json, hashlib, shutil, difflib
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
from content_v1 import ABSTRACT, RESULTS, DISCUSSION, CONCLUSION, NEW_REFERENCES

sys.stdout.reconfigure(encoding='utf-8')
W = Path(__file__).resolve().parent
R = W.parents[1]
S = R/'outputs/TED_Manuscript_Reframed_20261002'
O = R/'outputs/TED_Association_Manuscript_20261003'
V = O/'Verification'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
assert not O.exists(), 'Versioned output already exists; use a follow-up script for corrections.'
O.mkdir(); V.mkdir()
old_manifest = json.loads((S/'PACKAGE_SHA256.json').read_text(encoding='utf-8'))
for name, digest in old_manifest.items():
    assert sha(S/name) == digest, name
for p in S.iterdir():
    if p.name in ['Verification','PACKAGE_SHA256.json','Full_Review.pdf','READ_FIRST_KO.md']:
        continue
    if p.is_dir(): shutil.copytree(p,O/p.name)
    else: shutil.copy2(p,O/p.name)
shutil.copytree(S/'Verification', V/'Original_package_verification')
shutil.copy2(S/'PACKAGE_SHA256.json',V/'Original_package_manifest.json')

old = (S/'Manuscript.md').read_text(encoding='utf-8')
t = old
start = t.index('## Abstract\n') + len('## Abstract\n')
end = t.index('\n**Keywords:**',start)
t = t[:start]+'\n'+ABSTRACT+'\n'+t[end:]
old_intro = 'Our objective was to characterize the strength and consistency of genetic evidence, rather than to test a drug or classify GD and TED as genetically independent diseases.'
new_intro = 'Our objective was to compare the strength, direction and consistency of MR associations for *TSHR* and *IGF1R*, using the wider gene screen as context and regional colocalization to assess interpretation. The study did not test treatment efficacy or directly compare GD patients with and without TED.'
assert old_intro in t
t=t.replace(old_intro,new_intro)
t=t.replace('### Colocalization and robustness assessment','### Secondary regional analysis and robustness assessment')
needle='Colocalization used complete available regional cis-eQTL statistics'
assert needle in t
t=t.replace(needle, 'We report MR associations as the main findings and use colocalization to assess their interpretation. This reporting emphasis was adopted during post hoc reappraisal and was not a prospectively registered hierarchy. Colocalization used complete available regional cis-eQTL statistics')
t=t.replace('A small P value or high PP.H4','A small *P* value or high PP.H4')
start=t.index('## Results\n'); end=t.index('## Discussion\n',start)
t=t[:start]+'## Results\n\n'+RESULTS+'\n\n'+t[end:]
start=t.index('## Discussion\n')+len('## Discussion\n');end=t.index('### Limitations',start)
t=t[:start]+'\n'+DISCUSSION+'\n\n'+t[end:]
start=t.index('### Conclusions\n')+len('### Conclusions\n');end=t.index('## Declarations',start)
t=t[:start]+'\n'+CONCLUSION+'\n\n'+t[end:]
body, refs = t.split('## References\n',1)
refmap={int(n):s.strip() for n,s in re.findall(r'(?ms)^(\d+)\. (.*?)(?=^\d+\. |\Z)',refs)}
refmap.update(NEW_REFERENCES)
order=[]
for group in re.findall(r'\[([0-9, ]+)\]',body):
    for n in map(int,re.findall(r'\d+',group)):
        assert n in refmap
        if n not in order:order.append(n)
assert set(order)==set(refmap)
mapping={n:i+1 for i,n in enumerate(order)}
body=re.sub(r'\[([0-9, ]+)\]',lambda m:'['+', '.join(str(mapping[int(n)]) for n in re.findall(r'\d+',m[1]))+']',body)
t=body+'## References\n\n'+'\n\n'.join(f'{mapping[n]}. {refmap[n]}' for n in order)+'\n'
(O/'Manuscript.md').write_text(t,encoding='utf-8')

def inline(p,text):
    text=text.replace('\\*','＊')
    for x in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text):
        if not x:continue
        if x.startswith('**') and x.endswith('**'):p.add_run(x[2:-2]).bold=True
        elif x.startswith('*') and x.endswith('*'):p.add_run(x[1:-1]).italic=True
        else:p.add_run(x.replace('＊','*'))

def rebuild(name,text):
    d=Document(S/(name+'.docx'))
    for e in list(d._element.body):
        if e.tag!=qn('w:sectPr'):d._element.body.remove(e)
    for block in re.split(r'\n\s*\n',text.strip()):
        block=block.strip()
        if not block:continue
        if block.startswith('# '):
            p=d.add_paragraph(style='Title');inline(p,block[2:])
        elif block.startswith('## '):
            p=d.add_paragraph(style='Heading 1');inline(p,block[3:])
            if block[3:] in ['Abstract','Introduction','References']:p.paragraph_format.page_break_before=True
        elif block.startswith('### '):
            p=d.add_paragraph(style='Heading 2');inline(p,block[4:])
        else:
            p=d.add_paragraph();inline(p,block.replace('  \n','\n'))
        p.paragraph_format.widow_control=True
        for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)
    for e in list(d._element.xpath('.//w:pBdr')):e.getparent().remove(e)
    for style in d.styles:
        if style.name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
            style.font.color.rgb=RGBColor(0,0,0)
            for e in list(style.element.xpath('.//w:pBdr')):e.getparent().remove(e)
    d.save(O/(name+'.docx'))
    extracted='\n'.join(p.text for p in d.paragraphs)
    assert '##' not in extracted and '**' not in extracted

rebuild('Manuscript',t)
cover=(S/'Cover_Letter_Draft.md').read_text(encoding='utf-8')
cover=cover.replace('TSHR showed consistent MR associations, but shared-variant evidence varied across disease definitions and prior assumptions; IGF1R evidence remained limited.', 'Higher genetically proxied blood TSHR expression showed consistent associations with lower disease odds, while IGF1R estimates indicated higher odds with limited statistical support. CTLA4, HLA-DQA2 and TUBB also met the full-screen threshold in both GD datasets. Secondary colocalization qualified expression-mediated interpretation rather than replacing the association findings.')
cover=cover.replace('The manuscript distinguishes susceptibility evidence from treatment efficacy and does not claim TED-specific effects within GD.', 'The Discussion integrates experimental TSHR A-subunit immunization and orbital fibroblast studies as biological context, without claiming that these mechanisms were established by our blood genetic data. The manuscript distinguishes susceptibility evidence from treatment efficacy and does not claim TED-specific effects within GD.')
(O/'Cover_Letter_Draft.md').write_text(cover,encoding='utf-8');rebuild('Cover_Letter_Draft',cover)
check=Document(S/'STROBE_MR_Checklist.docx')
replacements={
'1 Title and abstract':'Association-focused MR findings and secondary colocalization stated in Abstract; title identifies genes and diseases.',
'4 Study design and sources':'Methods; Table 1; original vs post hoc analyses and revised reporting emphasis distinguished. No prospective registration.',
'11 Main results':'Results: TSHR/IGF1R associations and corroborating genes; Table 2 (15 comparisons), Table 3A–C, Figures 1–3 and full workbook.',
'14 Key results':'Opening Discussion and Conclusions: consistent TSHR associations and limited IGF1R evidence; no expression-mediated causal proof.',
'16 Interpretation':'Eight Discussion paragraphs; A-subunit and orbital fibroblast evidence is external biological context. No GD/TED independence or treatment-prediction claim.'}
for row in check.tables[0].rows:
    if row.cells[0].text in replacements:
        p=row.cells[1].paragraphs[0];p.clear();inline(p,replacements[row.cells[0].text])
check.save(O/'STROBE_MR_Checklist.docx')

wordcount=lambda s:len(re.sub(r'\[[0-9, ]+\]','',s.replace('*','')).split())
abstract_n=wordcount(ABSTRACT)
main=t.split('## Introduction\n',1)[1].split('## Declarations',1)[0]
main_n=wordcount('\n'.join(l for l in main.splitlines() if not l.startswith('#')))
audit={'stage':'AUTHORED_PENDING_RENDER_AND_FINAL_CHECKS','source_package':str(S),'source_manifest_verified':True,'abstract_words_including_labels':abstract_n,'main_words_excluding_headings':main_n,'discussion_paragraphs':len(DISCUSSION.split('\n\n')),'limitation_paragraphs':len(t.split('### Limitations\n',1)[1].split('### Conclusions',1)[0].strip().split('\n\n')),'references':len(order),'reference_old_to_new':mapping,'no_new_statistical_analysis':True}
(V/'editorial_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
(V/'manuscript_changes.diff').write_text(''.join(difflib.unified_diff(old.splitlines(True),t.splitlines(True),fromfile='20261002/Manuscript.md',tofile='20261003/Manuscript.md')),encoding='utf-8')
unchanged={}
for p in O.rglob('*'):
    if not p.is_file() or 'Verification' in p.parts:continue
    rel=p.relative_to(O)
    if (S/rel).exists() and sha(p)==sha(S/rel):unchanged[str(rel).replace('\\','/')]=sha(p)
(V/'unchanged_artifacts.json').write_text(json.dumps(unchanged,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
assert abstract_n<=250,abstract_n
assert len(DISCUSSION.split('\n\n'))==8

from pathlib import Path
from copy import deepcopy
import ast,re,json,sys,shutil,difflib
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1]
S=R/'outputs/TED_Manuscript_Reframed_20261002';O=R/'outputs/TED_Association_Manuscript_20261003'
shutil.copytree(W/'render',W/'render_before_polish')
t=(O/'Manuscript.md').read_text(encoding='utf-8')
a,b=t.split('**Keywords:**',1)
a=a.replace('UK Biobank hyperthyroidism','UK Biobank (UKB) hyperthyroidism').replace('per approximately one SD','per approximately one standard deviation')
t=a+'**Keywords:**'+b
t=t.replace('PLINK clumping used r² < 0.001','PLINK clumping used a linkage disequilibrium (LD) threshold of r² < 0.001')
t=t.replace('Five MHC-region genes among','Five major histocompatibility complex (MHC) region genes among')
t=t.replace('within a broader autoimmune gene screen','within a broader screen of druggable genes')
(O/'Manuscript.md').write_text(t,encoding='utf-8')
tree=ast.parse((W/'build_v1.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in ['inline','rebuild']],type_ignores=[]),str(W/'build_v1.py'),'exec'))
rebuild('Manuscript',t)
d=Document(O/'STROBE_MR_Checklist.docx');original=Document(S/'STROBE_MR_Checklist.docx')
for row,orig in zip(d.tables[0].rows,original.tables[0].rows):
    for c,s in zip(row.cells,orig.cells):
        props=s.paragraphs[0].runs[0]._r.rPr
        for p in c.paragraphs:
            for run in p.runs:
                if run._r.rPr is not None:run._r.remove(run._r.rPr)
                if props is not None:run._r.insert(0,deepcopy(props))
d.save(O/'STROBE_MR_Checklist.docx')
audit=json.loads((O/'Verification/editorial_audit.json').read_text(encoding='utf-8'))
audit['abstract_words_including_labels']=len(a.split('## Abstract\n',1)[1].replace('*','').split())
main=t.split('## Introduction\n',1)[1].split('## Declarations',1)[0]
audit['main_words_excluding_headings']=len(re.sub(r'\[[0-9, ]+\]','','\n'.join(l for l in main.splitlines() if not l.startswith('#'))).replace('*','').split())
assert audit['abstract_words_including_labels']<=250
(O/'Verification/editorial_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
old=(S/'Manuscript.md').read_text(encoding='utf-8')
(O/'Verification/manuscript_changes.diff').write_text(''.join(difflib.unified_diff(old.splitlines(True),t.splitlines(True),fromfile='20261002/Manuscript.md',tofile='20261003/Manuscript.md')),encoding='utf-8')
print(audit['abstract_words_including_labels'],audit['main_words_excluding_headings'])

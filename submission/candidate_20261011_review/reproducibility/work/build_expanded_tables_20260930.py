from pathlib import Path
import sys,json,shutil,hashlib,re
W=Path(__file__).resolve().parent;sys.path.insert(0,str(W/'pydeps'))
import pandas as pd,numpy as np
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=W/'expanded_analysis_20260930';R=T/'results';out=W.parent/'outputs/TED_TwoStage_Analysis_20260930';out.mkdir(exist_ok=True,parents=True)
mr=pd.read_csv(R/'Stage1_13_discovery_genes.csv');primary=pd.read_csv(R/'Stage2_14_genes_primary.csv');allpp=pd.read_csv(R/'Stage2_all_504_colocalization_settings.csv');summary=pd.read_csv(R/'Stage2_integrated_summary.csv');full=pd.read_csv(R/'Stage1_all_2544_genes_3_outcomes.csv')
genes=summary.gene.tolist();hits=genes[:-1];mhc=['HLA-A','HLA-DQA2','C4A','TUBB','PSMB8'];ocs={'BBJ_Graves':'BBJ Graves disease','UKB_hyperthyroid':'UKB hyperthyroidism','FinnGen_GO':'FinnGen Graves ophthalmopathy'}
assert (R/'independent_R_verification.txt').read_text().strip().endswith('PASS')
cmp=allpp[allpp.scenario.eq('original_reference')].merge(allpp[allpp.scenario.eq('available_eqtlgen')],on=['gene','outcome','p12'],suffixes=('_ref','_af'))
cross=int(((cmp['PP.H4_ref']>=.8)!=(cmp['PP.H4_af']>=.8)).sum());delta=float((cmp['PP.H4_ref']-cmp['PP.H4_af']).abs().max())
report=json.loads((R/'verification.json').read_text());report.update(status='PASS',independent_R='504 rows; max absolute posterior difference 5.51e-14; independent base-R equations',frequency_threshold_crossings=cross,frequency_max_abs_H4_delta=delta)
(R/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf8')
wide=full.pivot(index='gene_symbol',columns='outcome',values=['n_iv','beta','se','or','or_ci_lower','or_ci_upper','pvalue','status']);wide.columns=['__'.join(c) for c in wide.columns];wide.to_csv(out/'Stage1_ALL_2544_genes.csv',encoding='utf-8-sig')
for f in ['Stage1_13_discovery_genes.csv','Stage2_14_genes_primary.csv','Stage2_all_504_colocalization_settings.csv','Stage2_integrated_summary.csv']:
 pd.read_csv(R/f).to_csv(out/f,index=False,encoding='utf-8-sig')
def newdoc():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.left_margin=sec.right_margin=Inches(.65);sec.top_margin=sec.bottom_margin=Inches(.65)
 st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(11);st.paragraph_format.space_after=Pt(6)
 for nm in ['Title','Heading 1','Heading 2']:
  d.styles[nm].font.name='Times New Roman';d.styles[nm].font.color.rgb=RGBColor(0,0,0)
 f=sec.footer.paragraphs[0];f.alignment=2;r=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
 return d
def pval(v):
 if v<.001:
  mant,exp=f'{v:.2e}'.split('e');return mant+'×10'+str(int(exp)).translate(str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹'))
 return f'{v:.3g}'
def prob(v):return '<0.001' if v<.001 else f'{v:.3f}'
def table(d,headers,rows,widths,bold_cells=set()):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.style='Table Grid'
 tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)*1440)))
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for cell,txt in zip(t.rows[0].cells,headers):cell.text=txt
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for ri,row in enumerate(t.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci])
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=ri==0 or (ri-1,ci) in bold_cells
 return t
d=newdoc()
for idx,(oc,label) in enumerate(ocs.items()):
 if idx:d.add_page_break()
 d.add_heading('Table A. Stage 1: all 13 discovery genes',1);d.add_paragraph(label,style='Heading 2')
 rows=[];bold=set()
 for g in hits:
  r=mr[(mr.gene_symbol==g)&(mr.outcome==oc)].iloc[0];digits=3 if g=='MAPKAPK5' else 2
  if pd.isna(r.pvalue):
   rows.append([g,'NA','Not estimable','NA','MHC' if g in mhc else 'Other locus']);continue
  rows.append([g,int(r.n_iv),f'{r["or"]:.{digits}f} ({r.or_ci_lower:.{digits}f}–{r.or_ci_upper:.{digits}f})',pval(r.pvalue),'MHC' if g in mhc else 'Other locus'])
  if r.pvalue<(.05/2544 if idx==0 else .05):bold.add((len(rows)-1,3))
 table(d,['Gene','IVs','OR (95% CI)','P value','Region'],rows,[1.1,.45,2.2,1.25,1.95],bold)
 d.add_paragraph('All 13 genes were selected solely by BBJ discovery P < 0.05/2,544 = 1.965×10⁻⁵. Their results are displayed in all three disease datasets, regardless of significance. The complete 2,544-gene screen, including non-estimable results, is supplied in Stage1_ALL_2544_genes.csv.')
 d.add_paragraph('ORs are per approximately one SD higher reconstructed, genetically proxied blood expression. One IV: Wald ratio; multiple IVs: multiplicative random-effects IVW with under-dispersion correction. Two-sided P values. Bold: corrected discovery significance in BBJ; nominal P < 0.05 in UKB/FinnGen. Phenotypes differ across datasets; these are not interchangeable replications. Original reference-frequency estimates are shown.')
 d.add_paragraph('MHC = major histocompatibility complex; IV = instrumental variant. Correlated regional associations do not establish independent causal genes. IGF1R was selected for biological comparison and did not meet the discovery threshold; it is included in Stage 2 and the full Stage 1 data.')
d.save(out/'Table_A_Stage1_Discovery.docx')
d=newdoc()
for idx,(oc,label) in enumerate(ocs.items()):
 if idx:d.add_page_break()
 d.add_heading('Table B. Stage 2: all 14 evaluated genes',1);d.add_paragraph(label,style='Heading 2')
 rows=[];bold=set()
 for g in genes:
  r=primary[(primary.gene==g)&(primary.outcome==oc)].iloc[0];z=allpp[(allpp.gene==g)&(allpp.outcome==oc)]
  low=z[z.scenario.eq('original_reference')&z.p12.eq(1e-6)].iloc[0];af=z[z.scenario.eq('available_eqtlgen')&z.p12.eq(1e-5)].iloc[0]
  rows.append([g,int(r.n_overlap),prob(r['PP.H2']),prob(r['PP.H3']),prob(r['PP.H4']),prob(low['PP.H4']),prob(af['PP.H4'])])
  for col,val in [(4,r['PP.H4']),(5,low['PP.H4']),(6,af['PP.H4'])]:
   if val>=.8:bold.add((len(rows)-1,col))
 table(d,['Gene','SNPs','H2','H3','H4: main','H4: lower prior','H4: cohort AF'],rows,[1.1,.55,.7,.7,1,1.2,1.7],bold)
 d.add_paragraph('All 13 BBJ discovery genes plus the biologically selected IGF1R comparator were evaluated. Five MHC genes (HLA-A, HLA-DQA2, C4A, TUBB and PSMB8) were added post hoc following review; the existing nine-gene calculations were reproduced before extension.')
 d.add_paragraph('SNPs: overlapping variants in the primary reference-frequency analysis. H2: expression-only model; H3: distinct disease/expression causal variants; H4: shared-variant model. Values are posterior probabilities, not P values. H0/H1 account for the remaining probability. Values <0.001 are not zero; full-precision H0–H4 values are supplied in the CSV.')
 d.add_paragraph('Main: p1 = p2 = 10⁻⁴, p12 = 10⁻⁵. Lower prior: p12 = 10⁻⁶. Cohort AF: eQTLGen allele frequencies, all usable variants, p12 = 10⁻⁵; includes frequency and coverage changes. Same-variant frequency comparisons and intermediate p12 = 5×10⁻⁶ are also provided in the 504-setting data. Bold H4 ≥ 0.80.')
 d.add_paragraph('Only TSHR meets H4 ≥ 0.80 in both BBJ and FinnGen under the main prior. None meets both under p12 = 10⁻⁶. H3 dominance is conditional on the single-causal-variant model: it neither proves absence of a gene effect nor resolves multiple signals, particularly within MHC. No multi-signal colocalization was performed.')
d.save(out/'Table_B_Stage2_Colocalization.docx')
notes=f'''# 1·2단계 분석 확장 — 2026-09-30

1단계 전체 2,544개 유전자 × 3개 질환의 7,632개 조합을 정리했습니다. BBJ 2,234개, UKB 2,505개, FinnGen 2,480개는 추정 가능하며 나머지는 음성 결과가 아닌 미추정입니다. 기존 원본 MR 추정치를 그대로 사용했고 MR 전체를 새로 계산한 것은 아닙니다.

Table A는 BBJ 보정 기준을 통과한 13개 유전자 전부의 OR·95% CI·P값을 질환별로 표시합니다. 전체 2,544개 결과는 CSV 한 행에 한 유전자씩 정리했습니다.

2단계는 13개 발견 유전자 + IGF1R = 14개를 대상으로 실제 다시 계산했습니다. 기존 9개 결과를 재현하고, 미평가였던 MHC 5개에 동일한 beta/varbeta ABF 공위치화 계산을 확장했습니다. 14 × 3질환 × 3사전확률 × 4빈도/가용변이 조건 = 504개 설정입니다. 노출 분산 sdY=1, p1=p2=10⁻⁴, p12=10⁻⁵·5×10⁻⁶·10⁻⁶; 질환은 dataset1, 발현은 dataset2입니다. UKB beta·SE는 기존과 동일하게 μ(1−μ), μ=3731/484598로 나눴습니다.

MHC 원본 cis-eQTL 107,154행을 유의성 필터 없이 사용했습니다. 각 유전자의 P값이 0.9를 넘는 영역 전체 통계를 포함하며, eQTLGen GRCh37 GenePos ±1 Mb 이내임을 확인했습니다. GWAS는 GRCh38 원본에서 rsID로 연결하고 염색체와 대립유전자 쌍을 확인했습니다. 다른 빌드 좌표를 직접 비교하지 않았습니다. 빈도 파일의 GRCh37 좌표 및 allele-count 기반 빈도를 검증했습니다. ±1 Mb의 중심은 MHC 원본 GenePos로, 기존 9개 유전자의 동결된 입력 영역을 소급 변경하지 않았습니다.

## 결과

- MHC 5개는 모든 질환에서 H4 <0.80입니다. 가장 큰 값은 C4A/FinnGen의 H4=0.03024, H3=0.96976입니다. 나머지 MHC 조합도 H3가 우세합니다.
- 기본 사전확률에서 BBJ와 FinnGen 동시 기준을 통과한 유전자는 TSHR 하나입니다. p12=10⁻⁶에서는 TSHR/BBJ H4가 0.661로 내려가므로 동시 통과 유전자는 없습니다.
- 빈도 및 가용변이 치환으로 H4=0.80 판정이 바뀐 설정은 {cross}/126개, 최대 H4 절대 변화는 {delta:.8f}입니다.
- 기존 324개 설정의 최대 절대 차이는 {report['existing_324_max_abs_delta']:.3g}입니다. 독립적인 base-R의 별도 합산 방법으로 전체 504개 설정을 재검증했고 최대 차이 5.51×10⁻¹⁴로 일치했습니다. 로컬 coloc 패키지는 사용할 수 없어 패키지를 직접 실행한 것으로 표현하지 않습니다.

## 해석과 원고 반영 범위

유의한 MR 연관성과 공위치화 기준 충족은 별개입니다. MHC에서 H3가 높다는 것은 단일 인과변이 모형 아래 서로 다른 변이 설명을 더 지지한다는 뜻입니다. 이 계산으로 다중 신호를 해결하거나 독립적인 원인 유전자를 확정한 것은 아닙니다. 공식 방법 설명: https://chr1swallace.github.io/coloc/articles/a01_intro.html

이번 5개 유전자 추가는 리뷰 후 시행한 post hoc 확장입니다. 분석 대상 확대와 음성 결과를 모두 보고해야 합니다. 기존 제출 패키지는 9개/81개 설정 설명이 남아 있으므로, 이번 표를 최종 편입할 때 본문·Figure 3·보충문서·통합 데이터의 14개/126개 설정으로의 동기화가 필요합니다. 이 폴더는 확장 분석 검토본이며 기존 제출 패키지를 덮어쓰지 않았습니다.
'''
(out/'READ_ME_KO.md').write_text(notes,encoding='utf8')
shutil.copy2(R/'verification.json',out/'verification.json')
print('Created',out,'frequency threshold crossings',cross,'max delta',delta)

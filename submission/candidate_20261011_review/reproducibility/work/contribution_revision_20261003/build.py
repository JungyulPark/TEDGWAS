from pathlib import Path
import sys, re, json, hashlib, shutil, difflib
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
W = Path(__file__).resolve().parent
R = W.parents[1]
S = R/'outputs/TED_Association_Manuscript_20261003'
O = R/'outputs/TED_Research_Contribution_20261003'
V = O/'Verification'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
assert not O.exists(), 'Preserve existing revisions; do not overwrite.'
source_manifest = json.loads((S/'PACKAGE_SHA256.json').read_text(encoding='utf-8'))
for name, digest in source_manifest.items():
    assert sha(S/name) == digest, name
O.mkdir(); V.mkdir()
for p in S.iterdir():
    if p.name in ['Verification', 'PACKAGE_SHA256.json', 'Full_Review.pdf', 'READ_FIRST_KO.md']:
        continue
    if p.is_dir(): shutil.copytree(p, O/p.name)
    else: shutil.copy2(p, O/p.name)
shutil.copytree(S/'Verification', V/'Previous_revision')
shutil.copy2(S/'PACKAGE_SHA256.json', V/'Previous_revision_manifest.json')

old = (S/'Manuscript.md').read_text(encoding='utf-8')
intro_old = old.split('## Introduction\n',1)[1].split('## Methods\n',1)[0].strip()
ip = intro_old.split('\n\n')
assert len(ip)==4
method_explanation = ip[2].split(' Prior MR and multi-omics studies')[0]
prior_and_gap = """Previous genetic studies have already implicated *TSHR* in GD and Graves ophthalmopathy. Ji and colleagues integrated blood methylation, expression and protein associations with GD, reporting *TSHR* among the expression-associated genes also supported in a FinnGen ophthalmopathy dataset [9]. Other MR work has examined inflammatory proteins in ophthalmopathy [10], and recent integrative research has investigated *TSHR* across genetic and transcriptomic data in GD and ophthalmopathy [33]. These findings establish a basis for comparing clinically relevant receptors, but do not determine whether evidence for their blood expression is consistently supported across disease definitions under the same analysis framework. In particular, limited evidence for *IGF1R* could reflect imprecision or dependence on particular instruments, while a strong MR association for *TSHR* does not by itself establish that expression and disease share a causal variant. The remaining question is how robustly these two expression traits are associated with GD and related phenotypes, and how far those associations support biological interpretation."""
aim = """To address this question, we compared genetically proxied blood expression of *TSHR* and *IGF1R* within a screen of genes encoding potentially druggable proteins, using *CTLA4* and other associated genes as context. We used Biobank Japan (BBJ) GD for discovery, compared the original findings with UK Biobank (UKB) hyperthyroidism and FinnGen ophthalmopathy, and subsequently added FinnGen GD and a stricter ophthalmopathy definition. We evaluated effect direction, precision and statistical support across these datasets, then examined dependence on individual instruments, allele frequencies and colocalization assumptions. Our aim was to characterize the consistency and limits of expression-related genetic evidence for the two receptors. The analyses distinguish GD from broader hyperthyroidism and population-controlled ophthalmopathy; they do not directly estimate TED susceptibility among patients who already have GD or predict treatment response."""
intro = '\n\n'.join([ip[0], ip[1], method_explanation, prior_and_gap, aim])

discussion_old = old.split('## Discussion\n',1)[1].split('### Limitations\n',1)[0].strip()
dp = discussion_old.split('\n\n')
assert len(dp)==8
main_finding = """This study addresses how consistently blood expression-related genetic evidence supports two clinically relevant receptors across GD and related disease definitions. *TSHR* showed associations with lower disease odds in both GD datasets, with concordant directions in broader hyperthyroidism and ophthalmopathy. *IGF1R* estimates consistently indicated higher odds, but only the BBJ and UKB estimates reached uncorrected *P* < 0.05, and none met the discovery threshold. The contribution is a comparison of the reproducibility and limitations of these associations within the same analytical framework. It does not establish a difference in biological importance between the receptors, and statistical significance in one gene but not another is not a formal test of different effects."""
prior_comparison = """Our findings extend an existing body of genetic evidence. Ji and colleagues had already linked blood *TSHR* expression to GD and reported support in FinnGen ophthalmopathy [9]; recent integrative work has also investigated *TSHR* in GD and ophthalmopathy [33]. We add a focused comparison with *IGF1R*, a GD comparison spanning BBJ and FinnGen, and an assessment of both genes alongside the wider druggable-gene screen. The result is consistent association evidence for *TSHR*, limited evidence for *IGF1R*, and an explicit distinction between strong MR associations and variable regional support for expression-mediated causality. This comparison uses shared public resources, including eQTLGen and FinnGen, so agreement with earlier work is not wholly independent replication. Different MR methods, FinnGen releases and colocalization settings also preclude interpreting differences between reports as direct contradictions. Our additional evidence concerns the consistency and sensitivity of the associations, rather than discovery of a new susceptibility locus."""
orbital_context = """Experimental and clinical studies provide context for relating TSHR-directed immunity to orbital disease. TSHR can undergo cleavage and release its extracellular A-subunit. Expression of this subunit in an adenoviral immunization model induced greater thyroid-stimulating autoimmunity than a minimally cleaving receptor [28]; plasmid immunization subsequently reproduced orbital adipogenesis, fibrosis and muscle pathology in mice [29]. Daser and colleagues further reported enhanced inflammatory and hyaluronan responses to TNFα in orbital fibroblasts from TSHR-immunized mice [30]. Graves ophthalmopathy has also been documented in primarily euthyroid and hypothyroid patients [31], supporting a distinction between thyroid functional status and orbital immune activity. These observations make the ophthalmopathy associations biologically relevant, but they do not show that an A-subunit shed by human orbital fibroblasts, acting alone, causes TED independently of GD. Our instruments do not measure receptor cleavage or antibody activity. The ophthalmopathy comparisons use population controls and overlapping FinnGen endpoints, so neither their MR estimates nor their higher colocalization probabilities establish a TED-specific effect within GD or a difference between GD and TED."""
coloc = dp[6].replace(
    'Thus, association strength and evidence for a shared causal variant were not interchangeable.',
    'Thus, association strength and evidence for a shared causal variant were not interchangeable. Ji and colleagues used a shared-variant prior of 5×10⁻⁵ and a PP.H4 threshold above 0.50 [9], whereas our primary settings used 10⁻⁵ and at least 0.80. These differences further discourage comparing binary colocalization labels across studies without considering the model settings.'
)
implications = """The sensitivity analyses identify which qualifications matter for the central comparison. Frequency substitution preserved the original BBJ discovery set, whereas *C4B* and *HLA-G* changed significance classifications through altered effect scaling and instrument weighting. Instrument omission identified concentrated evidence for *CTLA4* and reduced statistical support for *IGF1R*. These checks reuse the same data and do not supply independent replication. For researchers evaluating TSHR and IGF-1R, the added answer is that blood expression-related associations are more consistently supported for *TSHR*, while its tissue mechanism and the extent of an inherited *IGF1R* contribution remain unresolved. Independent expression instruments and tissue-relevant regulatory data would help address these uncertainties. Comparing GD patients with and without TED, while accounting for follow-up and selection into the GD group, would address orbital susceptibility directly. The present results inform interpretation and follow-up research, but do not guide treatment selection."""
discussion = '\n\n'.join([main_finding, prior_comparison, dp[1], dp[2], orbital_context, dp[5], coloc, implications])

t = old.replace(intro_old, intro).replace(discussion_old, discussion)
body, refs = t.split('## References\n',1)
refmap = {int(n):s.strip() for n,s in re.findall(r'(?ms)^(\d+)\. (.*?)(?=^\d+\. |\Z)',refs)}
refmap[33] = 'Zhang H, Jiang S, Zhu T, Liu Y, Li J, Fang S, Li Y, Sun J, He X, Tong C, et al. Thyroid-stimulating hormone receptor mediates peripheral-central neuroimmune crosstalk in autoimmune thyroid diseases. BMC Med. 2026;24:404. doi:10.1186/s12916-026-04957-y.'
order=[]
for group in re.findall(r'\[([0-9, ]+)\]',body):
    for n in map(int,re.findall(r'\d+',group)):
        assert n in refmap
        if n not in order:order.append(n)
assert set(order)==set(refmap)
mapping={n:i+1 for i,n in enumerate(order)}
body=re.sub(r'\[([0-9, ]+)\]',lambda m:'['+', '.join(str(mapping[int(n)]) for n in re.findall(r'\d+',m[1]))+']',body)
t=body+'## References\n\n'+'\n\n'.join(f'{mapping[n]}. {refmap[n]}' for n in order)+'\n'
(O/'Manuscript.md').write_text(t,encoding='utf-8')

def inline(p,text):
    text=text.replace('\\*','＊')
    for x in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text):
        if not x:continue
        if x.startswith('**') and x.endswith('**'):p.add_run(x[2:-2]).bold=True
        elif x.startswith('*') and x.endswith('*'):p.add_run(x[1:-1]).italic=True
        else:p.add_run(x.replace('＊','*'))

def rebuild(name,text):
    # Retain the established manuscript layout, styles, headers and section settings.
    d=Document(S/(name+'.docx'))
    for e in list(d._element.body):
        if e.tag!=qn('w:sectPr'):d._element.body.remove(e)
    for block in re.split(r'\n\s*\n',text.strip()):
        if block.startswith('# '):p=d.add_paragraph(style='Title');inline(p,block[2:])
        elif block.startswith('## '):
            p=d.add_paragraph(style='Heading 1');inline(p,block[3:])
            if block[3:] in ['Abstract','Introduction','References']:p.paragraph_format.page_break_before=True
        elif block.startswith('### '):p=d.add_paragraph(style='Heading 2');inline(p,block[4:])
        else:p=d.add_paragraph();inline(p,block.replace('  \n','\n'))
        p.paragraph_format.widow_control=True
        for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)
    for e in list(d._element.xpath('.//w:pBdr')):e.getparent().remove(e)
    d.save(O/(name+'.docx'))

rebuild('Manuscript',t)
cover=(S/'Cover_Letter_Draft.md').read_text(encoding='utf-8')
cp=cover.split('\n\n')
assert cp[3].startswith('The clinical question')
cp[3]="Prior genetic studies have already implicated TSHR in Graves disease and ophthalmopathy. Our study examines how consistently blood expression-related evidence supports TSHR and IGF1R when evaluated within the same analysis framework across distinct disease definitions. We screened 2,544 potentially druggable genes and compared five disease datasets, explicitly acknowledging overlapping FinnGen participants. TSHR associations were consistent in direction, whereas IGF1R evidence remained limited. CTLA4, HLA-DQA2 and TUBB also passed the full-screen threshold in both GD datasets. Instrument and allele-frequency analyses qualify the robustness of these findings; secondary colocalization limits claims of expression-mediated causality. The contribution is a focused comparison and sensitivity assessment using public resources, not a claim to have discovered TSHR susceptibility or established a new orbital mechanism."
cp[4]="The manuscript distinguishes inherited expression-related susceptibility from treatment efficacy and does not claim a TED-specific effect within GD. External TSHR A-subunit and orbital fibroblast experiments provide biological context, while the small in-house orbital tissue series remains descriptive. Three main tables, three main figures, a supplementary document with three tables and two figures, and one detailed results workbook accompany the manuscript."
cover='\n\n'.join(cp)
(O/'Cover_Letter_Draft.md').write_text(cover,encoding='utf-8');rebuild('Cover_Letter_Draft',cover)

abstract=t.split('## Abstract\n',1)[1].split('**Keywords:**',1)[0]
main=t.split('## Introduction\n',1)[1].split('## Declarations',1)[0]
wordcount=lambda s:len(re.sub(r'\[[0-9, ]+\]','',s.replace('*','')).split())
audit={'status':'AUTHORED_PENDING_VERIFICATION','source_manifest_verified':True,'source_package':str(S),'abstract_words':wordcount(abstract),'main_words':wordcount('\n'.join(x for x in main.splitlines() if not x.startswith('#'))),'intro_paragraphs':5,'discussion_paragraphs':8,'limitation_paragraphs':2,'references':33,'reference_old_to_new':mapping,'new_statistical_analysis':False,'changed_sections':['Introduction','Discussion','References','Cover letter'],'preserved_sections':['Abstract','Methods apart from citation renumbering','Results apart from citation renumbering','Limitations','Conclusions','Declarations'],'literature_review_type':'Targeted primary-source comparison; not a systematic novelty review'}
(V/'editorial_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
(V/'manuscript_changes.diff').write_text(''.join(difflib.unified_diff(old.splitlines(True),t.splitlines(True),fromfile='Association_20261003/Manuscript.md',tofile='Contribution_20261003/Manuscript.md')),encoding='utf-8')
(V/'cover_letter_changes.diff').write_text(''.join(difflib.unified_diff((S/'Cover_Letter_Draft.md').read_text(encoding='utf-8').splitlines(True),cover.splitlines(True))),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))

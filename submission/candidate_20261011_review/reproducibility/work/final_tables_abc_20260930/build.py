from pathlib import Path
import csv, json, math, re, shutil, hashlib
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B = Path(__file__).resolve().parent
ROOT = B.parent.parent
SOURCE = ROOT/'outputs/TED_Paper_Tables_Figures_20260930_v2'
OUT = ROOT/'outputs/TED_Paper_Tables_Figures_20260930_v3'
OUT.mkdir(parents=True, exist_ok=True)
for source in SOURCE.rglob('*'):
    if source.is_file() and source.name not in {'Table3.docx','Table3.pdf','All_Tables.pdf','All_Figures.pdf','PACKAGE_SHA256.json','READ_ME_KO.md'}:
        dest = OUT/source.relative_to(SOURCE)
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, dest)

old_doc = Document(SOURCE/'Tables/Table3.docx')
assert len(old_doc.tables) == 3
old_rows = [[[c.text for c in r.cells] for r in t.rows[2:]] for t in old_doc.tables]
assert all(len(rows) == 14 for rows in old_rows)
genes = [r[0].rstrip('†') for r in old_rows[0]]
assert all([r[0].rstrip('†') for r in rr] == genes for rr in old_rows)
with (SOURCE/'Detailed_Data/Table3_integrated_14genes.csv').open(encoding='utf-8-sig', newline='') as f:
    data = {(r['gene'],r['outcome']):r for r in csv.DictReader(f)}
outcomes = ['BBJ_Graves','UKB_hyperthyroid','FinnGen_GO']

d = Document()
sec = d.sections[0]
# The three disease blocks form ONE native Word table on ONE wide page.
# Preserve readable numeric type rather than shrinking a 22-column table to A4.
sec.page_width, sec.page_height = Mm(420), Mm(297)
widths = [.92] + [.39,1.40,.80,.51,.51,.51,.60]*3
sec.left_margin = sec.right_margin = Inches((420/25.4-sum(widths))/2)
sec.top_margin = sec.bottom_margin = Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st = d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    for node in list(st._element.iter(qn('w:pBdr'))): node.getparent().remove(node)
    rpr=st._element.get_or_add_rPr(); fonts=rpr.find(qn('w:rFonts'))
    if fonts is None: fonts=OxmlElement('w:rFonts'); rpr.append(fonts)
    for attr in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+attr),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: fonts.attrib.pop(qn('w:'+attr),None)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(15)
d.styles['Title'].font.bold=True
d.styles['Title'].paragraph_format.space_after=Pt(10)
d.add_paragraph('Table 3 Integrated Mendelian randomization and colocalization results across 14 genes',style='Title')

pattern=re.compile(r'\b(?:'+'|'.join(re.escape(g) for g in sorted(genes,key=len,reverse=True))+r'|P)\b')
def rich(p,text,size=10.5,bold=False):
    pos=0
    for m in pattern.finditer(text):
        for s,italic in [(text[pos:m.start()],False),(m[0],True)]:
            if s:
                r=p.add_run(s);r.italic=italic;r.bold=bold;r.font.name='Times New Roman';r.font.size=Pt(size)
        pos=m.end()
    if text[pos:]:
        r=p.add_run(text[pos:]);r.bold=bold;r.font.name='Times New Roman';r.font.size=Pt(size)

def edge(cell,side,val='single',sz=4,color='D5D5D5'):
    pr=cell._tc.get_or_add_tcPr(); bs=pr.find(qn('w:tcBorders'))
    if bs is None:bs=OxmlElement('w:tcBorders');pr.append(bs)
    e=bs.find(qn('w:'+side))
    if e is None:e=OxmlElement('w:'+side);bs.append(e)
    for k,v in [('val',val),('sz',str(sz)),('color',color)]:e.set(qn('w:'+k),v)

t=d.add_table(rows=3,cols=22);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)*1440)))
for i,w in enumerate(widths):t.columns[i].width=Inches(w)
for row in t.rows:
    for i,c in enumerate(row.cells):c.width=Inches(widths[i])
    rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
t.cell(0,0).merge(t.cell(2,0)).text='Gene'
for i,label in enumerate(['A  BBJ Graves disease','B  UKB hyperthyroidism','C  FinnGen Graves ophthalmopathy']):
    start=1+7*i
    t.cell(0,start).merge(t.cell(0,start+6)).text=label
    t.cell(1,start).merge(t.cell(1,start+2)).text='Mendelian randomization'
    t.cell(1,start+3).merge(t.cell(1,start+6)).text='Colocalization'
    for j,label2 in enumerate(['SNPs\nn','OR (95% CI)','P value','PP.H2','PP.H3','PP.H4','PP.H4\nlower prior']):
        t.cell(2,start+j).text=label2

display=[]
for i,g in enumerate(genes):
    vals=[old_rows[0][i][0]]
    for block in old_rows: vals.extend(block[i][1:])
    row=t.add_row()
    for c,txt in zip(row.cells,vals):c.text=txt
    display.append(vals)

seen=set()
for ri,row in enumerate(t.rows):
    no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
    for ci,c in enumerate(row.cells):
        if c._tc in seen:continue
        seen.add(c._tc)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ri>=3:c.width=Inches(widths[ci])
        pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
        for side,val in [('top',115 if ri>=3 else 90),('bottom',115 if ri>=3 else 90),('left',32),('right',32)]:
            n=OxmlElement('w:'+side);n.set(qn('w:w'),str(val));n.set(qn('w:type'),'dxa');mar.append(n)
        pr.append(mar)
        for side in ['top','bottom','left','right']:edge(c,side,'nil')
        if ri==0:edge(c,'top',sz=8,color='000000')
        edge(c,'bottom',sz=7 if ri in [2,16] else 3,color='000000' if ri in [2,16] else 'D5D5D5')
        if ci in [8,15]:edge(c,'left',sz=5,color='A0A0A0')
        bold=ri<3
        if ri>=3 and ci>0:
            block=(ci-1)//7;sub=(ci-1)%7;r=data[(genes[ri-3],outcomes[block])]
            if sub==2 and r['pvalue']:bold=float(r['pvalue'])<(.05/2544 if block==0 else .05)
            if sub==5:bold=float(r['PP.H4'])>=.8
            if sub==6:bold=float(r['H4_lower_prior_1e-6'])>=.8
        txt=c.text
        for p in c.paragraphs:
            p.clear();p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            rich(p,txt,size=11 if ri==0 else 9.5 if ri<3 else 10.5,bold=bold)

notes=[
 'The 13 BBJ discovery genes and IGF1R† are displayed together in panels A–C. †Biologically selected comparator that did not meet the BBJ discovery threshold. Genes are ordered by BBJ discovery P values, with IGF1R last. BBJ, BioBank Japan; UKB, UK Biobank; GD, Graves disease; TED, thyroid eye disease.',
 'SNPs, n is the number of genetic instrumental variables (IVs; single-nucleotide polymorphisms) used in each MR estimate, not the number of participants. One SNP: Wald ratio; multiple SNPs: multiplicative random-effects inverse-variance weighting with under-dispersion correction. ORs are per approximately one SD higher reconstructed, genetically proxied blood expression. CI, confidence interval. P values are two-sided.',
 'Bold P values denote P < 0.05/2,544 = 1.965×10⁻⁵ in BBJ and nominal P < 0.05 in UKB/FinnGen. PP, posterior probability; H, hypothesis. Dataset 1 is disease and dataset 2 is blood expression: H2, expression only; H3, distinct causal variants; H4, a shared causal variant. H0 (neither trait) and H1 (disease only) account for the undisplayed probability. Values <0.001 are not zero.',
 'Primary probabilities use p₁ = p₂ = 10⁻⁴ and p₁₂ = 10⁻⁵. Lower prior means p₁₂ = 10⁻⁶. Bold PP.H4 values are ≥0.80 and are not P-value significance. Only TSHR meets this threshold in both BBJ and FinnGen under the primary prior; none meets both under the lower prior. Full-precision H0–H4, regional variant counts and all frequency/prior sensitivity settings are retained in the detailed data.',
 'MHC genes HLA-A, HLA-DQA2, C4A, TUBB and PSMB8 were added to colocalization post hoc. H3 dominance is conditional on the single-causal-variant model and does not resolve multiple signals or exclude a gene effect. Disease definitions differ; FinnGen population controls cannot isolate TED susceptibility within GD. CTLA4 estimates in UKB/FinnGen depend mainly on rs13030124 in leave-one-out analysis. NA for VKORC1 FinnGen indicates non-estimable MR, not a negative result; regional colocalization remained estimable.'
]
for i,note in enumerate(notes):
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(7 if i==0 else 0);p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.03
    rich(p,note,size=10)
d.save(OUT/'Tables/Table3.docx')
(B/'display_rows.json').write_text(json.dumps(display,ensure_ascii=False,indent=2),encoding='utf8')

# Rebuild the sole inconsistent figure title from its verified plotting source.
original=ROOT/'work/review_8dce23c/submission/candidate_20260911_review'
record=json.loads((original/'provenance/leave_one_out_figure_sources.json').read_text(encoding='utf8'))
for ext in ['pdf','png']:
    assert record[ext+'_sha256']==hashlib.sha256((SOURCE/'Figures'/('FigureS2.'+ext)).read_bytes()).hexdigest()
script=(original/'reproducibility/build_leave_one_out_figure.py').read_text(encoding='utf8')
old="D=Path(__file__).resolve().parents[1];P=D/'provenance';F=D/'figures'"
assert script.count(old)==1
script=script.replace(old,'D=Path('+repr(original.as_posix())+');P=Path('+repr(B.as_posix())+');F=Path('+repr((OUT/'Figures').as_posix())+')')
line="fig.text(.035,.975,'Figure S2',fontsize=14,fontweight='bold',va='top')"
assert script.count(line)==1
script=script.replace(line,'# Figure numbering is given in the filename and legend, consistently with the other figures.')
script='import sys\nsys.path.insert(0, '+repr((ROOT/'work/pydeps').as_posix())+')\n'+script
(B/'build_figure_s2.py').write_text(script,encoding='utf8')

readme='''# 최종 표와 그림 정리본 v3

Table 3은 A BBJ GD, B UKB 갑상선기능항진증, C FinnGen Graves 안병증을 가로로 나란히 배치한 하나의 Word 표입니다. 공통 유전자 행으로 MR와 공위치화 결과를 비교하며 제목과 각주는 한 번만 표시합니다. 전체 42개 조합의 수치를 유지하고 글자를 읽을 수 있도록 A3 가로 1쪽으로 제작했습니다.

IVs는 instrumental variables(도구변수)의 복수형입니다. 이번 MR에서는 SNP를 도구변수로 사용하므로 표 머리글을 SNPs, n으로 풀어 적었습니다. 환자 수나 임상 보정변수 개수가 아닙니다.

## 구성

- 본문 Tables 1–3와 보충 Tables S1–S3의 Word/PDF
- Figures 1–3와 Figures S1–S3의 PDF/PNG
- All_Tables.pdf와 All_Figures.pdf 검토 합본
- Figure_Legends와 모든 기존 상세 데이터

그림 내부에는 Figure 번호를 넣지 않는 방식으로 통일했습니다. Figure S2에만 있던 번호를 기존 검증된 plotting source에서 제거했고, 나머지 그림은 그대로 유지했습니다. 번호는 Figure1.pdf, FigureS1.pdf 등의 파일명과 Figure 1, Figure S1 등의 범례 제목으로 표시합니다. 패널 A/B 등의 표시는 유지합니다.

통계수치는 변경하지 않았습니다. 독립 Table 4는 없습니다. 본문·발표자료·통합 Supplementary Dataset workbook을 재빌드한 전체 제출 패키지는 아니며, 그 자료의 확장 분석 범위와 표 참조 동기화는 별도입니다.
'''
(OUT/'READ_ME_KO.md').write_text(readme,encoding='utf8')
print(json.dumps({'out':str(OUT),'native_word_tables':len(d.tables),'data_rows':len(display),'numeric_blocks':42,'page_size':'A3 landscape','figure_number_style':'filename and legend only'},ensure_ascii=False))

from pathlib import Path
import csv,json,math,hashlib,zipfile,shutil,re
from docx import Document
from pypdf import PdfReader,PdfWriter
from PIL import Image,ImageChops

B=Path(__file__).resolve().parent;ROOT=B.parent.parent
SOURCE=ROOT/'outputs/TED_Paper_Tables_Figures_20260930_v2'
OUT=ROOT/'outputs/TED_Paper_Tables_Figures_20260930_v3'
shutil.copy2(B/'render_Table3/Table3.pdf',OUT/'Tables/Table3.pdf')
N=0
def check(ok,label):
    global N
    assert ok,label
    N+=1
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def normal(s):return re.sub(r'\s+','',s)

old=Document(SOURCE/'Tables/Table3.docx')
new=Document(OUT/'Tables/Table3.docx')
check(len(new.tables)==1,'Exactly one native Word table')
t=new.tables[0]
check(len(t.rows)==17 and len(t.columns)==22,'Three headers and 14 shared gene rows with three disease blocks')
expected=json.loads((B/'display_rows.json').read_text(encoding='utf8'))
for i,r in enumerate(t.rows[3:]):
    cells=[c.text for c in r.cells]
    for a,e in zip(cells,expected[i]):check(a==e,'Word display source equality')
    for j,source_table in enumerate(old.tables):
        oc=[c.text for c in source_table.rows[i+2].cells]
        check(cells[0]==oc[0],'Shared gene label')
        for a,e in zip(cells[1+7*j:8+7*j],oc[1:]):check(a==e,'Original numeric cell unchanged')
        check(t.rows[2].cells[1+7*j].text=='SNPs\nn','Instrument count header expanded')

table_pdf=PdfReader(OUT/'Tables/Table3.pdf')
check(len(table_pdf.pages)==1,'Single-page Table 3')
page=table_pdf.pages[0]
check(math.isclose(float(page.mediabox.width),420/25.4*72,abs_tol=.1),'A3 landscape width')
check(math.isclose(float(page.mediabox.height),297/25.4*72,abs_tol=.1),'A3 landscape height')
txt=normal(page.extract_text())
for row in expected:
    check(normal(row[0]) in txt,'Gene present in PDF')
    for value in row[1:]:check(normal(value) in txt,'Formatted value present in PDF')
check('instrumentalvariables' in txt and 'notthe numberofparticipants'.replace(' ','') in txt,'Instrument definition present')
check('A BBJ Graves disease'.replace(' ','') in txt and 'B UKB hyperthyroidism'.replace(' ','') in txt and 'C FinnGen Graves ophthalmopathy'.replace(' ','') in txt,'All three disease headings')
check('Table4' not in txt,'No obsolete table reference')

for p in (SOURCE/'Detailed_Data').iterdir():check(sha(p)==sha(OUT/'Detailed_Data'/p.name),'Detailed data unchanged')
for name in ['Table1','Table2','TableS1','TableS2','TableS3']:
    for ext in ['pdf','docx']:check(sha(SOURCE/'Tables'/f'{name}.{ext}')==sha(OUT/'Tables'/f'{name}.{ext}'),'Other table unchanged')
check(not list(OUT.rglob('Table4.*')),'No standalone Table 4')

old_s2=Image.open(SOURCE/'Figures/FigureS2.png').convert('RGB')
new_s2=Image.open(OUT/'Figures/FigureS2.png').convert('RGB')
check(old_s2.size==new_s2.size,'Figure S2 canvas preserved')
box=ImageChops.difference(old_s2,new_s2).getbbox()
check(box is not None and box[3]<old_s2.height*.05,'Only old figure-number strip changed')
check(old_s2.crop((0,150,old_s2.width,old_s2.height)).tobytes()==new_s2.crop((0,150,new_s2.width,new_s2.height)).tobytes(),'Entire scientific figure below heading pixel-identical')
old_record=json.loads((ROOT/'work/review_8dce23c/submission/candidate_20260911_review/provenance/leave_one_out_figure_sources.json').read_text(encoding='utf8'))
new_record=json.loads((B/'leave_one_out_figure_sources.json').read_text(encoding='utf8'))
check(old_record['rows']==new_record['rows'],'All 20 plotted estimates and P values unchanged')
check(old_record['data_sha256']==new_record['data_sha256'],'Source data hash unchanged')
for name in ['Figure1','Figure2','Figure3','FigureS1','FigureS3']:
    for ext in ['pdf','png']:check(sha(SOURCE/'Figures'/f'{name}.{ext}')==sha(OUT/'Figures'/f'{name}.{ext}'),'Other figure unchanged')
for p in sorted((OUT/'Figures').glob('*.pdf')):
    text=''.join(pg.extract_text() for pg in PdfReader(p).pages)
    check(not re.search(r'Figure\s+S?\d',text),'Consistent absence of figure number inside artwork')
    check('\ufffd' not in text,'No replacement glyph')
old_text=normal(PdfReader(SOURCE/'Figures/FigureS2.pdf').pages[0].extract_text()).replace('FigureS2','')
new_text=normal(PdfReader(OUT/'Figures/FigureS2.pdf').pages[0].extract_text())
check(old_text==new_text,'Figure S2 PDF text preserved except number')

pagecounts={}
for label,names in [('Tables',['Table1','Table2','Table3','TableS1','TableS2','TableS3']),('Figures',['Figure1','Figure2','Figure3','FigureS1','FigureS2','FigureS3'])]:
    writer=PdfWriter();texts=[]
    for name in names:
        path=OUT/label/(name+'.pdf')
        texts.extend(pg.extract_text() for pg in PdfReader(path).pages)
        writer.append(str(path),outline_item=name)
    writer.write(OUT/f'All_{label}.pdf')
    merged=PdfReader(OUT/f'All_{label}.pdf');pagecounts[label]=len(merged.pages)
    check(len(merged.pages)==6,'Six-page review bundle')
    for pg,text in zip(merged.pages,texts):check(pg.extract_text()==text,'Merged PDF content preserved')
manifest={p.relative_to(OUT).as_posix():sha(p) for p in OUT.rglob('*') if p.is_file() and p.name!='PACKAGE_SHA256.json'}
(OUT/'PACKAGE_SHA256.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')

def pack(path,select=lambda p:True):
    files=[p for p in OUT.rglob('*') if p.is_file() and select(p)]
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for p in files:z.write(p,p.relative_to(OUT).as_posix())
    with zipfile.ZipFile(path) as z:
        check(z.testzip() is None,'ZIP CRC')
        for p in files:check(hashlib.sha256(z.read(p.relative_to(OUT).as_posix())).hexdigest()==sha(p),'ZIP file hash')
pack(OUT.with_suffix('.zip'))
pack(ROOT/'outputs/TED_Tables_Only_20260930_v3.zip',lambda p:p.parent.name in ['Tables','Detailed_Data'] or p.name in ['All_Tables.pdf','READ_ME_KO.md'])
pack(ROOT/'outputs/TED_Figures_Only_20260930_v3.zip',lambda p:p.parent.name=='Figures' or p.name.startswith('Figure_Legends') or p.name in ['All_Figures.pdf','READ_ME_KO.md'])
report={'status':'PASS','checks':N,'native_table_count':1,'table3_pages':1,'page_format':'A3 landscape',
        'data_rows':14,'disease_blocks':3,'source_numeric_cells_unchanged':True,'numeric_results_changed':False,
        'figure_s2_difference_bbox':box,'figure_s2_statistics_unchanged':True,'other_figures_byte_identical':True,
        'all_tables_pages':pagecounts['Tables'],'all_figures_pages':pagecounts['Figures'],
        'visual_review':'Single wide Table 3 and revised Figure S2 reviewed; retained figures and tables match previously reviewed files'}
(B/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False))

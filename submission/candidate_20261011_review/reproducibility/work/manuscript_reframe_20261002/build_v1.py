from pathlib import Path
import sys, re, json, hashlib, shutil, math
sys.stdout.reconfigure(encoding='utf-8')
R=Path(__file__).resolve().parents[2]; W=Path(__file__).resolve().parent
sys.path.insert(0,str(R/'work/pydeps'))
import pandas as pd, numpy as np, matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from content_v1 import *

O=R/'outputs/TED_Manuscript_Reframed_20261002'; O.mkdir(exist_ok=True)
for x in ['Tables','Figures','Verification']: (O/x).mkdir(exist_ok=True)
OLD=R/'work/review_8dce23c/submission/candidate_20260911_review'
master=(OLD.parents[1]/'MANUSCRIPT_TED_TRAP_v5_MASTER.md').read_text(encoding='utf-8')
F=R/'outputs/TED_Followup_Validation_20261002/Tables'; A=R/'outputs/TED_Reanalysis_20261002/Tables'
sources={}; display_audit=[]
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):
 sources[str(p.relative_to(R))]=sha(p);return pd.read_csv(p)
evidence=load(F/'Table_F1_24_genes_five_outcomes.csv')
mr=load(A/'All_2544_genes_five_outcomes.csv')
coloc=load(F/'All_24_genes_1440_coloc_settings.csv')
windows=load(F/'Table_F2_TSHR_windows_priors_frequencies.csv')
freq=load(F/'Table_F4_C4B_HLAG_frequency_MR.csv')
order=['BBJ_Graves','E4_GRAVES_STRICT','UKB_hyperthyroid','FinnGen_GO','E4_GRAVES_OPHT_STRICT']
short={'BBJ_Graves':'BBJ GD','E4_GRAVES_STRICT':'FinnGen GD','UKB_hyperthyroid':'UKB HT','FinnGen_GO':'FinnGen GO','E4_GRAVES_OPHT_STRICT':'FinnGen GO strict'}
colors=['#28658b','#43866b','#777777','#896348','#865290']
genes=['TSHR','IGF1R','CTLA4']
primary=evidence.set_index(['gene','outcome'])
hits=mr[mr.outcome.eq('BBJ_Graves')&mr.P_lt_0_05_over_2544].sort_values('pvalue').gene_symbol.tolist()
newgenes=['APOM','NEU1','IGLV2-18','IGLV3-12','TNF','HIST1H4E','LTA','LTB','C4B']
book=openpyxl.load_workbook(OLD/'Supplementary_Dataset.xlsx',read_only=True,data_only=True)
def olddf(n):
 vals=list(book[n].values);return pd.DataFrame(vals[1:],columns=vals[0])
oldloo=olddf('Leave_one_out')
newloo=load(R/'work/reanalysis_20261002/mr_v1/leave_one_out_14genes.csv')
def fmtp(x):
 if pd.isna(x):return 'NA'
 if x<.001:
  a,b=f'{x:.2e}'.split('e');return a+'×10'+str(int(b)).translate(str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹'))
 return f'{x:.3g}'
def pp(x):return '<0.001' if x<.001 else f'{x:.3f}'
def orci(row):
 if pd.isna(row.OR):return 'NA'
 digits=3 if row.get('gene',row.get('gene_symbol',''))=='MAPKAPK5' else 2
 return f'{row.OR:.{digits}f} ({row.CI_lower:.{digits}f}–{row.CI_upper:.{digits}f})'
def evrow(g,o):
 r=primary.loc[(g,o)].copy();r['gene']=g;return r
def tablemd(cols,rows):return '| '+' | '.join(cols)+' |\n|'+'|'.join(['---']*len(cols))+'|\n'+'\n'.join('| '+' | '.join(map(str,r))+' |' for r in rows)

t1cols=['Dataset / endpoint','Cases','Controls','Exposure or clinical role']
t1rows=[['eQTLGen blood expression','—','—','Exposure: up to 31,684; predominantly European'],['BBJ GD (GCST90018627)','2,809','172,656','Discovery; East Asian'],['FinnGen GD (E4_GRAVES_STRICT)','3,962','496,386','Post hoc GD comparison; Finnish'],['UKB hyperthyroidism (GCST90038636)','3,731','480,867','Self-report; ancestry not established as exclusively European'],['FinnGen GO (GRAVES_OPHT)','858','499,490','Original ophthalmopathy comparison; Finnish'],['FinnGen GO strict (E4_GRAVES_OPHT_STRICT)','753','499,595','Post hoc definition sensitivity; Finnish']]
t1note='GD, Graves disease; GO, Graves ophthalmopathy; HT, hyperthyroidism/thyrotoxicosis. The three FinnGen R12 endpoints overlap and use population non-cases for their respective definitions. Strict is the source endpoint label, not proof of perfect specificity. The descriptive orbital series (four TED patients, one control) is separate from the genetic datasets.'
t2cols=['Gene','Dataset','IVs','OR (95% CI)','P value','PP.H3','PP.H4']
t2rows=[]
for g in genes:
 for o in order:
  r=evrow(g,o);t2rows.append([g,short[o],int(r.n_iv),orci(r),fmtp(r.pvalue),pp(r['PP.H3']),pp(r['PP.H4'])])
t2note='IVs, instrumental variables (retained genetic variants). Wald ratio for one IV; random-effects IVW for multiple IVs. ORs are per approximately one SD higher reconstructed blood expression. Bold P values denote nominal P < 0.05 only. Discovery uses 0.05/2,544. PP.H3 supports distinct variants; PP.H4 supports one shared variant. Bold PP.H4 denotes ≥0.80 under p12 = 10⁻⁵. These are probabilities, not P values. CTLA4 multi-IV results are concentrated in rs13030124; TSHR is single-IV in all datasets.'
t3acols=['BBJ discovery gene','BBJ OR (95% CI)','BBJ P','FinnGen GD OR (95% CI)','FinnGen GD P','Focused support','GD PP.H4']
t3arows=[]
for g in hits:
 b=evrow(g,'BBJ_Graves'); f=evrow(g,'E4_GRAVES_STRICT');supported=pd.notna(f.pvalue) and f.pvalue<.05/13 and b.beta*f.beta>0
 t3arows.append([g,orci(b),fmtp(b.pvalue),orci(f),fmtp(f.pvalue),'Yes' if supported else ('NA' if pd.isna(f.pvalue) else 'No'),pp(f['PP.H4'])])
t3bcols=['Additional gene','GD IVs','GD OR (95% CI)','GD P','BBJ PP.H4','GD PP.H4','GO strict PP.H4']
t3brows=[]
for g in newgenes:
 f=evrow(g,'E4_GRAVES_STRICT');t3brows.append([g,int(f.n_iv),orci(f),fmtp(f.pvalue),pp(evrow(g,'BBJ_Graves')['PP.H4']),pp(f['PP.H4']),pp(evrow(g,'E4_GRAVES_OPHT_STRICT')['PP.H4'])])
t3ccols=['Gene / GD analysis','Reference OR (95% CI)','Reference P','Cohort-frequency OR (95% CI)','Cohort P','Interpretation']
t3crows=[]
for g in ['C4B','HLA-G']:
 a=freq[(freq.gene==g)&(freq.outcome=='E4_GRAVES_STRICT')&(freq.scenario=='original_reference')].iloc[0]
 b=freq[(freq.gene==g)&(freq.outcome=='E4_GRAVES_STRICT')&(freq.scenario=='reharmonized_eqtlgen')].iloc[0]
 t3crows.append([g,orci(a),fmtp(a.pvalue),orci(b),fmtp(b.pvalue),'Heterogeneous; loses screen significance' if g=='C4B' else 'Weight-sensitive; single-variant dependence'])
t3note='Panel A: focused support requires concordant direction and P < 0.05/13; 7 of 13 qualify. NA is unavailable, not a negative result. Panel B: nine additional reference-frequency FinnGen GD discoveries at P < 0.05/2,544; post hoc and not independent of the full screen. APOM/NEU1 share rs3101018; IGLV2-18/IGLV3-12 share rs5759410. The lead IGLV expression variant is absent in BBJ, weakening interpretation of its low PP.H4. Panel C: HLA-G was added because of frequency sensitivity, not primary discovery. Neither C4B nor HLA-G gains strong GD colocalization. All 24 genes and five datasets are available in the workbook.'

s1cols=['Dataset','Full cis SNPs','Primary PP.H4','Lower-prior PP.H4','PP.H4 range across windows']
s1rows=[]
alias={'BBJ':'BBJ_Graves','UKB':'UKB_hyperthyroid'}
for o in order:
 z=windows[windows.outcome.replace(alias).eq(o)&windows.scenario.eq('original_reference')&windows.p12.eq(1e-5)]
 r=evrow('TSHR',o);s1rows.append([short[o],int(r.n_overlap),pp(r['PP.H4']),pp(r.H4_low_prior),f'{z["PP.H4"].min():.3f}–{z["PP.H4"].max():.3f}'])
s1note='Table S1. TSHR regional and prior sensitivity. Lower prior: p12 = 10⁻⁶; p1 = p2 = 10⁻⁴ throughout. Window range uses the full source cis-region, ±500 kb, ±250 kb and five-outcome intersection at the primary prior. It is not a confidence interval. Full regional disease lead variants were retained. All 240 settings are in TSHR_windows.'
s2cols=['Gene / dataset','Frequency','IVs','OR (95% CI)','P value','Q-test P']
s2rows=[]
for _,r in freq.iterrows():s2rows.append([r.gene+' / '+short[r.outcome], 'Reference' if r.scenario=='original_reference' else 'eQTLGen',int(r.n_iv),orci(r),fmtp(r.pvalue),fmtp(r.Q_pvalue)])
s2note='Table S2. C4B and HLA-G frequency sensitivity. Variant sets and outcome effects were held fixed. Frequency changes rescaled expression estimates and IVW weights. Q tests heterogeneity, not directional pleiotropy; its P value is not a gene-association P value. All 72 omission estimates and instrument weights are provided in the workbook.'
oldpower=olddf('Detection_limits')
print('power scenarios',oldpower.scenario.unique())
s3cols=['Dataset','Frequency','Estimable genes','Median detectable OR (IQR)','Detect OR 1.5 (%)','Detect OR 2.0 (%)']
s3rows=[]
for o in ['BBJ_Graves','UKB_hyperthyroid','FinnGen_GO']:
 for scenario in ['original_reference','reharmonized_eqtlgen']:
  r=oldpower[(oldpower.outcome==o)&(oldpower.scenario==scenario)].iloc[0]
  s3rows.append([short[o],'Reference' if scenario=='original_reference' else 'eQTLGen',int(r.n_genes),f'{r.or_median:.2f} ({r.or_q1:.2f}–{r.or_q3:.2f})',f'{r.frac_OR1_5*100:.1f}',f'{r.frac_OR2*100:.1f}'])
s3note='Table S3. Previously completed MR detection benchmarks for the original three outcomes. Alpha = 0.05/2,544 for BBJ and 0.05 for UKB/original FinnGen GO; 80% power under a normal approximation with observed SE. Percentages refer to estimable genes in each scenario. These describe association precision, not colocalization power or the probability of passing a combined filter. New endpoints were not analyzed for this table.'

# Reorder references by first appearance in the full new main narrative.
refs=dict((int(a),b) for a,b in re.findall(r'^([0-9]+)\. (.+)$',master.split('## References')[1].split('## Figure Legends')[0],re.M))
body=INTRO+'\n'+METHODS+'\n'+RESULTS+'\n'+DISCUSSION
mapping={}
def replace_cites(m):
 nums=list(map(int,re.findall(r'\d+',m.group(1))))
 for n in nums:
  if n not in mapping:mapping[n]=len(mapping)+1
 return '['+', '.join(str(mapping[n]) for n in nums)+']'
body=re.sub(r'\[([0-9, ]+)\]',replace_cites,body)
reftext='## References\n\n'+'\n\n'.join(f'{new}. {refs[old]}' for old,new in mapping.items())
authors=master.split('## Abstract')[0].split('\n\n',2)[2]
declarations=master.split('## Declarations')[1].split('## References')[0].strip()
declarations=declarations.replace('FinnGen Release 12 Graves ophthalmopathy statistics','FinnGen Release 12 GD and ophthalmopathy statistics')
declarations+='\n\n**AI-assisted preparation.** AI-assisted tools supported code development, numerical checking and language editing. The authors are responsible for verifying the analyses, references, interpretations and final manuscript.\n'
main='# '+TITLE+'\n\n'+authors+'\n\n## Abstract\n\n'+ABSTRACT+'\n\n**Keywords:** Graves disease; thyroid eye disease; TSHR; IGF1R; Mendelian randomization; colocalization\n\n'+body+'\n## Declarations\n\n'+declarations+'\n'+reftext
wordcount=lambda x:len(re.findall(r'\S+',re.sub(r'\*|\[[\d, ]+\]|^#+.*$','',x,flags=re.M)))
counts={'abstract_words_including_labels':wordcount(ABSTRACT),'main_words_excluding_headings':wordcount(body),'references':len(mapping),'discussion_paragraphs_before_limitations':len([p for p in DISCUSSION.split('### Limitations')[0].split('\n\n') if p and not p.startswith('#')])}
print(counts)
(O/'Manuscript.md').write_text(main,encoding='utf-8')
(O/'Supplementary_Material.md').write_text(SUPPLEMENT,encoding='utf-8')
(O/'Tables.md').write_text('\n\n'.join(['# Main tables','## Table 1. Data sources',tablemd(t1cols,t1rows),t1note,'## Table 2. TSHR, IGF1R and CTLA4',tablemd(t2cols,t2rows),t2note,'## Table 3. Screen and supporting analyses','### A. BBJ discoveries and FinnGen GD comparison',tablemd(t3acols,t3arows),'### B. Additional FinnGen GD discoveries',tablemd(t3bcols,t3brows),'### C. Frequency-sensitive results',tablemd(t3ccols,t3crows),t3note]),encoding='utf-8')

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'text.color':'black','axes.labelcolor':'black','xtick.color':'black','ytick.color':'black','axes.spines.top':False,'axes.spines.right':False,'pdf.fonttype':42,'savefig.dpi':250})
def savefig(fig,name):
 for ext in ['png','pdf']:fig.savefig(O/'Figures'/f'{name}.{ext}',bbox_inches='tight',facecolor='white')
 plt.close(fig)

# Figure 1: all BBJ genes and explicit current study stages.
fig=plt.figure(figsize=(12,9)); gs=fig.add_gridspec(2,1,height_ratios=[1.05,2.6],hspace=.28)
flow=fig.add_subplot(gs[0]);flow.axis('off');flow.set_title('Figure 1. Study design and BBJ discovery',loc='left',fontweight='bold',fontsize=14,pad=12)
texts=['A  Blood expression instruments\n4,462 druggable genes considered\n2,544 instrumented genes; 6,135 records','BBJ GD discovery\n2,234 estimable genes; 13 discoveries\nThreshold P < 0.05/2,544','Disease-definition comparisons\nFinnGen GD, UKB hyperthyroidism\nTwo overlapping ophthalmopathy endpoints']
for i,txt in enumerate(texts):
 x=.01+i*.334;flow.add_patch(FancyBboxPatch((x,.33),.31,.59,boxstyle='round,pad=0.008',facecolor='#f1f4f5',edgecolor='#71838b',transform=flow.transAxes));flow.text(x+.015,.62,txt,ha='left',va='center',transform=flow.transAxes,fontsize=10,linespacing=1.7)
 if i<2:flow.annotate('',xy=(x+.331,.62),xytext=(x+.312,.62),xycoords='axes fraction',arrowprops={'arrowstyle':'->','color':'black'})
flow.text(.01,.1,'Regional follow-up: TSHR and IGF1R at the centre; CTLA4 comparator; 24 selected genes in total.',transform=flow.transAxes,fontsize=11)
ax=fig.add_subplot(gs[1]); inst=olddf('Instruments');inst['pvalue']=pd.to_numeric(inst.pvalue);inst=inst.sort_values('pvalue').drop_duplicates('gene_symbol')
disc=mr[mr.outcome=='BBJ_Graves'].merge(inst[['gene_symbol','chr','pos_hg19']],on='gene_symbol');disc=disc[disc.estimable]
disc['chr']=pd.to_numeric(disc['chr']);disc['pos_hg19']=pd.to_numeric(disc.pos_hg19)
chrommax=disc.groupby('chr').pos_hg19.max();offset={};total=0;ticks=[]
for c in range(1,23):offset[c]=total;sz=float(chrommax.get(c,1e8));ticks.append(total+sz/2);total+=sz+1e7
disc['x']=[offset[int(c)]+p for c,p in zip(disc.chr,disc.pos_hg19)];disc['y']=-np.log10(disc.pvalue)
for c,g in disc.groupby('chr'):ax.scatter(g.x,g.y,s=8,color='#bac7ce' if c%2 else '#92a6b2',alpha=.7,linewidths=0)
threshold=-np.log10(.05/2544);ax.axhline(threshold,color='#555555',ls='--',lw=1.2)
ax.text(total*.68,threshold+.45,'Bonferroni threshold: P < 1.965×10⁻⁵',fontsize=9)
offsets={'HLA-A':(0,12),'HLA-DQA2':(-28,14),'C4A':(-24,23),'TUBB':(-40,18),'PSMB8':(-35,-22),'IFNGR1':(25,25),'HSD3B7':(0,56),'VKORC1':(-20,36),'PRSS36':(22,14),'TNFSF14':(10,22),'MAPKAPK5':(-12,28),'CTLA4':(0,13),'TSHR':(0,13),'IGF1R':(0,20)}
for g in hits+['IGF1R']:
 r=disc[disc.gene_symbol==g].iloc[0]; color={'TSHR':'#146891','IGF1R':'#934912','CTLA4':'#73558a'}.get(g,'#344f5d');ax.scatter([r.x],[r.y],s=30 if g!='IGF1R' else 60,color=color,zorder=5)
 dx,dy=offsets[g];ax.annotate(g,(r.x,r.y),xytext=(dx,dy),textcoords='offset points',ha='center',fontsize=9,fontstyle='italic',color='black',arrowprops={'arrowstyle':'-','lw':.6,'color':'black'})
ax.set_xticks(ticks,[str(c) for c in range(1,23)]);ax.set_ylim(-.3,25.5);ax.set_xlim(-1e7,total);ax.set_ylabel('Association strength (−log₁₀ P)');ax.set_xlabel('Chromosome of strongest selected expression instrument (GRCh37)');ax.set_title('B  BBJ gene-level MR associations',loc='left',fontweight='bold');ax.grid(axis='y',alpha=.18)
savefig(fig,'Figure1')

fig=plt.figure(figsize=(14,9.2));gs=fig.add_gridspec(1,3,width_ratios=[3,1.7,2.7],wspace=.07);ax=fig.add_subplot(gs[0]);tx=fig.add_subplot(gs[1],sharey=ax);bar=fig.add_subplot(gs[2],sharey=ax)
fig.suptitle('Figure 2. TSHR and IGF1R across disease definitions, with CTLA4 as comparator',x=.05,ha='left',fontweight='bold',fontsize=14)
rows=[(g,o) for g in genes for o in order];labels=[]
for i,(g,o) in enumerate(rows):
 r=evrow(g,o);labels.append(g+'  |  '+short[o]);c=colors[order.index(o)]
 ax.errorbar(r.OR,i,xerr=[[r.OR-r.CI_lower],[r.CI_upper-r.OR]],fmt='o',color=c,ecolor=c,capsize=3,markersize=5)
 tx.text(.02,i,orci(r),ha='left',va='center',fontsize=9);tx.text(.73,i,fmtp(r.pvalue)+(' *' if r.pvalue<.05 else ''),ha='left',va='center',fontsize=9,fontweight='bold' if r.pvalue<.05 else 'normal')
 vals=[r['PP.H0']+r['PP.H1'],r['PP.H2'],r['PP.H3'],r['PP.H4']];left=0
 for v,col in zip(vals,['#eeeeee','#c3cbd0','#dcab80','#72a9bf']):bar.barh(i,v,left=left,height=.6,color=col,edgecolor='white',lw=.5);left+=v
 bar.text(1.025,i,pp(r['PP.H4']),va='center',fontsize=10,fontweight='bold' if r['PP.H4']>=.8 else 'normal')
for a in [ax,tx,bar]:
 a.set_ylim(14.8,-.8)
 for y in [4.5,9.5]:a.axhline(y,color='#cccccc',lw=.7)
ax.set_xscale('log');ax.set_xlim(.025,4);ax.set_xticks([.03,.1,.3,1,3],['0.03','0.1','0.3','1','3']);ax.axvline(1,color='black',ls='--',lw=.8);ax.set_yticks(range(15),labels,fontsize=10);ax.set_xlabel('OR per approximately 1 SD higher expression');ax.set_title('A  MR estimate',loc='left',fontweight='bold');tx.axis('off');tx.set_xlim(0,1.32);tx.set_title('OR (95% CI)             P value',loc='left',fontsize=10)
bar.tick_params(axis='y',left=False,labelleft=False);bar.set_xlim(0,1.18);bar.set_xticks([0,.5,1]);bar.axvline(.8,color='black',ls=':',lw=.8);bar.set_xlabel('Posterior probability');bar.set_title('B  Colocalization                         H4',loc='left',fontweight='bold')
from matplotlib.patches import Patch
fig.legend(handles=[Patch(facecolor=c,label=l) for c,l in zip(['#eeeeee','#c3cbd0','#dcab80','#72a9bf'],['H0/H1','H2: expression only','H3: distinct variants','H4: shared variant'])],loc='lower center',ncol=4,frameon=False,bbox_to_anchor=(.57,.015))
fig.subplots_adjust(left=.24,right=.95,top=.90,bottom=.12);savefig(fig,'Figure2')

fig=plt.figure(figsize=(12,11.5));gs=fig.add_gridspec(1,2,width_ratios=[1.65,1],wspace=.34);hm=fig.add_subplot(gs[0]);pr=fig.add_subplot(gs[1])
gorder=genes+[g for g in hits if g not in genes]+newgenes+['HLA-G'];assert len(gorder)==24
matrix=np.array([[evrow(g,o)['PP.H4'] for o in order] for g in gorder]);hm.imshow(matrix,cmap='Blues',vmin=0,vmax=1,aspect='auto',alpha=.48)
for i in range(24):
 for j in range(5):
  v=matrix[i,j];hm.text(j,i,pp(v),ha='center',va='center',fontsize=9,color='black')
  if v>=.8:hm.add_patch(Rectangle((j-.5,i-.5),1,1,fill=False,edgecolor='black',linewidth=1.2))
hm.set_xticks(range(5),['BBJ\nGD','FinnGen\nGD','UKB\nHT','FinnGen\nGO','FinnGen\nGO strict'],fontsize=9);hm.set_yticks(range(24),gorder,fontstyle='italic');hm.tick_params(length=0);hm.set_title('A  Shared-variant probability (PP.H4)',loc='left',fontweight='bold',pad=14)
for y in [2.5,13.5,22.5]:hm.axhline(y,color='black',lw=1)
for o,c in zip(order,colors):
 z=coloc[(coloc.gene=='TSHR')&(coloc.outcome==o)&(coloc.scenario=='original_reference')].sort_values('p12',ascending=False)
 pr.plot(range(3),z['PP.H4'],marker='o',lw=1.7,color=c,label=short[o])
pr.axhline(.8,color='black',ls='--',lw=.8);pr.set_ylim(0,1.05);pr.set_xticks(range(3),['10⁻⁵','5×10⁻⁶','10⁻⁶']);pr.set_xlabel('Shared-variant prior (p12)');pr.set_ylabel('PP.H4');pr.set_title('B  TSHR prior sensitivity',loc='left',fontweight='bold',pad=14);pr.legend(frameon=False,loc='center left',bbox_to_anchor=(0,.4),fontsize=10);pr.grid(alpha=.15)
fig.suptitle('Figure 3. Regional genetic evidence and TSHR prior sensitivity',x=.06,ha='left',fontweight='bold',fontsize=14);fig.subplots_adjust(left=.13,right=.97,top=.94,bottom=.07);savefig(fig,'Figure3')

for ext in ['png','pdf']:shutil.copy2(OLD/'figures'/f'FigureS1.{ext}',O/'Figures'/f'FigureS1.{ext}')
loo=oldloo.rename(columns={'excluded_SNP':'removed_snp'})
loo=loo[~loo.removed_snp.str.startswith('None')][['gene','outcome','removed_snp','OR','CI_lower','CI_upper','pvalue']]
extra=newloo[newloo.gene_symbol.isin(['IGF1R','CTLA4'])].rename(columns={'gene_symbol':'gene'})[loo.columns]
loo=pd.concat([loo,extra],ignore_index=True);assert len(loo)==25
fig,axes=plt.subplots(1,2,figsize=(14,10),gridspec_kw={'width_ratios':[1,1]});fig.subplots_adjust(left=.17,right=.86,wspace=.9,top=.87,bottom=.12)
for ax,g in zip(axes,['IGF1R','CTLA4']):
 z=loo[loo.gene==g].copy();z['sort']=z.outcome.map({o:i for i,o in enumerate(order)});z=z.sort_values(['sort','removed_snp']);labels=[]
 for i,(_,r) in enumerate(z.iterrows()):
  ax.errorbar(r.OR,i,xerr=[[r.OR-r.CI_lower],[r.CI_upper-r.OR]],fmt='o',color=colors[order.index(r.outcome)],capsize=3,markersize=4)
  labels.append(short[r.outcome]+'\n− '+r.removed_snp);ax.text(1.04,i,fmtp(r.pvalue),transform=ax.get_yaxis_transform(),va='center',fontsize=9,fontweight='bold' if r.pvalue<.05 else 'normal')
 ax.set_yticks(range(len(z)),labels,fontsize=8.5);ax.set_ylim(len(z)-.25,-.75);ax.set_xscale('log');ax.axvline(1,color='black',ls='--',lw=.8);ax.set_xlabel('OR after omission (95% CI)');ax.set_title(g+'                                         P value',loc='left',fontweight='bold',fontsize=11);ax.grid(axis='x',alpha=.15)
fig.suptitle('Figure S2. Sensitivity of IGF1R and CTLA4 estimates to individual instruments',x=.05,ha='left',fontweight='bold',fontsize=14);fig.text(.08,.035,'Each row omits the named SNP. Original reference frequencies; 25 correlated omission estimates.\nTSHR and BBJ CTLA4 are single-instrument analyses and cannot undergo omission.',fontsize=11);savefig(fig,'FigureS2')

legends='''# Figure legends

**Figure 1. Study design and BBJ discovery.** (A) Original blood-expression screen and subsequent disease-definition comparisons. FinnGen endpoints overlap and are not independent replications. (B) One point per estimable gene (n = 2,234), positioned by the GRCh37 coordinate of its strongest selected expression instrument. The horizontal line is P = 0.05/2,544, equivalent to −log10(P) approximately 4.707. Thirteen discoveries and IGF1R are labelled. This is a gene-level MR plot, not a SNP GWAS Manhattan plot. The threshold adjusts the original discovery screen, not every later exploratory analysis.

**Figure 2. TSHR and IGF1R across disease definitions, with CTLA4 as comparator.** (A) Wald/IVW odds ratios and 95% CIs for all 15 gene–dataset comparisons. The horizontal axis is logarithmic and the reference is OR = 1. Bold P values and an asterisk indicate nominal P < 0.05, not Bonferroni significance. ORs describe reconstructed blood expression, not a treatment intervention. (B) Aligned colocalization probabilities under p1 = p2 = 10⁻⁴ and p12 = 10⁻⁵. H0/H1 denotes neither trait or disease only; H2 expression only; H3 distinct variants; H4 a shared variant. Bold H4 numbers indicate ≥0.80. MR P values and posterior probabilities answer different questions.

**Figure 3. Regional evidence and TSHR prior sensitivity.** (A) Primary PP.H4 for 24 selected genes and five datasets. Borders mark PP.H4 ≥ 0.80, not P-value significance. All lettering is black, and probabilities <0.001 retain their full precision in the workbook. Row dividers separate the three focal/comparator genes, the remaining BBJ discoveries, nine additional FinnGen GD discoveries, and sensitivity-selected HLA-G. Only these selected genes underwent colocalization. The low BBJ IGLV probabilities are limited by absence of the lead expression variant. (B) TSHR PP.H4 at p12 = 10⁻⁵, 5×10⁻⁶ and 10⁻⁶; p1 = p2 = 10⁻⁴. The dashed line marks the reporting threshold of 0.80. FinnGen GO and GO strict are overlapping ophthalmopathy endpoints; GD is E4_GRAVES_STRICT.

**Figure S1. Descriptive orbital expression.** Previously reviewed sample-level transcript abundance for TSHR, IGF1R and CTLA4 in four TED patients and one control. Points are biological samples; horizontal lines show TED means. Fold changes use log2[(mean TED TPM + 0.01)/(control TPM + 0.01)]. Axes differ by gene. No inferential expression test was performed; these observations do not validate the genetic findings. The figure is unchanged from the earlier review package.

**Figure S2. Instrument-omission sensitivity.** Twenty-five reference-frequency SNP-omission estimates for IGF1R and CTLA4 across the five datasets. The named SNP is omitted; Wald ratios apply when one remains and IVW when at least two remain. Horizontal lines show 95% CIs; the vertical reference is OR = 1. Bold P values mark nominal P < 0.05 only. Original full-set estimates are in Table 2. TSHR and BBJ CTLA4 were ineligible because each has one instrument. These correlated post hoc diagnostics are not independent replication.
'''
(O/'Figure_Legends.md').write_text(legends,encoding='utf-8')

# A single consolidated results workbook, with exact source values and actual column inventory.
wb=openpyxl.Workbook();wb.remove(wb.active)
guide=[['Sheet','Scope / source']];dictrows=[['Sheet','Column','Meaning / scope']]; exported={}
def sheet(name,df,scope):
 ws=wb.create_sheet(name); ws.append([str(c) for c in df.columns]);guide.append([name,scope]);exported[name]=df
 for c in df.columns:
  meanings={'pvalue':'Two-sided MR P value unless this sheet explicitly describes a different test','PP.H0':'Posterior: neither trait','PP.H1':'Posterior: disease only','PP.H2':'Posterior: expression only','PP.H3':'Posterior: distinct causal variants','PP.H4':'Posterior: one shared causal variant','Q_pvalue':'Cochran Q heterogeneity P value','n_iv':'Retained genetic instrument count','beta':'Log-odds MR coefficient unless instrument-level/exposure sheet','se':'Standard error; see sheet scope','OR':'MR odds ratio','CI_lower':'Lower 95% CI','CI_upper':'Upper 95% CI','top_snp':'Highest probability variant conditional on H4, not necessarily disease lead','scenario':'Frequency, availability or harmonization scenario; see Readme','n_overlap':'Number of matched regional variants','outcome':'Disease phenotype code; see Table 1'}
  dictrows.append([name,str(c),meanings.get(c,'Source-exported field; interpret with the stated sheet scope and Supplementary Methods')])
 for row in df.itertuples(index=False,name=None):
  ws.append([None if pd.isna(x) else (x.item() if isinstance(x,np.generic) else x) for x in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(name='Arial',bold=True,color='000000');cell.fill=PatternFill('solid',fgColor='E7EDF1');cell.alignment=Alignment(wrap_text=True)
 ws.row_dimensions[1].height=32
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name='Arial',size=10,color='000000')
 for i,c in enumerate(df.columns,1):ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=min(34,max(15,len(str(c))+2))
sheet('MR_primary',mr,'Primary five-outcome export; all 2,544 eligible genes × five datasets = 12,720 rows, including unavailable estimates.')
sheet('Selected_evidence',evidence,'All 24 genes × five datasets; combines primary MR, H0–H4 and sensitivity summaries.')
sheet('Colocalization',coloc,'1,440 full-precision settings. Disease is dataset 1; expression is dataset 2.')
sheet('TSHR_windows',windows,'240 settings with source-cis, narrower and intersection windows; additional regional check.')
sheet('MR_original_freq',olddf('MR'),'Frequency scenarios for original BBJ, UKB and FinnGen GO; preserved from previous verified workbook.')
sheet('MR_new_freq',load(A/'New_MR_all_frequency_scenarios.csv'),'New GD and GO strict full screens under four frequency/harmonization scenarios.')
sheet('Instruments',olddf('Instruments'),'Selected original SNP–gene instrument records; positions GRCh37.')
sheet('MR_estimators_original',olddf('MR_estimators'),'Alternative estimators and diagnostics for original three outcomes only; not repeated for all extensions.')
sheet('LOO_original',oldloo,'Original IGF1R/CTLA4 omissions and full-set comparators; reference frequency.')
sheet('LOO_new_fixed14',newloo,'New-endpoint omissions for eligible genes in the original fixed 14-gene set.')
sheet('LOO_anchors_plot',loo,'The 25 omitted-SNP estimates displayed in Figure S2.')
sheet('Frequency_diagnostics',freq,'Eight C4B/HLA-G full models across new endpoints and two frequency sources.')
sheet('LOO_frequency',load(F/'C4B_HLAG_leave_one_out_72.csv'),'72 omitted-SNP C4B/HLA-G models under two frequency sources.')
sheet('Instrument_weights',load(F/'C4B_HLAG_SNP_weights_estimates.csv'),'C4B/HLA-G single-instrument estimates and normalized full-model IVW weights.')
sheet('Frequency_scaling',load(F/'C4B_HLAG_frequency_scaling.csv'),'C4B/HLA-G exposure frequencies, beta/SE and Z preservation diagnostics.')
sheet('Region_coverage',load(F/'All_11_genes_marginal_signal_coverage.csv'),'TSHR and ten new candidates: marginal disease/expression signal coverage.')
sheet('Detection_original',oldpower,'Earlier precision benchmarks for original three outcomes only; not formal colocalization power.')
readme=wb.create_sheet('Readme',0)
for row in [['Supplementary Dataset','Five-outcome TSHR/IGF1R reappraisal — 2026-10-02'],['Values','Frozen analysis outputs imported from the source files, not hand-entered recalculations.'],['Scenarios','original_reference; paired_reference/paired_eqtlgen for identical coloc variants; available_eqtlgen for expanded coverage; MR reharmonized_eqtlgen reruns alignment.'],['Interpretation','P is association evidence; PP.H4 is shared-variant model probability. FinnGen endpoints overlap. Unavailable values are blank, not zero.'],['Scope','No row describes outcome_min_p unless that exact column occurs in its sheet.'],['Precision','Primary MR and earlier leave-one-out exports can use 1.96 versus exact normal quantile for CI rounding. Both preserve source numerical conventions.']]+guide:readme.append(row)
readme.column_dimensions['A'].width=28;readme.column_dimensions['B'].width=110;readme.freeze_panes='A8'
dd=wb.create_sheet('Column_dictionary',1)
for row in dictrows:dd.append(row)
dd.column_dimensions['A'].width=28;dd.column_dimensions['B'].width=30;dd.column_dimensions['C'].width=95;dd.freeze_panes='A2';dd.auto_filter.ref=dd.dimensions
for ws in [readme,dd]:
 for row in ws:
  for c in row:c.font=Font(name='Arial',size=10,color='000000');c.alignment=Alignment(wrap_text=True,vertical='top')
 for c in ws[1]:c.font=Font(name='Arial',size=11,bold=True,color='000000')
wb.save(O/'Supplementary_Dataset.xlsx')

# Plain academic DOCX with deliberate table widths and repeated headers.
def document(landscape=False):
 d=Document();s=d.sections[0];setpage(s,landscape)
 for n in ['Normal','Body Text']:
  st=d.styles[n];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.25
 for n,sz in [('Title',17),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[n];st.font.name='Times New Roman';st.font.size=Pt(sz);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
 d.core_properties.title=TITLE;d.core_properties.author='Jungyul Park and colleagues';return d
def setpage(s,landscape=False):
 s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT;s.page_width=Inches(11 if landscape else 8.5);s.page_height=Inches(8.5 if landscape else 11)
 s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75);s.header_distance=s.footer_distance=Inches(.3)
def inline(p,text):
 for seg in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text):
  if not seg:continue
  r=p.add_run(seg.strip('*') if seg.startswith('*') else seg);r.bold=seg.startswith('**');r.italic=seg.startswith('*') and not seg.startswith('**')
def markdown(d,text):
 for block in text.strip().split('\n\n'):
  if not block.strip():continue
  if block.startswith('# '):p=d.add_paragraph(style='Title');inline(p,block[2:])
  elif block.startswith('### '):p=d.add_paragraph(style='Heading 2');inline(p,block[4:])
  elif block.startswith('## '):
   if block.strip() in ['## Abstract','## Introduction','## References']:d.add_page_break()
   p=d.add_paragraph(style='Heading 1');inline(p,block[3:])
  else:p=d.add_paragraph();inline(p,block.replace('  \n','\n').replace('\\*','*'))
def addtable(d,title,cols,rows,widths,note,pcols=(),h4cols=()):
 d.add_heading(title,level=1);t=d.add_table(rows=1,cols=len(cols));t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,h in enumerate(cols):t.rows[0].cells[j].text=h
 header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
 for vals in rows:
  cells=t.add_row().cells
  for j,val in enumerate(vals):cells[j].text=str(val)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);tcPr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');m.append(el)
   tcPr.append(m)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');tcPr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
    for rr in p.runs:rr.font.size=Pt(9.5);rr.bold=i==0;rr.font.color.rgb=RGBColor(0,0,0)
   if i>0:
    if j in pcols:
     txt=str(rows[i-1][j]);bold='×10' in txt or (txt!='NA' and float(txt)<.05)
     for rr in c.paragraphs[0].runs:rr.bold=bold
    if j in h4cols and not str(rows[i-1][j]).startswith('<'):
     for rr in c.paragraphs[0].runs:rr.bold=float(rows[i-1][j])>=.8
   borders=OxmlElement('w:tcBorders')
   for edge in ['bottom']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'BFC6CA');borders.append(b)
   tcPr.append(borders)
 p=d.add_paragraph(note);p.paragraph_format.space_before=Pt(8)
 for rr in p.runs:rr.font.size=Pt(9)

d=document();markdown(d,main);d.add_page_break();markdown(d,legends.split('**Figure S1.')[0]);d.save(O/'Manuscript.docx')
d=document(True)
addtable(d,'Table 1. Data sources',t1cols,t1rows,[3.6,.8,.9,4.2],t1note)
d.add_page_break();addtable(d,'Table 2. TSHR, IGF1R and CTLA4 across five disease datasets',t2cols,t2rows,[.7,1.7,.4,2,1.7,1,2],t2note,pcols=[4],h4cols=[6])
d.add_page_break();addtable(d,'Table 3. Screen and supporting analyses — A',t3acols,t3arows,[1,1.5,1.2,1.6,1.3,1.2,1.1],t3note.split('Panel B:')[0],pcols=[2,4],h4cols=[6])
d.add_page_break();addtable(d,'Table 3. Screen and supporting analyses — B',t3bcols,t3brows,[1.2,.5,1.8,1.5,1.3,1.4,1.2],t3note.split('Panel B:')[1].split('Panel C:')[0],pcols=[3],h4cols=[4,5,6])
d.add_page_break();addtable(d,'Table 3. Screen and supporting analyses — C',t3ccols,t3crows,[1.1,1.6,1.1,1.8,1.3,2],t3note.split('Panel C:')[1],pcols=[2,4]);d.save(O/'Tables.docx')
d=document();markdown(d,SUPPLEMENT)
sec=d.add_section(WD_SECTION_START.NEW_PAGE);setpage(sec,True)
addtable(d,'Table S1. TSHR regional and prior sensitivity',s1cols,s1rows,[2,1.2,1.6,1.9,2.8],s1note,h4cols=[2,3])
d.add_page_break();addtable(d,'Table S2. Frequency-sensitive exploratory candidates',s2cols,s2rows,[2.3,1.2,.5,2.1,1.7,1.7],s2note,pcols=[4])
d.add_page_break();addtable(d,'Table S3. Original association detection benchmarks',s3cols,s3rows,[1.4,1.2,1.3,2.6,1.5,1.5],s3note)
for n in ['S1','S2']:
 d.add_page_break();d.add_heading('Figure '+n,level=1)
 from PIL import Image
 with Image.open(O/'Figures'/f'Figure{n}.png') as im:w,h=im.size
 picwidth=min(9.2,5.2*w/h)
 d.add_picture(str(O/'Figures'/f'Figure{n}.png'),width=Inches(picwidth))
 caption=legends.split(f'**Figure {n}.')[1].split('\n\n')[0];markdown(d,'**Figure '+n+'.'+caption)
d.save(O/'Supplementary_Material.docx')

cover='''# Cover letter

Dear Editor,

Please consider our original research manuscript, “'''+TITLE+'''”.

The clinical question is how genetic evidence for blood TSHR and IGF1R expression relates to GD and TED-related susceptibility. We screened 2,544 potentially druggable genes, compared five disease datasets and assessed regional colocalization in 24 selected genes. TSHR showed consistent MR associations, but shared-variant evidence varied across disease definitions and prior assumptions; IGF1R evidence remained limited. The manuscript distinguishes susceptibility evidence from treatment efficacy and does not claim TED-specific effects within GD.

The work includes transparent negative and sensitivity findings, checks of regional coverage, cohort-frequency substitution and instrument omission. FinnGen endpoints are explicitly acknowledged as overlapping. The small orbital tissue series is descriptive and is not presented as independent validation. Three main tables, three main figures, a supplementary document with three tables and two figures, and one detailed results workbook accompany the manuscript.

This version is prepared for author review. Before submission, the corresponding author will confirm the target journal, all-author approval, originality and exclusive-submission declarations, current funding and conflict disclosures, and the status and extent of overlap with any related orbital-transcriptome manuscript. These declarations should not be inferred from completion of the computational analysis.

Sincerely,

Suk-Woo Yang, MD, PhD\nCorresponding author, on behalf of Jungyul Park, Min-Seon Kim, Kyung-Hwa Shin and Suk-Woo Yang\nDepartment of Ophthalmology, Seoul St. Mary's Hospital, College of Medicine, The Catholic University of Korea\nE-mail: yswoph@catholic.ac.kr
'''
(O/'Cover_Letter_Draft.md').write_text(cover,encoding='utf-8');d=document();markdown(d,cover);d.save(O/'Cover_Letter_Draft.docx')
checkcols=['STROBE-MR item(s)','Location / reporting status']
checkrows=[['1 Title and abstract','Genetic evidence, MR and colocalization design stated in Abstract; title identifies genes and diseases.'],['2–3 Background and objectives','Introduction; comparative TSHR/IGF1R susceptibility question, not treatment efficacy.'],['4 Study design and sources','Methods; Table 1; original vs post hoc analyses distinguished. No prospective registration.'],['5 Assumptions','Introduction, Methods and Limitations: relevance, pleiotropy limitations, ancestry transfer and tissue relevance.'],['6 Statistical methods','Wald/IVW, harmonization, effect reconstruction, thresholds, coloc hypotheses and priors in Methods/Supplementary Methods.'],['7 Assessment of assumptions','F statistics, Q, original alternative estimators, variant coverage; direct confounding/pleiotropy validation unavailable.'],['8 Sensitivity analyses','Frequency, windows, priors, instrument omissions. No multi-signal coloc or directionality analysis.'],['9 Software and registration','Methods; pinned R/package versions; no prospective study registration.'],['10 Descriptive data','Table 1; 2,544 eligible genes; estimable counts; case/control counts. Individual demographics not available.'],['11 Main results','Table 2 (all 15 selected-gene comparisons), Table 3A–C, Figures 1–3, full workbook.'],['12 Assumption assessment results','Instrument heterogeneity/concentration and coloc discordance described in Results; limitations retained.'],['13 Sensitivity results','Tables S1–S3; Figure S2; frequency/LOO/window sheets. Single-IV exclusions stated.'],['14 Key results','Opening Discussion and Conclusions; no expression-mediated causal proof.'],['15 Limitations','Two paragraphs; selection, overlap, ancestry, tissue, coverage, instruments, single-signal model and unperformed tests.'],['16 Interpretation','Eight Discussion paragraphs; no claim of GD/TED independence or treatment prediction.'],['17 Generalizability','Blood exposure, ancestry and phenotype limitations; no patient-level subgroup conclusions.'],['18 Funding','Declarations; retained from existing manuscript, subject to author confirmation.'],['19 Data and sharing','Supplementary Dataset and source repositories; verification records and scripts in review package.'],['20 Conflicts','Declarations; retained from existing manuscript, subject to author confirmation.']]
d=document();addtable(d,'STROBE-MR reporting crosswalk',checkcols,checkrows,[1.55,5.45],'This crosswalk documents what is and is not reported; it is not certification of methodological validity or an official journal form.');d.save(O/'STROBE_MR_Checklist.docx')

# Basic audit now; renderer and page-level review follow in a separate versioned script.
assert len(mr)==12720 and len(evidence)==120 and len(coloc)==1440 and len(windows)==240
assert len(hits)==13 and sum(r[5]=='Yes' for r in t3arows)==7
assert len(set(evidence.gene))==24 and (evidence[evidence.outcome=='E4_GRAVES_STRICT']['PP.H4']<.8).all()
assert np.max(abs(coloc[[f'PP.H{i}' for i in range(5)]].sum(axis=1)-1))<1e-8
assert counts['abstract_words_including_labels']<=250
assert counts['discussion_paragraphs_before_limitations']==8
first=[]
for group in re.findall(r'\[([0-9, ]+)\]',body):
 for n in map(int,re.findall(r'\d+',group)):
  if n not in first:first.append(n)
assert first==list(range(1,len(refs.keys()&mapping.keys())+1))
for g,o in rows:
 r=evrow(g,o);s=mr[(mr.gene_symbol==g)&(mr.outcome==o)].iloc[0]
 for c in ['OR','CI_lower','CI_upper','pvalue','beta','se','n_iv']:assert np.isclose(r[c],s[c],rtol=1e-12,atol=0),(g,o,c)
audit={'status':'NUMERICAL_AND_STRUCTURE_PASS_RENDER_PENDING','counts':counts,'source_hashes':sources,'reference_mapping':mapping,'main_table2_comparisons':15,'primary_mr_rows':len(mr),'coloc_settings':len(coloc),'TSHR_window_settings':len(windows),'figureS2_omissions':len(loo),'table3_panelA_supported':7,'figureS1_preserved_sha256':sha(O/'Figures/FigureS1.png'),'limitations':'Independent exposure replication, multi-signal colocalization and direct GD-with/without-TED analysis were not performed. No all-author approval or journal acceptance is implied.'}
(O/'Verification/build_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
for f in [Path(__file__),W/'content_v1.py']:shutil.copy2(f,O/'Verification'/f.name)
print('BUILD COMPLETE',O)

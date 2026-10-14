"""Final clarity edits and removal of an observed blank page; analysis unchanged."""
from pathlib import Path
import re,json,sys
from docx import Document
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1];O=R/'outputs/TED_Manuscript_Reframed_20261002'
changes={
 'under a lower shared-variant prior':'when p12 decreased from 10⁻⁵ to 10⁻⁶',
 'Mean orbital transcript levels exceeded the single control.':'Mean orbital TSHR and IGF1R transcript levels exceeded the single control.',
 'Population-control comparisons cannot isolate TED susceptibility within GD or predict treatment response.':'The study cannot isolate TED susceptibility within GD or predict treatment response.'
}
assumptions='MR interpretation assumes that instruments predict expression, are not associated with exposure–disease confounders, and influence disease only through the modeled exposure. A small P value or high PP.H4 cannot by itself establish these assumptions.'
ai='AI-assisted tools supported code development, numerical checking and drafting; author verification and responsibility are stated in the Declarations.'
p=O/'Manuscript.md';text=p.read_text(encoding='utf-8')
for a,b in changes.items():text=text.replace(a,b)
text=text.replace('Where applicable, sensitivity assessment',assumptions+'\n\nWhere applicable, sensitivity assessment')
text=text.replace('## Results\n',ai+'\n\n## Results\n')
p.write_text(text,encoding='utf-8')
d=Document(O/'Manuscript.docx')
def reset(p,t):
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 for seg in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',t):
  if not seg:continue
  r=p.add_run(seg.strip('*') if seg.startswith('*') else seg);r.bold=seg.startswith('**');r.italic=seg.startswith('*') and not r.bold
# Rebuild abstract paragraphs from the approved markdown, preserving emphasis.
newabs=text.split('## Abstract\n\n')[1].split('\n\n**Keywords:')[0].split('\n\n')
inabs=False;i=0
for para in d.paragraphs:
 if para.text=='Abstract':inabs=True;continue
 if inabs and para.text.startswith('Keywords:'):break
 if inabs and para.text.strip():reset(para,newabs[i]);i+=1
assert i==4
for para in d.paragraphs:
 if para.text.startswith('Where applicable, sensitivity assessment'):para.insert_paragraph_before(assumptions)
 if para.text=='Results':para.insert_paragraph_before(ai)
for para in d.paragraphs:
 if para.text=='Figure legends':
  prev=para._p.getprevious()
  if prev is not None and prev.tag==qn('w:p') and prev.findall('.//'+qn('w:br')):prev.getparent().remove(prev)
  para.paragraph_format.page_break_before=True
d.save(O/'Manuscript.docx')
a=json.loads((O/'Verification/build_audit.json').read_text(encoding='utf-8'))
count=lambda s:len(re.findall(r'\S+',re.sub(r'\*|\[[\d, ]+\]|^#+.*$','',s,flags=re.M)))
a['counts']['abstract_words_including_labels']=count('\n\n'.join(newabs));a['counts']['main_words_excluding_headings']=count(text.split('## Introduction')[1].split('## Declarations')[0]);a['editorial_revision']='v3: clearer abstract; explicit MR assumptions and Methods AI disclosure; blank-page fix'
assert a['counts']['abstract_words_including_labels']<=250
(O/'Verification/build_audit.json').write_text(json.dumps(a,ensure_ascii=False,indent=2),encoding='utf-8');print(a['counts'])

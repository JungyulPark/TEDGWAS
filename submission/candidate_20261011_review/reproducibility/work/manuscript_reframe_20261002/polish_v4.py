from pathlib import Path
import re,json,zipfile
from lxml import etree
from docx import Document
W=Path(__file__).resolve().parent;R=W.parents[1];O=R/'outputs/TED_Manuscript_Reframed_20261002'
changes={
 'when p12 decreased from 10⁻⁵ to 10⁻⁶':'when the shared-variant prior decreased from 10⁻⁵ to 10⁻⁶',
 'The exposure was whole-blood expression from eQTLGen, comprising':'The exposure used whole-blood expression quantitative trait locus (eQTL) statistics from eQTLGen, comprising',
 '6,135 SNP–gene records for 2,544 genes.':'6,135 single-nucleotide polymorphism (SNP)–gene records for 2,544 genes.',
 'Nominal BBJ and UKB associations did not survive the discovery criterion or establish a shared causal variant, and the FinnGen GD estimate was close to the null with a CI spanning both directions.':'The BBJ association did not meet the discovery criterion, and neither its nominal association nor that in UKB established a shared causal variant. The FinnGen GD estimate was close to the null, with a CI spanning both directions.'
}
p=O/'Manuscript.md';t=p.read_text(encoding='utf-8')
for a,b in changes.items():assert a in t,a;t=t.replace(a,b)
p.write_text(t,encoding='utf-8')
d=Document(O/'Manuscript.docx')
for para in d.paragraphs:
 for a,b in changes.items():
  if a in para.text:
   # These replacements fall within plain-text runs, preserving gene/P italics elsewhere.
   done=False
   for run in para.runs:
    if a in run.text:run.text=run.text.replace(a,b);done=True;break
   assert done,a
d.save(O/'Manuscript.docx')
d=Document(O/'Tables.docx')
for para in d.paragraphs:
 if para.text.startswith(('Panel A:','nine additional',' HLA-G','HLA-G was added')):
  r=para.add_run(' Bold P values denote nominal P < 0.05; screen and focused thresholds are stated separately.');r.font.size=__import__('docx').shared.Pt(9)
d.save(O/'Tables.docx')
for p in O.glob('*.docx'):
 with zipfile.ZipFile(p) as z:entries={n:z.read(n) for n in z.namelist()}
 changed=False
 for name in ['word/document.xml','word/styles.xml']:
  root=etree.fromstring(entries[name]);ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  for el in root.xpath('//w:pBdr',namespaces=ns):el.getparent().remove(el);changed=True
  entries[name]=etree.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
 if changed:
  with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:
   for n,b in entries.items():z.writestr(n,b)
a=json.loads((O/'Verification/build_audit.json').read_text(encoding='utf-8'))
count=lambda s:len(re.findall(r'\S+',re.sub(r'\*|\[[\d, ]+\]|^#+.*$','',s,flags=re.M)))
a['counts']['abstract_words_including_labels']=count(t.split('## Abstract')[1].split('**Keywords:**')[0]);a['counts']['main_words_excluding_headings']=count(t.split('## Introduction')[1].split('## Declarations')[0]);a['editorial_revision']='v4: abbreviations; BBJ-specific discovery wording; black unadorned title; nominal-P footnotes'
(O/'Verification/build_audit.json').write_text(json.dumps(a,ensure_ascii=False,indent=2),encoding='utf-8');print(a['counts'])

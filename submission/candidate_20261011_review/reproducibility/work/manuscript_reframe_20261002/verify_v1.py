from pathlib import Path
import sys,json,math,re,hashlib
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd,openpyxl
from docx import Document
W=Path(__file__).resolve().parent;R=W.parents[1];O=R/'outputs/TED_Manuscript_Reframed_20261002'
old=R/'work/review_8dce23c/submission/candidate_20260911_review/Supplementary_Dataset.xlsx'
F=R/'outputs/TED_Followup_Validation_20261002/Tables';A=R/'outputs/TED_Reanalysis_20261002/Tables'
w=openpyxl.load_workbook(O/'Supplementary_Dataset.xlsx',read_only=True,data_only=False)
b=openpyxl.load_workbook(old,read_only=True,data_only=True)
paths={'MR_primary':A/'All_2544_genes_five_outcomes.csv','Selected_evidence':F/'Table_F1_24_genes_five_outcomes.csv','Colocalization':F/'All_24_genes_1440_coloc_settings.csv','TSHR_windows':F/'Table_F2_TSHR_windows_priors_frequencies.csv','MR_new_freq':A/'New_MR_all_frequency_scenarios.csv','LOO_new_fixed14':R/'work/reanalysis_20261002/mr_v1/leave_one_out_14genes.csv','Frequency_diagnostics':F/'Table_F4_C4B_HLAG_frequency_MR.csv','LOO_frequency':F/'C4B_HLAG_leave_one_out_72.csv','Instrument_weights':F/'C4B_HLAG_SNP_weights_estimates.csv','Frequency_scaling':F/'C4B_HLAG_frequency_scaling.csv','Region_coverage':F/'All_11_genes_marginal_signal_coverage.csv'}
oldnames={'MR_original_freq':'MR','Instruments':'Instruments','MR_estimators_original':'MR_estimators','LOO_original':'Leave_one_out','Detection_original':'Detection_limits'}
count=0;report=[]
def same(a,b):
 if a is None or (not isinstance(a,str) and pd.isna(a)):return b is None
 if isinstance(a,(float,int)) and not isinstance(a,bool):return isinstance(b,(float,int)) and math.isclose(a,b,rel_tol=3e-14,abs_tol=0)
 return a==b
for name in list(paths)+list(oldnames):
 if name in paths:
  d=pd.read_csv(paths[name]); expected=[list(d.columns)]+list(d.itertuples(index=False,name=None))
 else:expected=list(b[oldnames[name]].values)
 rows=w[name].iter_rows(values_only=True);seen=0
 for i,expectedrow in enumerate(expected):
  actual=next(rows);assert len(actual)==len(expectedrow),(name,i)
  for j,(a,z) in enumerate(zip(expectedrow,actual)):
   assert same(a,z),(name,i,j,a,z);count+=1
  seen+=1
 assert next(rows,None) is None,name
 report.append({'sheet':name,'data_rows':seen-1,'source_columns':len(expected[0])})
 print(name,'PASS',seen-1,flush=True)
for ws in w:
 for row in ws:
  for c in row:
   assert c.data_type not in ['f','e'],(ws.title,c.coordinate,c.value)
dictionary=list(w['Column_dictionary'].values)[1:];actualcols={(s.title,str(c)) for s in w if s.title not in ['Readme','Column_dictionary'] for c in next(s.values)}
assert {(s,c) for s,c,_ in dictionary}==actualcols
mr=pd.read_csv(paths['MR_primary']);e=pd.read_csv(paths['Selected_evidence']);c=pd.read_csv(paths['Colocalization'])
assert e.groupby('outcome').size().eq(24).all()
assert c.groupby(['gene','outcome']).size().eq(12).all()
assert abs(c[[f'PP.H{i}' for i in range(5)]].sum(axis=1)-1).max()<1e-8
assert sum((e.outcome=='E4_GRAVES_STRICT')&(e['PP.H4']>=.8))==0
doc=Document(O/'Tables.docx');tab=doc.tables[1]
order=['BBJ_Graves','E4_GRAVES_STRICT','UKB_hyperthyroid','FinnGen_GO','E4_GRAVES_OPHT_STRICT']
sup=str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹')
def pv(x):
 if x<.001:
  a,b=f'{x:.2e}'.split('e');return a+'×10'+str(int(b)).translate(sup)
 return f'{x:.3g}'
for i,(g,o) in enumerate([(g,o) for g in ['TSHR','IGF1R','CTLA4'] for o in order],1):
 r=e[(e.gene==g)&(e.outcome==o)].iloc[0];cells=[z.text for z in tab.rows[i].cells]
 assert cells[0]==g and cells[2]==str(int(r.n_iv))
 assert cells[3]==f'{r.OR:.2f} ({r.CI_lower:.2f}–{r.CI_upper:.2f})' and cells[4]==pv(r.pvalue)
 for j,col in [(5,'PP.H3'),(6,'PP.H4')]:assert cells[j]==f'{r[col]:.3f}'
text=(O/'Manuscript.md').read_text(encoding='utf-8');bef=text.split('## References')[0];used=[]
for group in re.findall(r'\[([0-9, ]+)\]',bef):
 for n in map(int,re.findall(r'\d+',group)):
  if n not in used:used.append(n)
assert used==list(range(1,29))
assert 'nine selected genes' not in text and 'Both UKB and FinnGen are broadly European' not in text
result={'status':'PASS','workbook_source_cells_compared':count,'source_sheets_checked':report,'formula_cells':0,'error_cells':0,'dictionary_matches_actual_columns':True,'table2_source_verified_rows':15,'reference_first_appearance':used,'clinical_scope':'TSHR/IGF1R question retained; no direct TED-within-GD or treatment claim.'}
(O/'Verification/final_numeric_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print('ALL CHECKS PASS',count,flush=True)

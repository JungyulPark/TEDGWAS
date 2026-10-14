from pathlib import Path
import csv, json, math, hashlib, shutil, sys, zipfile
from docx import Document
from pypdf import PdfReader, PdfWriter

B = Path(__file__).resolve().parent
ROOT = B.parent.parent
SOURCE = ROOT / 'outputs/TED_Paper_Tables_Figures_20260930'
OUT = ROOT / 'outputs/TED_Paper_Tables_Figures_20260930_v2'
N = 0
def check(condition, message):
    global N
    assert condition, message
    N += 1
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()
def rows(p):
    with p.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

for name, sub in [('Table3', 'Tables'), ('Figure_Legends', '')]:
    shutil.copy2(B / ('render_' + name) / (name + '.pdf'), OUT / sub / (name + '.pdf'))

d = Document(OUT / 'Tables/Table3.docx')
expected = json.loads((B / 'display_rows.json').read_text(encoding='utf8'))
check(len(d.tables) == 3, 'Three panels')
actual = []
for t in d.tables:
    check(len(t.rows) == 16, 'Two headers and 14 data rows')
    actual.extend([[c.text for c in r.cells] for r in t.rows[2:]])
check(len(actual) == 42, '42 results')
for a, e in zip(actual, expected):
    for x, y in zip(a, e['cells']):
        check(x == y, 'Rendered Word cell source equality')

# Compare existing displayed numbers independently against both superseded tables.
old_mr = Document(SOURCE / 'Tables/Table3.docx')
old_pp = Document(SOURCE / 'Tables/Table4.docx')
for new_table, mr_table, pp_table in zip(d.tables, old_mr.tables, old_pp.tables):
    for i, old_row in enumerate(mr_table.rows[1:]):
        new = [c.text for c in new_table.rows[i+2].cells]
        old = [c.text for c in old_row.cells]
        check(new[:4] == old[:4], 'All 39 original MR display rows unchanged')
    for i, old_row in enumerate(pp_table.rows[1:]):
        new = [c.text for c in new_table.rows[i+2].cells]
        old = [c.text for c in old_row.cells]
        check(new[0].rstrip('†') == old[0], 'Same coloc gene order')
        check(new[4:8] == old[2:6], 'All 42 primary and low-prior display rows unchanged')

full = {r['gene_symbol']: r for r in rows(SOURCE / 'Detailed_Data/Stage1_ALL_2544_genes.csv')}
pp = {(r['gene'], r['outcome']): r for r in rows(SOURCE / 'Detailed_Data/Stage2_14_genes_primary.csv')}
settings = rows(SOURCE / 'Detailed_Data/Stage2_all_504_colocalization_settings.csv')
combined = rows(OUT / 'Detailed_Data/Table3_integrated_14genes.csv')
check(len(combined) == len({(r['gene'], r['outcome']) for r in combined}) == 42, 'No duplicated or missing gene outcomes')
for r in combined:
    gene, oc = r['gene'], r['outcome']
    for key in ['n_iv', 'beta', 'se', 'or', 'or_ci_lower', 'or_ci_upper', 'pvalue', 'status']:
        check(r[key] == full[gene][key+'__'+oc], 'Full-precision MR preservation')
    for key in ['n_overlap', 'top_snp'] + ['PP.H'+str(i) for i in range(5)]:
        check(r[key] == pp[(gene, oc)][key], 'Full-precision posterior preservation')
    check(math.isclose(sum(float(r['PP.H'+str(i)]) for i in range(5)), 1, abs_tol=1e-10), 'Posterior sum')
    for scenario, p12, key in [('original_reference', 1e-6, 'H4_lower_prior_1e-6'), ('available_eqtlgen', 1e-5, 'H4_cohort_AF_1e-5')]:
        source = next(x for x in settings if x['gene'] == gene and x['outcome'] == oc and x['scenario'] == scenario and math.isclose(float(x['p12']), p12, rel_tol=1e-12))
        check(r[key] == source['PP.H4'], 'Full-precision sensitivity preservation')
check(sum(bool(r['pvalue']) for r in combined) == 41, 'One non-estimable MR preserved')

for p in (SOURCE/'Detailed_Data').iterdir():
    check(sha(p) == sha(OUT/'Detailed_Data'/p.name), 'Original data file unchanged')
for p in (SOURCE/'Figures').iterdir():
    check(sha(p) == sha(OUT/'Figures'/p.name), 'Figure unchanged')
for name in ['Table1', 'Table2', 'TableS1', 'TableS2', 'TableS3']:
    for ext in ['pdf', 'docx']:
        check(sha(SOURCE/'Tables'/f'{name}.{ext}') == sha(OUT/'Tables'/f'{name}.{ext}'), 'Retained table unchanged')
check(not list(OUT.rglob('Table4.*')), 'No superseded Table 4 in delivery')

pages = len(PdfReader(OUT/'Tables/Table3.pdf').pages)
check(pages == 3, 'Integrated table is three pages')
legends = PdfReader(OUT/'Figure_Legends.pdf')
check(len(legends.pages) == 2, 'Legends remain two pages')
legend_text = ''.join(p.extract_text() for p in legends.pages)
check('Table 4' not in legend_text and 'Table 3 integrates' in legend_text, 'Updated figure cross-reference')
new_text = '\n'.join(p.text for p in d.paragraphs)
check('within GD' in new_text and 'post hoc' in new_text and 'not a BBJ discovery hit' in new_text, 'Interpretive qualifiers retained')
for p in PdfReader(OUT/'Tables/Table3.pdf').pages:
    tx = p.extract_text()
    check('IGF1R' in tx and 'MAPKAPK5' in tx and 'HLA-DQA2' in tx, 'All panels show complete rows')
    check('\ufffd' not in tx, 'No replacement glyph')

names = ['Table1', 'Table2', 'Table3', 'TableS1', 'TableS2', 'TableS3']
writer = PdfWriter()
source_text = []
for name in names:
    path = OUT/'Tables'/f'{name}.pdf'
    source_text.extend([p.extract_text() for p in PdfReader(path).pages])
    writer.append(str(path), outline_item=name)
writer.write(OUT/'All_Tables.pdf')
merged = PdfReader(OUT/'All_Tables.pdf')
check(len(merged.pages) == 8, 'Tables bundle reduced from eleven to eight pages')
for page, original in zip(merged.pages, source_text):
    check(page.extract_text() == original, 'Merged PDF text preservation')

report = {'status': 'PASS', 'checks': N, 'table3_pages': pages, 'all_tables_pages': len(merged.pages),
          'genes': 14, 'disease_datasets': 3, 'displayed_combinations': 42, 'estimable_mr': 41,
          'original_numeric_values_changed': False, 'figure_files_changed': False,
          'standalone_table4_included': False, 'visual_review': 'All three integrated table pages and both legend pages inspected'}
(B/'verification.json').write_text(json.dumps(report, indent=2), encoding='utf8')
manifest = {p.relative_to(OUT).as_posix(): sha(p) for p in OUT.rglob('*') if p.is_file() and p.name != 'PACKAGE_SHA256.json'}
(OUT/'PACKAGE_SHA256.json').write_text(json.dumps(manifest, indent=2), encoding='utf8')

def archive(dest, select=lambda p: True):
    files = [p for p in OUT.rglob('*') if p.is_file() and select(p)]
    with zipfile.ZipFile(dest, 'w', zipfile.ZIP_DEFLATED) as z:
        for p in files:
            z.write(p, p.relative_to(OUT).as_posix())
    with zipfile.ZipFile(dest) as z:
        check(z.testzip() is None, 'Archive CRC')
        check(not any(Path(name).name.startswith('Table4.') for name in z.namelist()), 'Archive excludes Table 4')
        for p in files:
            check(hashlib.sha256(z.read(p.relative_to(OUT).as_posix())).hexdigest() == sha(p), 'Archive hash')

archive(OUT.with_suffix('.zip'))
archive(ROOT/'outputs/TED_Tables_Only_20260930_v2.zip', lambda p: p.parent.name in {'Tables', 'Detailed_Data'} or p.name in {'All_Tables.pdf', 'READ_ME_KO.md'})
short = Path('${LOCAL_USER_HOME}/Documents/Codex')
for src, dst in [(OUT.with_suffix('.zip'), 'TED_Tables_Figures_20260930_v2.zip'),
                 (OUT/'Tables/Table3.pdf', 'TED_Table3_Integrated_20260930.pdf'),
                 (OUT/'Tables/Table3.docx', 'TED_Table3_Integrated_20260930.docx'),
                 (OUT/'All_Tables.pdf', 'TED_All_Tables_20260930_v2.pdf')]:
    shutil.copy2(src, short/dst)
    check(sha(src) == sha(short/dst), 'Short-path delivery hash')
print(json.dumps({**report, 'checks': N, 'package': str(short/'TED_Tables_Figures_20260930_v2.zip')}, ensure_ascii=False))

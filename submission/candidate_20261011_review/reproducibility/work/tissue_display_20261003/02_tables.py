from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd,math,json,hashlib,shutil,sys
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1];A=R/'outputs/TED_Tissue_Extension_20261003';O=R/'outputs/TED_Tissue_Tables_Figures_20261003';O.mkdir(exist_ok=True)
av=pd.read_csv(A/'Tables/Tissue_availability_10_combinations.csv');d=pd.read_csv(A/'Tables/Tissue_MR_all_estimable.csv')
doc=Document();sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
sec.top_margin=sec.bottom_margin=Inches(.53);sec.left_margin=sec.right_margin=Inches(.70)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(10)
 st.paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].font.bold=True
doc.core_properties.title='TSHR and IGF1R tissue results';doc.core_properties.author='';doc.core_properties.subject='Fixed tissue panel with full result reporting'
tn={'Thyroid':'Thyroid','Cells_Cultured_fibroblasts':'Cultured fibroblasts','Adipose_Subcutaneous':'Subcutaneous adipose','Adipose_Visceral_Omentum':'Visceral adipose','Whole_Blood':'Whole blood'}
on={'E4_GRAVES_STRICT':'FinnGen GD','BBJ_Graves':'BBJ GD','UKB_hyperthyroid':'UKB HT','FinnGen_GO':'FinnGen GO','E4_GRAVES_OPHT_STRICT':'FinnGen GO strict'}
sup=str.maketrans('0123456789-','⁰¹²³⁴⁵⁶⁷⁸⁹⁻')
def fp(x):
 if pd.isna(x):return '—'
 if x<.001:
  m,e=f'{x:.2e}'.split('e');return m+'×10'+str(int(e)).translate(sup)
 return f'{x:.4f}' if x<.1 else f'{x:.3f}'
def note(t):
 p=doc.add_paragraph(t);p.paragraph_format.space_after=Pt(4)
 for r in p.runs:r.font.size=Pt(9)
 return p
def table(headers,widths,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.LEFT
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,(h,w) in enumerate(zip(headers,widths)):
  c=t.rows[0].cells[i];c.width=Inches(w);c.text=h
  s=OxmlElement('w:shd');s.set(qn('w:fill'),'E8EEF4');c._tc.get_or_add_tcPr().append(s)
  for r in c.paragraphs[0].runs:r.bold=True
 hp=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hp)
 for row in rows:
  cells=t.add_row().cells
  for c,w,v in zip(cells,widths,row):c.width=Inches(w);c.text=str(v)
 for row in t.rows:
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for n in ['top','bottom','left','right']:
    e=OxmlElement('w:'+n);e.set(qn('w:w'),'48' if n in ['top','bottom'] else '70');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.size=Pt(9.5);r.font.color.rgb=RGBColor(0,0,0)
 return t
doc.add_paragraph('TSHR and IGF1R tissue results','Title')
doc.add_paragraph('Fixed panel of five GTEx v8 tissues and five disease datasets')
doc.add_paragraph('A  Instrument availability','Heading 1')
note('Four of ten gene–tissue combinations yielded one qualifying instrument each. The remaining six combinations were unestimable; they were retained in the planned testing denominators.')
rows=[]
for _,a in av.iterrows():
 hit=d[(d.gene==a.gene)&(d.tissue==a.tissue)];f='—' if hit.empty else f'{hit.iloc[0].F_min:.1f}'
 status='One SNP retained' if len(hit) else 'No eQTL at P < 5×10⁻⁸' if a.tested else 'Gene absent from tested-gene file'
 rows.append([a.gene,tn[a.tissue],int(a.n_exposure),int(a.n_selected_instruments),f,fp(a.minimum_eQTL_P),status])
t=table(['Gene','Tissue','N','IVs','F','Minimum eQTL P','Availability'],[.8,2.05,.6,.5,.6,1.3,3.3],rows)
for row in t.rows[1:]:
 for r in row.cells[0].paragraphs[0].runs:r.italic=True
doc.add_paragraph()
note('N is the GTEx tissue eQTL sample size. IVs are retained instrumental genetic variants. Eligibility required cis ±1 Mb, eQTL P < 5×10⁻⁸, F > 10, MAF ≥ 0.01, and EUR-reference LD clumping at r² < 0.001 within 10 Mb. No threshold was relaxed.')
note('The thyroid TSHR minimum eQTL P was 9.60474×10⁻⁸ and did not qualify. The cultured-fibroblast TSHR gene was absent from the tested-gene file. Neither situation is a negative MR result. Cultured fibroblasts and adipose tissues are not orbital tissue. GTEx whole blood (N = 670) is distinct from the original eQTLGen exposure (up to 31,684 participants); lack of a GTEx instrument does not invalidate the original blood analysis.')
note('Planned tests: ten primary comparisons against FinnGen GD, and forty secondary comparisons against the other four datasets. Twenty estimates were available; all thirty unestimable comparisons remain in Supplementary Dataset, Tissue_extension.')
doc.add_page_break();doc.add_paragraph('B  All estimable disease associations','Heading 1')
note('All estimates use one SNP and the Wald ratio. FinnGen GD is the primary comparison within this post hoc tissue extension, not a new primary endpoint of the original study.')
order=[('TSHR','Adipose_Subcutaneous'),('TSHR','Adipose_Visceral_Omentum'),('IGF1R','Thyroid'),('IGF1R','Cells_Cultured_fibroblasts')]
oc=list(on);rr=[];ordered=[]
for g,ti in order:
 for ocv in oc:
  v=d[(d.gene==g)&(d.tissue==ti)&(d.outcome==ocv)].iloc[0];ordered.append(v)
  label='Corrected support*' if v.statistical_category=='family_corrected' else 'Nominal only' if v.statistical_category=='nominal_only' else 'Not significant'
  rr.append([g+'\n'+tn[ti],on[ocv],f'{v.OR:.2f} ({v.CI_lower:.2f}–{v.CI_upper:.2f})',fp(v.P),fp(v.P_adjusted),label])
t=table(['Gene and tissue','Disease dataset','OR (95% CI)','P value','Adjusted P','Statistical support'],[2.15,1.45,1.55,1.15,1.15,2.1],rr)
for j in range(4):
 c=t.cell(j*5+1,0).merge(t.cell(j*5+5,0));c.text=order[j][0]+'\n'+tn[order[j][1]];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for i,p in enumerate(c.paragraphs):
  p.paragraph_format.space_after=Pt(0)
  for r in p.runs:r.font.size=Pt(10)
for i,v in enumerate(ordered,1):
 if v.family=='primary':
  for r in t.cell(i,1).paragraphs[0].runs:r.bold=True
 if v.statistical_category=='family_corrected':
  for j in [3,4]:
   for r in t.cell(i,j).paragraphs[0].runs:r.bold=True
doc.add_paragraph().paragraph_format.space_after=Pt(0)
note('* Bold P values indicate family-corrected support. Primary FinnGen GD: raw P < 0.005, adjusted P = min(1, 10P). Secondary datasets: raw P < 0.00125, adjusted P = min(1, 40P). Nominal only means P < 0.05 without passing the relevant corrected threshold. The correction denominators retain unestimable comparisons.')
note('ORs are per unit of normalized, inverse-normal-transformed expression. Effect magnitudes are not directly comparable to reconstructed eQTLGen units. GD, Graves disease; HT, hyperthyroidism; GO, Graves ophthalmopathy. FinnGen endpoints overlap and use population controls. No tissue-specific colocalization was performed. Non-significance does not establish absence of an effect.')
# Remove paragraph-border residue in the bundled default template.
for root in [doc._element,doc.styles.element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(O/'Tissue_result_tables.docx')
pd.DataFrame(rr,columns=['Gene and tissue','Disease dataset','OR (95% CI)','P value','Adjusted P','Statistical support']).to_csv(O/'Tissue_results_display.csv',index=False,encoding='utf-8-sig')
for name in ['Tissue_MR_50_planned_combinations.csv','Tissue_availability_10_combinations.csv','Selected_instruments_verified.csv']:
 q=O/'Source_data'/name;q.parent.mkdir(exist_ok=True);shutil.copy2(A/'Tables'/name,q)
for ext in ['pdf','png']:shutil.copy2(A/'Figures'/('Tissue_extension_forest.'+ext),O/'Figures'/('Tissue_all_estimates.'+ext))
shutil.copy2(R/'outputs/TED_Tissue_Integrated_20261003/TableS3_Tissue_extension.pdf',O/'Table_S3_current.pdf')
print('Created two-page result table document and copied complete source records')

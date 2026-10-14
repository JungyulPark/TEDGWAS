from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from docx import Document
from pypdf import PdfReader,PdfWriter
import pandas as pd,math,re,json,hashlib,shutil,sys
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1];O=R/'outputs/TED_Tissue_Tables_Figures_20261003';A=R/'outputs/TED_Tissue_Extension_20261003';old=R/'outputs/TED_Tissue_Integrated_20261003'
doc=Document(O/'Tissue_result_tables.docx');d=pd.read_csv(A/'Tables/Tissue_MR_all_estimable.csv');av=pd.read_csv(A/'Tables/Tissue_availability_10_combinations.csv');checks=[]
def check(n,v):checks.append(n);assert v,n
def pv(t):
 if '×10' in t:
  a,b=t.split('×10');b=b.translate(str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻','0123456789-'));return float(a)*10**int(b)
 return float(t)
check('Two table panels',len(doc.tables)==2)
check('Availability ten rows',len(doc.tables[0].rows)==11)
for row,(_,a) in zip(doc.tables[0].rows[1:],av.iterrows()):
 c=[c.text for c in row.cells];check('Availability '+a.gene+a.tissue,c[0]==a.gene and int(c[2])==a.n_exposure and int(c[3])==a.n_selected_instruments)
 if a.tested:check('eQTL P '+a.gene+a.tissue,math.isclose(pv(c[5]),a.minimum_eQTL_P,rel_tol=.006))
check('All twenty rows',len(doc.tables[1].rows)==21)
order=[('TSHR','Adipose_Subcutaneous'),('TSHR','Adipose_Visceral_Omentum'),('IGF1R','Thyroid'),('IGF1R','Cells_Cultured_fibroblasts')]
ocs=['E4_GRAVES_STRICT','BBJ_Graves','UKB_hyperthyroid','FinnGen_GO','E4_GRAVES_OPHT_STRICT'];ridx=1
for g,ti in order:
 for oc in ocs:
  v=d[(d.gene==g)&(d.tissue==ti)&(d.outcome==oc)].iloc[0];row=doc.tables[1].rows[ridx];c=[c.text for c in row.cells]
  nums=list(map(float,re.findall(r'\d+\.\d+',c[2])))
  check('OR CI row '+str(ridx),nums==[round(v.OR,2),round(v.CI_lower,2),round(v.CI_upper,2)])
  for j,val in [(3,v.P),(4,v.P_adjusted)]:
   check('P row '+str(ridx)+' col '+str(j),math.isclose(pv(c[j]),val,rel_tol=.006,abs_tol=.0005 if val>.1 else .00005 if val>.001 else 0))
   check('P emphasis row '+str(ridx)+' col '+str(j),all(bool(r.bold)==(v.statistical_category=='family_corrected') for r in row.cells[j].paragraphs[0].runs))
  check('Classification row '+str(ridx),('Corrected support' in c[5])==(v.statistical_category=='family_corrected'))
  ridx+=1
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
for p in (O/'Source_data').iterdir():check('Immutable copied data '+p.name,sha(p)==sha(A/'Tables'/p.name))
previous=json.loads((old/'MANIFEST_SHA256.json').read_text(encoding='utf-8'))['files']
check('Previous integrated package unchanged',all(sha(old/n)==meta['sha256'] for n,meta in previous.items()))
pdf=W/'render_release/Tissue_result_tables.pdf';check('Two rendered table pages',len(PdfReader(pdf).pages)==2);shutil.copy2(pdf,O/'Tissue_result_tables.pdf')
writer=PdfWriter();start=0
for label,p in [('Tissue result tables',O/'Tissue_result_tables.pdf'),('Availability and primary comparison',O/'Figures/Tissue_primary_summary.pdf'),('All twenty estimates',O/'Figures/Tissue_all_estimates.pdf')]:
 reader=PdfReader(p)
 for page in reader.pages:writer.add_page(page)
 writer.add_outline_item(label,start);start+=len(reader.pages)
writer.add_metadata({'/Title':'TSHR and IGF1R tissue results tables and figures'});writer.write(O/'Tissue_results_review.pdf')
check('Four-page review',start==4)
ver=O/'Verification';ver.mkdir(exist_ok=True)
report={'status':'PASS','checks':len(checks),'all_20_estimates_reconciled':True,'all_10_availability_rows_reconciled':True,'source_tables_unchanged':True,'previous_manuscript_package_unchanged':True,'table_pages':2,'review_pages':4,'validation_scope':'Reporting and formatting only; no new MR or colocalization was run.'}
(ver/'checks.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
for p in W.glob('*'):
 if p.suffix in ['.py','.R']:shutil.copy2(p,ver/p.name)
(O/'README.md').write_text('''# Tissue extension display package

Read PLACEMENT_KO.md for the recommended manuscript placement. Tissue_results_review.pdf combines two table pages, the new primary-summary figure and the full twenty-estimate forest plot. Tissue_result_tables.docx is editable. Table_S3_current.pdf is the compact supplementary table already incorporated in the manuscript.

The new tables and figures are review and presentation materials. No new main-table, main-figure or supplementary numbering has been assigned. The existing manuscript package was preserved unchanged. Full planned cells are copied into Source_data and remain in the single official Supplementary Dataset workbook.

Filled forest squares and emphasized P values indicate family-corrected support. Availability squares indicate a qualifying instrument, not disease significance. GTEx whole blood differs from the original eQTLGen exposure. The report explains the single-instrument and non-orbital-tissue limitations.
''',encoding='utf-8')
print(json.dumps(report,indent=2))

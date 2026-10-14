from pathlib import Path
import os,sys,runpy,re
from docx import Document
from docx.shared import RGBColor
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1];O=R/'outputs/TED_Tissue_Integrated_20261003'
# Repair the two edited abstract labels using explicit bold/italic tokens.
doc=Document(O/'Manuscript.docx');md=(O/'Manuscript.md').read_text(encoding='utf-8')
for t in md.split('## Abstract')[1].split('**Keywords:**')[0].strip().split('\n\n'):
 if not t.startswith(('**Methods:**','**Results:**')):continue
 p=next(p for p in doc.paragraphs if p.text.replace('*','')==t.replace('*',''))
 p.clear()
 for s in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',t):
  if not s:continue
  bold=s.startswith('**') and s.endswith('**');ital=s.startswith('*') and s.endswith('*') and not bold
  r=p.add_run(s[2:-2] if bold else s[1:-1] if ital else s);r.bold=bold;r.italic=ital;r.font.color.rgb=RGBColor(0,0,0)
doc.save(O/'Manuscript.docx')
os.environ['PATH']=str(R/'work/tools/lo/LibreOffice/program')+os.pathsep+'${LOCAL_USER_HOME}/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin'+os.pathsep+os.environ['PATH']
renderer='${LOCAL_USER_HOME}/.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents/render_docx.py'
for name in ['Manuscript','Supplementary_Material','STROBE_MR_Checklist']:
 folder=W/'render'/name;assert not list(folder.glob('page-*.png'))
 sys.argv=[renderer,str(O/(name+'.docx')),'--output_dir',str(folder),'--emit_pdf'];runpy.run_path(renderer,run_name='__main__')
 print(name,len(list(folder.glob('page-*.png'))),'pages',flush=True)

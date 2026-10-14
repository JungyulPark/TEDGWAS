from pathlib import Path
from docx import Document
import os,sys,runpy
W=Path(__file__).resolve().parent;R=W.parents[1];O=R/'outputs/TED_Tissue_Integrated_20261003'
sys.stdout.reconfigure(encoding='utf-8')
doc=Document(O/'Manuscript.docx');md=(O/'Manuscript.md').read_text(encoding='utf-8')
assert '*IGF1R* estimates consistently indicated higher odds' in md.split('## Introduction')[0]
md=md.replace('*IGF1R* estimates consistently indicated higher odds','Blood *IGF1R* estimates consistently indicated higher odds',1)
p=next(p for p in doc.paragraphs if 'IGF1R estimates consistently indicated higher odds' in p.text)
ir=next(i for i,r in enumerate(p.runs) if r.text=='IGF1R' and 'estimates consistently' in ''.join(t.text for t in p.runs[i+1:]))
p.runs[ir-1].text+='Blood '
changes=[
 ('GTEx is predominantly European ancestry, not exclusively European.','GTEx participants were predominantly of European ancestry, not exclusively European.'),
 ('Complete regional summaries were unavailable for this extension, so tissue colocalization was not performed;','Complete regional summaries were not obtained for this extension, so tissue colocalization was not performed;'),
 ('blood *cis*-eQTL data from the eQTLGen Consortium,','blood *cis*-eQTL data from the eQTLGen Consortium, tissue eQTL data from GTEx v8,'),
 ('We thank the eQTLGen Consortium, Biobank Japan,','We thank the eQTLGen Consortium, the GTEx Consortium, Biobank Japan,')
]
for old,new in changes:
 assert md.count(old)==1,old
 md=md.replace(old,new)
 plainold=old.replace('*','');plainnew=new.replace('*','')
 p=next(p for p in doc.paragraphs if plainold in p.text)
 # Insert/replace only text within its existing run, retaining surrounding formatting.
 if old.startswith('blood '):
  old='eQTL data from the eQTLGen Consortium,';new='eQTL data from the eQTLGen Consortium, tissue eQTL data from GTEx v8,'
 else:old=plainold;new=plainnew
 found=False
 for r in p.runs:
  if old in r.text:r.text=r.text.replace(old,new);found=True;break
 assert found,old
(O/'Manuscript.md').write_text(md,encoding='utf-8');doc.save(O/'Manuscript.docx')
os.environ['PATH']=str(R/'work/tools/lo/LibreOffice/program')+os.pathsep+'${LOCAL_USER_HOME}/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin'+os.pathsep+os.environ['PATH']
renderer='${LOCAL_USER_HOME}/.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents/render_docx.py'
folder=W/'render_release/Manuscript';assert not folder.exists()
sys.argv=[renderer,str(O/'Manuscript.docx'),'--output_dir',str(folder),'--emit_pdf'];runpy.run_path(renderer,run_name='__main__')

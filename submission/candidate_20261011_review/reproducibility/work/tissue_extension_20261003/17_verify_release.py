from pathlib import Path
from zipfile import ZipFile
from docx import Document
from pypdf import PdfReader
import pandas as pd
import re,json,hashlib,difflib,math,sys
sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parents[1]
O=R/'outputs/TED_Tissue_Integrated_20261003';A=R/'outputs/TED_Tissue_Extension_20261003';S=R/'outputs/TED_Allele_Integrated_20261003';V=O/'Verification'
checks=[]
def check(name,ok):
 checks.append({'check':name,'pass':bool(ok)});assert ok,name
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
md=(O/'Manuscript.md').read_text(encoding='utf-8');doc=Document(O/'Manuscript.docx')
dt='\n\n'.join(p.text for p in doc.paragraphs)
clean=lambda t:re.sub(r'\s+',' ',t.replace('*','').replace('##','').replace('#','').replace('\\','')).strip()
abstract=md.split('## Abstract')[1].split('**Keywords:**')[0]
words=len(re.findall(r'\S+',abstract.replace('*','')))
check('Abstract <=250 words',words<=250)
main=md.split('## Introduction')[1].split('## Declarations')[0]
mainwords=len(re.findall(r'\S+',re.sub(r'#+[^\n]+\n','',main).replace('*','')))
check('Main text <=5000 words',mainwords<=5000)
discussion=md.split('## Discussion')[1].split('### Limitations')[0]
check('Eight discussion paragraphs',len([p for p in discussion.strip().split('\n\n') if p.strip()])==8)
lim=md.split('### Limitations')[1].split('### Conclusions')[0]
check('Two limitation paragraphs',len([p for p in lim.strip().split('\n\n') if p.strip()])==2)
before,refs=md.split('## References')
seen=[]
for match in re.finditer(r'\[([\d, ]+)\]',before):
 for n in map(int,match.group(1).split(',')):
  if n not in seen:seen.append(n)
check('All 33 references first-cited in order',seen==list(range(1,34)))
check('33 numbered reference records',re.findall(r'^(\d+)\. ',refs,re.M)==list(map(str,range(1,34))))
for p in [p for p in md.split('\n\n') if p and not p.startswith('#')]:
 if any(k in p for k in ['GTEx','Adipose *TSHR*','0.36, 95% CI','thyroid/fibroblast','Detection_original']) or re.match(r'^\d+\. ',p):
  check('DOCX/MD text: '+p[:75],clean(p) in clean(dt))
check('Abstract tissue scope and blood scope explicit','Adipose TSHR–GD associations passed correction' in dt and 'Blood IGF1R estimates consistently' in dt)
for name in ['Instrument_checks.json','MR_checks.json']:
 c=json.loads((A/name).read_text());check(name+' passed',c.get('all_pass',all(x['pass'] for x in c['checks'])))
mr=pd.read_csv(A/'Tables/Tissue_MR_50_planned_combinations.csv');av=pd.read_csv(A/'Tables/Tissue_availability_10_combinations.csv');rv=pd.read_csv(A/'R_independent_verification.csv')
check('50 cells; 20 estimable; 30 unavailable',len(mr)==50 and (mr.status=='estimated').sum()==20 and (mr.status=='not_estimable').sum()==30)
check('20 independent R estimates pass',len(rv)==20 and rv.passed.all() and rv.max_relative_difference.max()<1e-12)
check('Ten primary / forty secondary denominators',(mr.family=='primary').sum()==10 and (mr.family=='secondary').sum()==40 and set(mr.planned_family_tests)=={10,40})
check('TSHR eight corrected, one nominal, one nonsignificant',mr[(mr.gene=='TSHR')&(mr.status=='estimated')].statistical_category.value_counts().to_dict()=={'family_corrected':8,'nominal_only':1,'not_nominally_significant':1})
check('All 10 IGF1R P values >0.05',(mr[(mr.gene=='IGF1R')&(mr.status=='estimated')].P>.05).all())
check('Four selected instruments, one per estimable combination',len(pd.read_csv(A/'Tables/Selected_instruments_verified.csv'))==4 and mr.loc[mr.status=='estimated','n_retained_instruments'].eq(1).all())
check('TSHR thyroid not relaxed',av[(av.gene=='TSHR')&(av.tissue=='Thyroid')].iloc[0].minimum_eQTL_P>5e-8)
supp=Document(O/'Supplementary_Material.docx');oldsupp=Document(S/'Supplementary_Material.docx')
check('Exactly three supplementary tables',len(supp.tables)==3)
for i in [0,1]:check('Supplementary Table S'+str(i+1)+' unchanged',supp.tables[i]._tbl.xml==oldsupp.tables[i]._tbl.xml)
st=supp.tables[2];check('S3 includes ten planned combinations',len(st.rows)==11)
tn={'Thyroid':'Thyroid','Cells_Cultured_fibroblasts':'Cultured fibroblasts','Adipose_Subcutaneous':'Subcutaneous adipose','Adipose_Visceral_Omentum':'Visceral adipose','Whole_Blood':'Whole blood'}
for i,(_,r) in enumerate(mr[mr.family=='primary'].iterrows(),1):
 cells=[c.text for c in st.rows[i].cells]
 check('S3 identity '+str(i),cells[0]==r.gene+' / '+tn[r.tissue] and int(cells[1])==r.n_exposure and int(cells[2])==r.n_retained_instruments)
 if r.status=='estimated':
  nums=list(map(float,re.findall(r'\d+\.\d+',cells[3])))
  check('S3 OR/CI '+str(i),nums==[round(r.OR,2),round(r.CI_lower,2),round(r.CI_upper,2)])
  for cell,v in [(cells[4],r.P),(cells[5],r.P_adjusted)]:check('S3 P '+str(i)+' '+cell,math.isclose(float(cell),v,rel_tol=.006,abs_tol=.0005 if v>.01 else 0))
 else:check('S3 unavailable '+str(i),'No qualifying IV' in cells[3] or 'Not tested' in cells[3])
for fn in ['Tables.docx','Tables.pdf','Tables.md','Figure_Legends.md','All_Figures.pdf']:
 check(fn+' unchanged',sha(O/fn)==sha(S/fn))
for p in (S/'Figures').iterdir():check(p.name+' unchanged',sha(p)==sha(O/'Figures'/p.name))
with ZipFile(S/'Supplementary_Material.docx') as old,ZipFile(O/'Supplementary_Material.docx') as new:
 for n in old.namelist():
  if n.startswith('word/media/'):check('Supplementary figure media '+n+' unchanged',old.read(n)==new.read(n))
check('Old detection data retained',json.loads((V/'workbook_checks.json').read_text())['unchanged_data_worksheets']==18)
for name,folder in [('Manuscript',W/'render_release/Manuscript'),('Supplementary_Material',W/'render/Supplementary_Material'),('STROBE_MR_Checklist',W/'render/STROBE_MR_Checklist')]:
 pdf=PdfReader(folder/(name+'.pdf'));text='\n'.join(p.extract_text() for p in pdf.pages)
 check(name+' has no blank PDF pages',all(len(p.extract_text().strip())>15 for p in pdf.pages))
 check(name+' has no replacement-glyph text','\ufffd' not in text)
 if name=='Manuscript':check('Rendered manuscript contains new results','Adipose TSHR' in text and 'Tissue-specific extension' in text)
 if name=='Supplementary_Material':check('Rendered Table S3 contains estimates','1.28e-34' in text and '4.09e-19' in text)
report={'status':'PASS','checks':checks,'abstract_words':words,'main_words':mainwords,'discussion_paragraphs':8,'limitations_paragraphs':2,'references':33,'independent_R_max_relative_difference':rv.max_relative_difference.max(),'document_pages':{'Manuscript':15,'Supplementary_Material':9,'STROBE_MR_Checklist':1},'validation_scope':'Numerical implementation, source identities, preservation, citation order and rendering; not proof of instrumental validity or tissue mediation.'}
(V/'final_checks.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
for fn in ['Manuscript.md','Supplementary_Material.md']:
 diff=difflib.unified_diff((S/fn).read_text(encoding='utf-8').splitlines(True),(O/fn).read_text(encoding='utf-8').splitlines(True),fromfile='previous/'+fn,tofile='tissue_extension/'+fn)
 (V/(Path(fn).stem+'_changes.diff')).write_text(''.join(diff),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='checks'},indent=2));print('Checks:',len(checks))
